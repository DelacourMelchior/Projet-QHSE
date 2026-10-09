# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_RENTABILITE_DELACOUR"
TEMPLATES_DIR = os.path.join(BASE_DIR, "01_TEMPLATES_VIERGES")
GUIDES_DIR = os.path.join(BASE_DIR, "02_GUIDES_METHODOLOGIQUES")

os.makedirs(TEMPLATES_DIR, exist_ok=True)
os.makedirs(GUIDES_DIR, exist_ok=True)

# BRAND COLORS
BRAND_GREEN = RGBColor(27, 77, 62)     # #1B4D3E
DARK_GREEN = RGBColor(21, 40, 34)      # #152822
LIGHT_GREEN = RGBColor(74, 108, 92)    # #4A6C5C
GOLD = RGBColor(197, 160, 101)         # #C5A065
WHITE = RGBColor(255, 255, 255)
GREEN_TEXT = RGBColor(27, 94, 32)      # #1B5E20
RED_TEXT = RGBColor(183, 28, 28)       # #B71C1C
MUTED_TEXT = RGBColor(120, 120, 120)

CREAM_BG = "F9F9F7"
WHY_BG = "FDFBF7"
CASH_BG = "F4F8F5"
HEADER_BG = "1B4D3E"
SUBHDR_BG = "E8F0EC"
BORDER_ACCENT = "E2DCCC"

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

HEADER_IMG_PATH = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def create_base_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)
        
        # Header / Footer
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_h = p_head.add_run("CABINET DELACOUR - PERFORMANCE & RENTABILITE PME")
        r_h.font.name = FONT_SANS
        r_h.font.size = Pt(7.5)
        r_h.font.color.rgb = MUTED_TEXT

        footer = section.footer
        p_ft = footer.paragraphs[0]
        pPr = p_ft._p.get_or_add_pPr()
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right')
        tab.set(qn('w:pos'), '10512')
        tabs.append(tab)
        pPr.append(tabs)
        
        r_l = p_ft.add_run("Cabinet Delacour - contact@cabinetdelacour.com")
        r_l.font.name = FONT_SANS
        r_l.font.size = Pt(7.5)
        r_l.font.color.rgb = MUTED_TEXT
        
        r_tab = p_ft.add_run("\tPage ")
        r_tab.font.name = FONT_SANS
        r_tab.font.size = Pt(7.5)
        r_tab.font.color.rgb = MUTED_TEXT
        
        fld_pg = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        p_ft._p.append(fld_pg)
    return doc

def add_header_brand(doc, tool_code, tool_title, doc_type_label):
    if os.path.exists(HEADER_IMG_PATH):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(HEADER_IMG_PATH, width=Inches(7.35))
    
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_title.autofit = False
    c_title = tbl_title.cell(0, 0)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=70, bottom=70, left=120, right=120)
    set_cell_border(c_title, left={'sz': 24, 'val': 'single', 'color': 'C5A065'},
                             top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p_title = c_title.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(2)
    r_cl = p_title.add_run(f"PROGRAMME CASH & RENTABILITE DELACOUR - {tool_code} - {doc_type_label}\n")
    r_cl.font.name = FONT_SANS
    r_cl.font.size = Pt(8)
    r_cl.font.bold = True
    r_cl.font.color.rgb = GOLD
    
    r_main = p_title.add_run(tool_title)
    r_main.font.name = FONT_SERIF
    r_main.font.size = Pt(13)
    r_main.font.bold = True
    r_main.font.color.rgb = BRAND_GREEN
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(3)

def add_section_header(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title_text)
    r.font.name = FONT_SERIF
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = BRAND_GREEN

def add_cash_counter_box(doc, enjeu_description, indicateur_cle, formule_calcul=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    set_explicit_cell_width(c, 10200)
    set_cell_background(c, CASH_BG)
    set_cell_margins(c, top=60, bottom=60, left=100, right=100)
    set_cell_border(c, left={'sz': 20, 'val': 'single', 'color': '1B4D3E'},
                      top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run("ENJEU OPERATIONNEL & IMPACT FINANCIER REEL : ")
    r1.font.name = FONT_SANS
    r1.font.size = Pt(8)
    r1.font.bold = True
    r1.font.color.rgb = BRAND_GREEN
    r2 = p.add_run(enjeu_description)
    r2.font.name = FONT_SANS
    r2.font.size = Pt(8)
    
    p_ind = c.add_paragraph()
    p_ind.paragraph_format.space_after = Pt(0)
    r_ind_lbl = p_ind.add_run("Indicateur cle de suivi : ")
    r_ind_lbl.font.name = FONT_SANS
    r_ind_lbl.font.size = Pt(7.5)
    r_ind_lbl.font.bold = True
    r_ind_lbl.font.color.rgb = BRAND_GREEN
    r_ind_val = p_ind.add_run(f"{indicateur_cle}   |   Montant reel constate entreprise : [ ________ EUR ]")
    r_ind_val.font.name = FONT_SANS
    r_ind_val.font.size = Pt(7.5)
    r_ind_val.font.color.rgb = GREEN_TEXT
    
    if formule_calcul:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r_f = p2.add_run(f"Formule de calcul du ROI : {formule_calcul}")
        r_f.font.name = FONT_SANS
        r_f.font.size = Pt(7.5)
        r_f.font.italic = True
        r_f.font.color.rgb = LIGHT_GREEN
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)

def add_metadata_box(doc, fields_dict):
    tbl = doc.add_table(rows=len(fields_dict), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_w = [2600, 7600]
    for idx, (lbl, val) in enumerate(fields_dict.items()):
        c0, c1 = tbl.cell(idx, 0), tbl.cell(idx, 1)
        set_explicit_cell_width(c0, col_w[0])
        set_explicit_cell_width(c1, col_w[1])
        set_cell_background(c0, "F0F4F2")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=35, bottom=35, left=60, right=60)
        set_cell_margins(c1, top=35, bottom=35, left=60, right=60)
        set_cell_border(c0, top={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            bottom={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            left={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            right={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
        set_cell_border(c1, top={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            bottom={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            left={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                            right={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.font.name = FONT_SANS
        r0.font.size = Pt(7.5)
        r0.font.bold = True
        r0.font.color.rgb = BRAND_GREEN
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = FONT_SANS
        r1.font.size = Pt(7.5)
        if "[" in val or "____" in val:
            r1.font.italic = True
            r1.font.color.rgb = MUTED_TEXT
            
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(3)

def format_word_table(doc, headers, rows_data, col_widths_dxa=None, subheaders=None, center_cols=None):
    tbl = doc.add_table(rows=0, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row 0
    hdr_cells0 = tbl.add_row().cells
    for i, text in enumerate(headers):
        hdr_cells0[i].text = text
        if col_widths_dxa and i < len(col_widths_dxa):
            set_explicit_cell_width(hdr_cells0[i], col_widths_dxa[i])
        set_cell_background(hdr_cells0[i], HEADER_BG)
        set_cell_margins(hdr_cells0[i], top=50, bottom=50, left=60, right=60)
        set_cell_border(hdr_cells0[i], top={'sz': 4, 'val': 'single', 'color': HEADER_BG},
                                      bottom={'sz': 4, 'val': 'single', 'color': HEADER_BG},
                                      left={'sz': 4, 'val': 'single', 'color': HEADER_BG},
                                      right={'sz': 4, 'val': 'single', 'color': HEADER_BG})
        for p in hdr_cells0[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE
                
    if subheaders:
        hdr_cells1 = tbl.add_row().cells
        for i, text in enumerate(subheaders):
            hdr_cells1[i].text = text
            if col_widths_dxa and i < len(col_widths_dxa):
                set_explicit_cell_width(hdr_cells1[i], col_widths_dxa[i])
            set_cell_background(hdr_cells1[i], SUBHDR_BG)
            set_cell_margins(hdr_cells1[i], top=40, bottom=40, left=60, right=60)
            for p in hdr_cells1[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7)
                    r.font.bold = True
                    r.font.color.rgb = BRAND_GREEN
                    
    for r_idx, row_data in enumerate(rows_data):
        row_cells = tbl.add_row().cells
        bg_hex = CREAM_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val
            if col_widths_dxa and c_idx < len(col_widths_dxa):
                set_explicit_cell_width(row_cells[c_idx], col_widths_dxa[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=45, bottom=45, left=60, right=60)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                if center_cols and c_idx in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)
                    if "[" in val or "____" in val:
                        r.font.italic = True
                        r.font.color.rgb = MUTED_TEXT
                    elif any(w in val for w in ["OUI", "CONFORME", "VALIDE", "GAGNE", "CRITIQUE", "LEVE"]):
                        if "CRITIQUE" in val:
                            r.font.bold = True
                            r.font.color.rgb = RED_TEXT
                        else:
                            r.font.bold = True
                            r.font.color.rgb = GREEN_TEXT
                    elif any(w in val for w in ["NON", "REFUSE", "ALERTE", "PERDU", "BLOQUANT"]):
                        r.font.bold = True
                        r.font.color.rgb = RED_TEXT
                        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(3)

def add_validation_signoff_box(doc, signoff_roles):
    tbl = doc.add_table(rows=len(signoff_roles)+1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_w = [2500, 2700, 2500, 2500]
    hdrs = ["ROLE / ACTEUR", "NOM & FONCTION", "DATE & DECISION", "VISA / SIGNATURE"]
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(hdrs):
        hdr_cells[i].text = h
        set_explicit_cell_width(hdr_cells[i], col_w[i])
        set_cell_background(hdr_cells[i], HEADER_BG)
        set_cell_margins(hdr_cells[i], top=40, bottom=40, left=50, right=50)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(7.5)
                r.font.bold = True
                r.font.color.rgb = WHITE
                
    for idx, role in enumerate(signoff_roles):
        row_cells = tbl.rows[idx+1].cells
        row_cells[0].text = role
        row_cells[1].text = "[Nom du responsable]"
        row_cells[2].text = "[Date] - [Bon pour accord]"
        row_cells[3].text = "[Signature / Cachet]"
        for c_idx in range(4):
            set_explicit_cell_width(row_cells[c_idx], col_w[c_idx])
            set_cell_background(row_cells[c_idx], "FFFFFF")
            set_cell_margins(row_cells[c_idx], top=45, bottom=45, left=50, right=50)
            set_cell_border(row_cells[c_idx], top={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)
                    if "[" in r.text:
                        r.font.italic = True
                        r.font.color.rgb = MUTED_TEXT
                    else:
                        r.font.bold = True
                        r.font.color.rgb = BRAND_GREEN
                        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)

def add_next_steps_box(doc, steps):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    set_explicit_cell_width(c, 10200)
    set_cell_background(c, WHY_BG)
    set_cell_margins(c, top=60, bottom=60, left=100, right=100)
    set_cell_border(c, left={'sz': 20, 'val': 'single', 'color': 'C5A065'},
                      top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r_t = p.add_run("ACTION IMMEDIATE APRES COMPLETION DU FORMULAIRE :")
    r_t.font.name = FONT_SANS
    r_t.font.size = Pt(8)
    r_t.font.bold = True
    r_t.font.color.rgb = GOLD
    for s in steps:
        p_s = c.add_paragraph()
        p_s.paragraph_format.space_after = Pt(1)
        r_s = p_s.add_run(f"- {s}")
        r_s.font.name = FONT_SANS
        r_s.font.size = Pt(7.5)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(3)

# GUIDE SPECIFIC HELPERS
def add_why_box(doc, gains_list, pertes_list):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    c_title = tbl.cell(0, 0)
    c_title.merge(tbl.cell(0, 1))
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, WHY_BG)
    set_cell_margins(c_title, top=70, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("POURQUOI CET OUTIL TRANSFORME LA RENTABILITE DE L'ENTREPRISE")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN
    
    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, WHY_BG)
    set_cell_margins(c1, top=50, bottom=70, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("Ce que l'entreprise GAGNE (Cash direct) :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = GREEN_TEXT
    for g in gains_list:
        p = c1.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"- {g}")
        r.font.name = FONT_SANS
        r.font.size = Pt(8)
        
    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, WHY_BG)
    set_cell_margins(c2, top=50, bottom=70, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Ce que l'entreprise PERD sans cet outil (Gaspillages) :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RED_TEXT
    for pe in pertes_list:
        p = c2.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"- {pe}")
        r.font.name = FONT_SANS
        r.font.size = Pt(8)
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_consultant_box(doc, posture_text, rules_list):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    c_head = tbl.cell(0, 0)
    set_explicit_cell_width(c_head, 10200)
    set_cell_background(c_head, CREAM_BG)
    set_cell_margins(c_head, top=70, bottom=50, left=120, right=120)
    set_cell_border(c_head, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_h = c_head.paragraphs[0]
    p_h.paragraph_format.space_after = Pt(2)
    r_t = p_h.add_run("CONSEIL CONSULTANT & POSTURE TERRAIN (RECUPERATION DU CASH)")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN
    
    r_sub = p_h.add_run(f"\n{posture_text}")
    r_sub.font.name = FONT_SANS
    r_sub.font.size = Pt(8)
    r_sub.font.italic = True
    
    c_body = tbl.cell(1, 0)
    set_explicit_cell_width(c_body, 10200)
    set_cell_background(c_body, CREAM_BG)
    set_cell_margins(c_body, top=50, bottom=70, left=120, right=120)
    set_cell_border(c_body, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                           top={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
    p_b = c_body.paragraphs[0]
    p_b.paragraph_format.space_after = Pt(1)
    for title, desc in rules_list:
        p = c_body.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"- {title} : ")
        r1.font.name = FONT_SANS
        r1.font.size = Pt(8)
        r1.font.bold = True
        r1.font.color.rgb = BRAND_GREEN
        r2 = p.add_run(desc)
        r2.font.name = FONT_SANS
        r2.font.size = Pt(8)
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_guide_box(doc, steps_list):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    set_explicit_cell_width(c, 10200)
    set_cell_background(c, "F5F7F6")
    set_cell_margins(c, top=70, bottom=70, left=120, right=120)
    set_cell_border(c, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                      top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                      right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run("PROTOCOLE DE DEPLOIEMENT PAS-A-PAS (REGLE DES 2 MINUTES & CHRONO)")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(9)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN
    
    for i, st in enumerate(steps_list):
        p_st = c.add_paragraph()
        p_st.paragraph_format.space_after = Pt(1)
        r_num = p_st.add_run(f"Etape {i+1} : ")
        r_num.font.name = FONT_SANS
        r_num.font.bold = True
        r_num.font.size = Pt(8)
        r_num.font.color.rgb = BRAND_GREEN
        r_desc = p_st.add_run(st)
        r_desc.font.name = FONT_SANS
        r_desc.font.size = Pt(8)
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_case_study_box(doc, entreprise, situation, action, roi_text):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    c_head = tbl.cell(0, 0)
    set_explicit_cell_width(c_head, 10200)
    set_cell_background(c_head, "F4F6F5")
    set_cell_margins(c_head, top=60, bottom=50, left=100, right=100)
    set_cell_border(c_head, left={'sz': 16, 'val': 'single', 'color': '1B4D3E'},
                           top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                           right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                           bottom={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
    p_h = c_head.paragraphs[0]
    p_h.paragraph_format.space_after = Pt(1)
    r_t = p_h.add_run(f"CAS PRATIQUE REEL PME : {entreprise}")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(9)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN
    
    c_b = tbl.cell(1, 0)
    set_explicit_cell_width(c_b, 10200)
    set_cell_background(c_b, "FFFFFF")
    set_cell_margins(c_b, top=50, bottom=60, left=100, right=100)
    set_cell_border(c_b, left={'sz': 16, 'val': 'single', 'color': '1B4D3E'},
                         bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 2, 'val': 'single', 'color': BORDER_ACCENT})
    
    p1 = c_b.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("1. Situation Initiale (Fuite de cash) : ")
    r1.font.bold = True
    r1.font.size = Pt(8)
    r1.font.color.rgb = RED_TEXT
    r1_t = p1.add_run(situation)
    r1_t.font.size = Pt(8)
    
    p2 = c_b.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("2. Action Commando Implante : ")
    r2.font.bold = True
    r2.font.size = Pt(8)
    r2.font.color.rgb = BRAND_GREEN
    r2_t = p2.add_run(action)
    r2_t.font.size = Pt(8)
    
    p3 = c_b.add_paragraph()
    p3.paragraph_format.space_after = Pt(1)
    r3 = p3.add_run("3. ROI Financier Concret Mesure : ")
    r3.font.bold = True
    r3.font.size = Pt(8)
    r3.font.color.rgb = GREEN_TEXT
    r3_t = p3.add_run(roi_text)
    r3_t.font.size = Pt(8.5)
    r3_t.font.bold = True
    r3_t.font.color.rgb = GREEN_TEXT
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(4)

def add_audit_checklist_box(doc, criteria_list):
    headers = ["CRITERE D'EVALUATION DE L'OUTIL", "CIBLE ATTENDUE", "CONFORME ?", "ACTION CORRECTIVE SI NON"]
    col_w = [4200, 2400, 1400, 2200]
    rows = []
    for crit, cible, action in criteria_list:
        rows.append([crit, cible, "[ OUI / NON ]", action])
    format_word_table(doc, headers, rows, col_widths_dxa=col_w, center_cols=[2])

def add_apple_header(doc, tool_code, apple_name, punchline):
    if os.path.exists(HEADER_IMG_PATH):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(HEADER_IMG_PATH, width=Inches(7.35))
    
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_title.autofit = False
    c_title = tbl_title.cell(0, 0)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=50, bottom=50, left=120, right=120)
    set_cell_border(c_title, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                             top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p_title = c_title.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(1)
    r_cl = p_title.add_run(f"CABINET DELACOUR — INSTRUMENT DE RENTABILITÉ #{tool_code}\n")
    r_cl.font.name = FONT_SANS
    r_cl.font.size = Pt(8)
    r_cl.font.bold = True
    r_cl.font.color.rgb = GOLD
    
    r_main = p_title.add_run(apple_name + " ")
    r_main.font.name = FONT_SERIF
    r_main.font.size = Pt(13)
    r_main.font.bold = True
    r_main.font.color.rgb = BRAND_GREEN
    
    r_punch = p_title.add_run(f"— « {punchline} »")
    r_punch.font.name = FONT_SERIF
    r_punch.font.size = Pt(9.5)
    r_punch.font.italic = True
    r_punch.font.color.rgb = MUTED_TEXT
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(2)

def add_apple_protocol_page(doc, regles_dor, scenario_3_etapes, pieges_mortels, test_30s):
    doc.add_page_break()
    
    if os.path.exists(HEADER_IMG_PATH):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(HEADER_IMG_PATH, width=Inches(7.35))
        
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(4)
    p_h.paragraph_format.space_after = Pt(2)
    r_h = p_h.add_run("LE PROTOCOLE TERRAIN & LES RÉFLEXES DE SURVIE (GUIDAGE INTÉGRÉ)")
    r_h.font.name = FONT_SERIF
    r_h.font.size = Pt(11)
    r_h.font.bold = True
    r_h.font.color.rgb = BRAND_GREEN
    
    # 1. Les 3 Règles d'Or
    add_section_header(doc, "1. LES 3 RÈGLES D'OR NON NÉGOCIABLES")
    tbl_r = doc.add_table(rows=len(regles_dor), cols=2)
    tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_r.autofit = False
    col_w_r = [2200, 8000]
    for i, (titre, desc) in enumerate(regles_dor):
        c0, c1 = tbl_r.cell(i, 0), tbl_r.cell(i, 1)
        set_explicit_cell_width(c0, col_w_r[0])
        set_explicit_cell_width(c1, col_w_r[1])
        set_cell_background(c0, HEADER_BG)
        set_cell_background(c1, "FFFFFF" if i%2==0 else CREAM_BG)
        set_cell_margins(c0, 30, 30, 60, 60)
        set_cell_margins(c1, 30, 30, 80, 80)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(titre)
        r0.font.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = BRAND_GREEN
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(desc)
        r1.font.size = Pt(8)
    
    # 2. Le Scénario en 3 Étapes Chrono
    add_section_header(doc, "2. LE SCÉNARIO TERRAIN EN 3 ÉTAPES (RÈGLE DES 2 MINUTES)")
    tbl_s = doc.add_table(rows=len(scenario_3_etapes), cols=2)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s.autofit = False
    col_w_s = [2200, 8000]
    for i, (etape_num, action_text) in enumerate(scenario_3_etapes):
        c0, c1 = tbl_s.cell(i, 0), tbl_s.cell(i, 1)
        set_explicit_cell_width(c0, col_w_s[0])
        set_explicit_cell_width(c1, col_w_s[1])
        set_cell_background(c0, SUBHDR_BG)
        set_cell_background(c1, "FFFFFF" if i%2==0 else CREAM_BG)
        set_cell_margins(c0, 30, 30, 60, 60)
        set_cell_margins(c1, 30, 30, 80, 80)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(etape_num)
        r0.font.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = BRAND_GREEN
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(action_text)
        r1.font.size = Pt(8)
        
    # 3. Pièges Mortels & Test 30s
    add_section_header(doc, "3. LES 3 PIÈGES MORTELS & LE TEST FLASH DU DIRIGEANT")
    tbl_p = doc.add_table(rows=2, cols=1)
    tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p.autofit = False
    cp0, cp1 = tbl_p.cell(0, 0), tbl_p.cell(1, 0)
    set_explicit_cell_width(cp0, 10200)
    set_explicit_cell_width(cp1, 10200)
    set_cell_background(cp0, WHY_BG)
    set_cell_background(cp1, CASH_BG)
    set_cell_margins(cp0, 40, 40, 80, 80)
    set_cell_margins(cp1, 40, 40, 80, 80)
    set_cell_border(cp0, left={'sz': 16, 'val': 'single', 'color': '8A2BE2'})
    set_cell_border(cp1, left={'sz': 16, 'val': 'single', 'color': '1B4D3E'})
    
    p_pieges = cp0.paragraphs[0]
    p_pieges.paragraph_format.space_after = Pt(1)
    rp_t = p_pieges.add_run("PIÈGES MORTELS À ÉVITER : ")
    rp_t.font.bold = True
    rp_t.font.size = Pt(8)
    rp_t.font.color.rgb = RED_TEXT
    for piege in pieges_mortels:
        pp = cp0.add_paragraph()
        pp.paragraph_format.space_after = Pt(1)
        rpp = pp.add_run(f"• {piege}")
        rpp.font.size = Pt(7.8)
        
    p_test = cp1.paragraphs[0]
    p_test.paragraph_format.space_after = Pt(0)
    rpt_t = p_test.add_run("LE TEST FLASH DE 30 SECONDES DU DIRIGEANT : ")
    rpt_t.font.bold = True
    rpt_t.font.size = Pt(8)
    rpt_t.font.color.rgb = BRAND_GREEN
    rpt_v = p_test.add_run(test_30s)
    rpt_v.font.size = Pt(8)

print("commando_style.py initialized.")
