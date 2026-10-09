#!/usr/bin/env python3
"""
Script d'exécution : Remplissage de la Grille d'Audit SMQ pour TechIndustrie SAS
Conforme à l'architecture 3 couches Delacour.
Lit la trame DOCX originale et remplit toutes les données d'audit factuelles
extraites du document d'audit (Impôt Invisible de 907 866 € / an).
"""

import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_runs(cell, text, font_name="Arial", font_size_pt=8.5, bold=False, color_rgb=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    """
    Définit le texte d'une cellule docx en conservant un style propre et homogène.
    """
    cell.text = ""  # Vider le contenu existant
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)

def add_cell_paragraph(cell, text, font_name="Arial", font_size_pt=8.5, bold=False, color_rgb=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    """
    Ajoute un paragraphe dans une cellule.
    """
    p = cell.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)
    return p

def main():
    source_path = "OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE/03_INTERVENTION_SUR_SITE_JOUR_J/GRILLE_D_AUDIT_SMQ_DELACOUR.docx"
    output_dir = "DOSSIER_AUDIT_TECHINDUSTRIE_SAS"
    os.makedirs(output_dir, exist_ok=True)
    target_path = os.path.join(output_dir, "GRILLE_D_AUDIT_SMQ_DELACOUR.docx")

    print(f"Chargement de la trame source : {source_path}")
    doc = docx.Document(source_path)

    # 1. TABLE 0 : En-tête de cadrage
    t0 = doc.tables[0]
    set_cell_runs(t0.rows[0].cells[0], "Société Client : TechIndustrie SAS", font_size_pt=9, bold=True)
    set_cell_runs(t0.rows[0].cells[1], "Date d'audit : 18/08/2026", font_size_pt=9, bold=True)
    set_cell_runs(t0.rows[1].cells[0], "Dirigeant / Décideur d'entreprise : Dirigeant TechIndustrie SAS (50 collab. / CA 3,5 M€)", font_size_pt=9, bold=True)
    set_cell_runs(t0.rows[1].cells[1], "Auditeur référent : Cabinet Delacour (Stratégie & Performance QHSE)", font_size_pt=9, bold=True)

    # Données des 29 questions d'audit (Clauses 4 à 10) pour TechIndustrie SAS
    audit_data = {
        # CLAUSE 4 : CONTEXTE DE L'ORGANISME
        "4.1": {
            "comment": "Aucune analyse PESTEL/SWOT chiffrée. Absence de certification ISO 9001 bloquant l'accès au marché médical grands comptes (300 k€ de perte de marché) et opportunités de croissance bridées à 150 k€ par manque de structuration.",
            "note": "0"
        },
        "4.2": {
            "comment": "Cartographie des parties prenantes non formalisée. Clients historiques médical bien servis mais relations avec sous-traitants rang 2 défaillantes (cahiers des charges flous générant 45 k€/an d'avoirs et litiges).",
            "note": "1"
        },
        "4.3": {
            "comment": "Périmètre d'activité d'usinage et sous-traitance industrielle compris empiriquement mais non documenté. Tâches en doublon et zones de flou d'attribution.",
            "note": "1"
        },
        "4.4": {
            "comment": "Goulots d'étranglement fréquents entre bureau d'études et production. Absence de cartographie des flux générant des retards et 5% de retouches en atelier (coût: 7 000 €/an).",
            "note": "1"
        },

        # CLAUSE 5 : LEADERSHIP & ENGAGEMENT
        "5.1": {
            "comment": "Le Dirigeant est 'sous l'eau' dans du micro-management quotidien (15h/sem. perdues, impact 141 000 €/an). Absence de tableau de bord consolidé et de revues de direction formelles.",
            "note": "0"
        },
        "5.2": {
            "comment": "Pas de Politique Qualité rédigée ni diffusée. Démarche qualité perçue par les équipes comme une contrainte administrative plutôt qu'un levier économique.",
            "note": "0"
        },
        "5.3": {
            "comment": "Organigramme non mis à jour, fiches de poste incomplètes et pas de grille RACI. Forte dépendance exclusive envers le Dirigeant et 1 cadre clé de production.",
            "note": "1"
        },

        # CLAUSE 6 : PLANIFICATION
        "6.1": {
            "comment": "Pas de registre des risques formalisé ni de cotation financière des vulnérabilités. L'Impôt Invisible de 907 866 €/an (2 487 €/jour) n'était pas mesuré jusqu'à présent.",
            "note": "0"
        },
        "6.2": {
            "comment": "Objectifs opérationnels vagues et non reliés à la rentabilité. Retards de traitement comptable provoquant un décalage de 15 jours sur la clôture mensuelle (impact: 92 055 €/an).",
            "note": "1"
        },
        "6.3": {
            "comment": "Évolutions d'organisation ou de lignes de fabrication réalisées à chaud sans méthodologie de conduite du changement ni mise à jour des standards.",
            "note": "1"
        },

        # CLAUSE 7 : SUPPORT & RESSOURCES
        "7.1": {
            "comment": "Maintenance préventive partielle. L'organisation subit de plein fouet un taux d'absentéisme de 8,0% générant désorganisation et surcoûts (126 000 €/an).",
            "note": "1"
        },
        "7.2": {
            "comment": "Matrice des compétences absente. Intégration informelle des nouveaux. Remplacement coûteux de 1 cadre et 4 opérationnels sur l'année (turnover: 32 250 €/an).",
            "note": "1"
        },
        "7.3": {
            "comment": "Les collaborateurs font 'au mieux' mais n'ont pas conscience de l'impact financier direct des rebuts et des heures de retouche non facturées.",
            "note": "1"
        },
        "7.4": {
            "comment": "Communication interne principalement orale et informelle. Déformation fréquente des consignes entre management et atelier.",
            "note": "1"
        },
        "7.5": {
            "comment": "Absence de GED structurée. Pilotage reposant sur 12 fichiers Excel critiques et manuels (opacité: 14 561 €/an). Cadres perdant 1 à 2h/jour à chercher l'information.",
            "note": "0"
        },

        # CLAUSE 8 : RÉALISATION DES ACTIVITÉS OPÉRATIONNELLES
        "8.1": {
            "comment": "Gammes opératoires insuffisamment standardisées. Variabilité de la qualité d'exécution et des temps de réglage selon les opérateurs en poste.",
            "note": "1"
        },
        "8.2": {
            "comment": "Pas de revue de devis chiffrée systématique. Contrats sous-traitants rang 2 aux exigences floues engendrant des ventes à marge quasi-nulle ou négative.",
            "note": "1"
        },
        "8.3": {
            "comment": "Les retours d'expérience sur la fabrication de pièces médicales en petites séries ne sont pas capitalisés dans des dossiers de conception réutilisables.",
            "note": "1"
        },
        "8.4": {
            "comment": "Absence de grille d'évaluation et de suivi annuel de la performance des prestataires et sous-traitants. Non-qualité externe subie sans recadrage.",
            "note": "0"
        },
        "8.5": {
            "comment": "5% de la production affectée par des retouches et reprises manuelles (7 000 €/an). Présence de double-saisies sur supports papier et Excel.",
            "note": "1"
        },
        "8.6": {
            "comment": "Contrôles dimensionnels effectués avant expédition, mais traçabilité et PV de recette non systématiquement adossés au dossier client.",
            "note": "2"
        },
        "8.7": {
            "comment": "Traitement au cas par cas des pièces défectueuses. Pas de zone de quarantaine physique matérialisée dans l'atelier pour stopper la perte financière.",
            "note": "1"
        },

        # CLAUSE 9 : ÉVALUATION DES PERFORMANCES
        "9.1": {
            "comment": "Pilotage dans l'opacité sur 12 fichiers Excel manuels (impact 14 561 €/an). Aucun tableau de bord automatisé en temps réel pour détecter les dérives.",
            "note": "0"
        },
        "9.1.2": {
            "comment": "Aucun suivi NPS ou indicateur factuel de satisfaction client. Traitement réactif des réclamations via l'octroi d'avoirs (45 000 €/an).",
            "note": "1"
        },
        "9.2": {
            "comment": "Aucun audit interne réalisé au cours des 12 derniers mois. Absence de contrôle périodique de la robustesse des processus.",
            "note": "0"
        },
        "9.3": {
            "comment": "Absence de Revue de Direction structurée. Les réunions existantes traitent uniquement des urgences opérationnelles quotidiennes sans arbitrage ROI.",
            "note": "0"
        },

        # CLAUSE 10 : AMÉLIORATION
        "10.1": {
            "comment": "Pas de démarche méthodique de chasse aux gaspillages (Lean / 5S). Idées d'amélioration du personnel terrain non formalisées ni tracées.",
            "note": "1"
        },
        "10.2": {
            "comment": "Traitement symptomatique des problèmes sans recherche des causes racines (méthode 5 Pourquoi inexistante). Récurrence chronique des mêmes pannes.",
            "note": "0"
        },
        "10.3": {
            "comment": "Le système manque d'agilité pour intégrer rapidement les retours d'expérience et adapter la rentabilité globale.",
            "note": "1"
        }
    }

    # Remplissage des Tables 3 à 9
    clause_tables = [doc.tables[3], doc.tables[4], doc.tables[5], doc.tables[6], doc.tables[7], doc.tables[8], doc.tables[9]]
    
    total_notes = []
    
    for t in clause_tables:
        for r in t.rows[1:]: # Ignorer l'en-tête
            code_iso = r.cells[0].text.strip()
            if code_iso in audit_data:
                data = audit_data[code_iso]
                set_cell_runs(r.cells[6], data["comment"], font_name="Arial", font_size_pt=8.0, bold=False)
                set_cell_runs(r.cells[7], data["note"], font_name="Arial", font_size_pt=9.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                total_notes.append(int(data["note"]))

    print(f"Nombre total de questions renseignées : {len(total_notes)}")
    score_global_pts = sum(total_notes)
    max_global_pts = len(total_notes) * 3
    score_global_pct = (score_global_pts / max_global_pts) * 100

    # 2. TABLE 10 : Matrice de priorisation (Quick Wins vs Projets de fond)
    t10 = doc.tables[10]
    
    # Quadrant 1 : Quick Wins
    set_cell_runs(t10.rows[0].cells[0], "Impact ROI Élevé / Difficulté Faible (Quick Wins - Priorité 1)", font_size_pt=9.5, bold=True, color_rgb=(15, 76, 58))
    qw_items = [
        "- Seuil de validation de marge obligatoire sur la revue de devis (Outil 08 - colmatage 45 k€)",
        "- Mode opératoire de secours et binômage pour la fonction du Dirigeant (Outil 20 - gain 15h/sem, 141 k€)",
        "- Centralisation et sécurisation des 12 fichiers Excel critiques (Outil 03 - 14,5 k€)",
        "- Fiche FNC 5 Pourquoi pour stopper les retouches atelier (Outil 19 - stop retouches 7 k€)"
    ]
    for item in qw_items:
        add_cell_paragraph(t10.rows[0].cells[0], item, font_size_pt=8.0)

    # Quadrant 2 : Projets de Fond
    set_cell_runs(t10.rows[0].cells[1], "Impact ROI Élevé / Difficulté Haute (Projets de Fond - Priorité 2)", font_size_pt=9.5, bold=True, color_rgb=(15, 76, 58))
    pf_items = [
        "- Certification ISO 9001 et structuration du SMQ (déblocage marché médical, 300 k€ + 150 k€)",
        "- Plan d'action QVT & RACI pour réduire le turnover et l'absentéisme de 8% (126 k€ + 32,2 k€)",
        "- Automatisation du processus de clôture mensuelle (réduction délai de 15j à J+2, 92 k€)",
        "- Cartographie macro et création des gammes opératoires visuelles (Outil 11/14)"
    ]
    for item in pf_items:
        add_cell_paragraph(t10.rows[0].cells[1], item, font_size_pt=8.0)

    # Quadrant 3 : Améliorations Secondaires
    set_cell_runs(t10.rows[1].cells[0], "Impact ROI Faible / Difficulté Faible (Améliorations Secondaires)", font_size_pt=9.5, bold=True)
    sec_items = [
        "- Standardisation des compte-rendus de réunion (Outil A2b)",
        "- Démarche 5S d'organisation des bureaux et ateliers (Outil A4)",
        "- Boîte à idées d'amélioration continue terrain"
    ]
    for item in sec_items:
        add_cell_paragraph(t10.rows[1].cells[0], item, font_size_pt=8.0)

    # Quadrant 4 : Fausses Bonnes Idées
    set_cell_runs(t10.rows[1].cells[1], "Impact ROI Faible / Difficulté Haute (Fausses Bonnes Idées - À Éviter)", font_size_pt=9.5, bold=True)
    fbi_items = [
        "- Rédaction de procédures administratives volumineuses non lues sur le terrain ('paperasse inutile')",
        "- Changement d'ERP lourd sans clarification préalable des processus et responsabilités",
        "- Multiplication des contrôles finaux sans traiter les causes racines en production"
    ]
    for item in fbi_items:
        add_cell_paragraph(t10.rows[1].cells[1], item, font_size_pt=8.0)

    # 3. TABLE 11 : Synthèse globale de robustesse & Calcul du score
    t11 = doc.tables[11]
    
    # Calcul des scores par pilier
    # Pilier 1 (8 questions : 4.1, 4.2, 5.2, 6.1, 7.1, 8.2, 8.4, 9.2) -> Notes: 0, 1, 0, 0, 1, 1, 0, 0 = 3/24 (12.5%)
    p1_pts = 3
    p1_max = 24
    p1_pct = (p1_pts / p1_max) * 100

    # Pilier 2 (7 questions : 4.3, 4.4, 5.3, 7.2, 7.5, 8.1, 8.5) -> Notes: 1, 1, 1, 1, 0, 1, 1 = 6/21 (28.6%)
    p2_pts = 6
    p2_max = 21
    p2_pct = (p2_pts / p2_max) * 100

    # Pilier 3 (5 questions : 5.1, 6.2, 9.1, 9.3, 10.2) -> Notes: 0, 1, 0, 0, 0 = 1/15 (6.7%)
    p3_pts = 1
    p3_max = 15
    p3_pct = (p3_pts / p3_max) * 100

    p_total_pts = p1_pts + p2_pts + p3_pts
    p_total_max = p1_max + p2_max + p3_max
    p_total_pct = (p_total_pts / p_total_max) * 100

    synthese_text = f"""📊 SYNTHÈSE GLOBALE DE ROBUSTESSE OPÉRATIONNELLE — TECHINDUSTRIE SAS

Formule de calcul du Score Global : Score Global SMQ (%) = (Somme des notes attribuées / (Nombre de questions validées × 3)) × 100

- Score Global SMQ (29 questions ISO 9001) : {score_global_pts} / {max_global_pts} pts ({score_global_pct:.1f} %)
- Score Synthétique 3 Piliers prioritaires : {p_total_pts} / {p_total_max} pts ({p_total_pct:.1f} %)
- ESTIMATION DE L'IMPÔT INVISIBLE ANNUEL : 907 866 € / an (soit 2 487 € / jour d'inaction)

Ventilation par Pilier Cabinet Delacour :
- Score Pilier 1 (Identification & Maîtrise des Risques) : {p1_pts} / {p1_max} pts ({p1_pct:.1f} %)
  Axe de gain principal : Stopper la fuite des coûts cachés de non-qualité, des litiges et des erreurs répétitives.
- Score Pilier 2 (Structuration & Souveraineté du Savoir) : {p2_pts} / {p2_max} pts ({p2_pct:.1f} %)
  Axe de gain principal : Gain de productivité collective, clarification RACI et suppression de la dépendance homme-clé.
- Score Pilier 3 (Pilotage par la Data & Agilité Décisionnelle) : {p3_pts} / {p3_max} pts ({p3_pct:.1f} %)
  Axe de gain principal : Automatisation du reporting, fin de l'opacité Excel et réaction décisionnelle sous 24 heures."""

    set_cell_runs(t11.rows[0].cells[0], synthese_text, font_size_pt=9.0, bold=False)

    # 4. TABLE 12 : Registre des vulnérabilités majeures (Points Rouges 48h)
    t12 = doc.tables[12]
    
    vuln_data = [
        ("Surcharge du Dirigeant en micro-management (15h/sem) & absence de suppléance", "141 000 €", "🔴 Haute (48h)", "Création fiche RACI (Outil 20), délégation des tâches d'atelier et binômage."),
        ("Absence de revue de devis chiffrée & contrats sous-traitants rang 2 toxiques", "45 000 €", "🔴 Haute (48h)", "Mise en place de la grille de revue de contrat (Outil 08) avec seuil de marge."),
        ("Opacité décisionnelle & dispersion des données sur 12 fichiers Excel manuels", "14 561 €", "🔴 Haute (48h)", "Centralisation et automatisation du Tableau de Bord Data (Outil 03)."),
        ("Croissance bridée & perte de marché médical par manque d'ISO 9001", "450 000 €", "🔴 Haute (48h)", "Lancement de la structuration SMQ et démarche ISO 9001 (Outil 00/01).")
    ]

    # Remplir les lignes du tableau 12
    for idx, (risque, cout, priorite, action) in enumerate(vuln_data, start=1):
        if idx < len(t12.rows):
            r = t12.rows[idx]
        else:
            r = t12.add_row()
        
        set_cell_runs(r.cells[0], risque, font_size_pt=8.5, bold=True)
        set_cell_runs(r.cells[1], cout, font_size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_runs(r.cells[2], priorite, font_size_pt=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_runs(r.cells[3], action, font_size_pt=8.5, bold=False)

    # Sauvegarder le fichier final
    doc.save(target_path)
    print(f"✅ Fichier d'audit complété généré avec succès dans le dossier dédié : {target_path}")

if __name__ == "__main__":
    main()
