import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

base_annex_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/OUTILS_ANNEXES_TRANSVERSAUX"
templates_dir = os.path.join(base_annex_dir, "01_TEMPLATES_VIERGES")
guides_dir = os.path.join(base_annex_dir, "02_GUIDES_METHODOLOGIQUES")
docx_dir = os.path.join(base_annex_dir, "DOCX")
html_dir = os.path.join(base_annex_dir, "HTML")
md_dir = os.path.join(base_annex_dir, "MD")

os.makedirs(templates_dir, exist_ok=True)
os.makedirs(guides_dir, exist_ok=True)
os.makedirs(docx_dir, exist_ok=True)
os.makedirs(html_dir, exist_ok=True)
os.makedirs(md_dir, exist_ok=True)

def save_template(doc, filename):
    doc.save(os.path.join(templates_dir, filename))
    doc.save(os.path.join(docx_dir, filename))

def save_guide(doc, filename):
    doc.save(os.path.join(guides_dir, filename))
    doc.save(os.path.join(docx_dir, filename))

# BRAND COLORS
BRAND_GREEN = RGBColor(27, 77, 62)     # #1B4D3E
DARK_GREEN = RGBColor(21, 40, 34)      # #152822
LIGHT_GREEN = RGBColor(74, 108, 92)    # #4A6C5C
GOLD = RGBColor(197, 160, 101)         # #C5A065
WHITE = RGBColor(255, 255, 255)
GREEN_TEXT = RGBColor(27, 94, 32)      # #1B5E20
RED_TEXT = RGBColor(183, 28, 28)       # #B71C1C
MUTED_TEXT = RGBColor(120, 120, 120)

CREAM_BG = "F9F9F7"
WHY_BG = "FDFBF7"
THEMES_BG = "F7F9F8"
BORDER_ACCENT = "E2DCCC"

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

HEADER_IMG_PATH = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
        else:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)

def create_base_document():
    doc = Document()
    section = doc.sections[0]
    
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.top_margin = Inches(0.95)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_head.paragraph_format.space_after = Pt(0)
    
    if os.path.exists(HEADER_IMG_PATH):
        r_img = p_head.add_run()
        r_img.add_picture(HEADER_IMG_PATH, width=Inches(7.3))
    
    # FOOTER OPENXML CLEAN
    footer = section.footer
    p_ft = footer.paragraphs[0]
    p_ft.text = ""
    p_ft.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ft.paragraph_format.space_after = Pt(0)
    
    pPr = p_ft._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), '10512')
    tabs.append(tab)
    pPr.append(tabs)
    
    r_l = p_ft.add_run("Cabinet Delacour - delacour.melchior@cabinetdelacour.com")
    r_l.font.name = FONT_SANS
    r_l.font.size = Pt(7.5)
    r_l.font.italic = True
    r_l.font.color.rgb = MUTED_TEXT
    
    r_tab = p_ft.add_run("\t\t")
    
    r_pg = p_ft.add_run("Page ")
    r_pg.font.name = FONT_SANS
    r_pg.font.size = Pt(7.5)
    r_pg.font.italic = True
    r_pg.font.color.rgb = MUTED_TEXT
    
    fld1 = OxmlElement('w:fldSimple')
    fld1.set(qn('w:instr'), 'PAGE')
    p_ft._p.append(fld1)
    
    r_sl = p_ft.add_run(" / ")
    r_sl.font.name = FONT_SANS
    r_sl.font.size = Pt(7.5)
    r_sl.font.italic = True
    r_sl.font.color.rgb = MUTED_TEXT
    
    fld2 = OxmlElement('w:fldSimple')
    fld2.set(qn('w:instr'), 'NUMPAGES')
    p_ft._p.append(fld2)
    
    return doc

def add_header_brand(doc, tool_code, tool_title, sub_meta):
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_title.autofit = False
    c_t = tbl_title.cell(0, 0)
    set_explicit_cell_width(c_t, 10200)
    set_cell_background(c_t, CREAM_BG)
    set_cell_margins(c_t, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_t, left={'sz': 24, 'val': 'single', 'color': 'C5A065'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p_title = c_t.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(2)
    r_cl = p_title.add_run(f"OUTILS ANNEXES TRANSVERSAUX - {tool_code} — {sub_meta}\n")
    r_cl.font.name = FONT_SANS
    r_cl.font.size = Pt(8.5)
    r_cl.font.bold = True
    r_cl.font.color.rgb = GOLD
    
    r_main_t = p_title.add_run(tool_title)
    r_main_t.font.name = FONT_SERIF
    r_main_t.font.size = Pt(14.5)
    r_main_t.font.bold = True
    r_main_t.font.color.rgb = BRAND_GREEN
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_consultant_box(doc, advice_text, themes_tuples):
    num_theme_rows = (len(themes_tuples) + 1) // 2
    total_rows = 2 + num_theme_rows

    tbl = doc.add_table(rows=total_rows, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, THEMES_BG)
    set_cell_margins(c_title, top=80, bottom=60, left=120, right=120)
    set_cell_border(c_title, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                             top={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                             right={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                             bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("CONSEILS CONSULTANT & PRINCIPES DE FORMULATION MÉTHODOLOGIQUE")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c_posture = tbl.cell(1, 0)
    c_posture_r = tbl.cell(1, 1)
    c_posture.merge(c_posture_r)
    set_explicit_cell_width(c_posture, 10200)
    set_cell_background(c_posture, THEMES_BG)
    set_cell_margins(c_posture, top=40, bottom=60, left=120, right=120)
    set_cell_border(c_posture, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                               right={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                               bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p_p = c_posture.paragraphs[0]
    p_p.paragraph_format.space_after = Pt(0)
    r_lbl = p_p.add_run("Posture & Animation : ")
    r_lbl.font.name = FONT_SANS
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(8)
    r_lbl.font.color.rgb = BRAND_GREEN
    r_txt = p_p.add_run(advice_text.replace("PME", "entreprise"))
    r_txt.font.name = FONT_SANS
    r_txt.font.italic = True
    r_txt.font.size = Pt(8)

    for idx, (t_head, t_body) in enumerate(themes_tuples):
        row_idx = 2 + (idx // 2)
        col_idx = idx % 2
        is_last_row = (row_idx == total_rows - 1)
        is_left = (col_idx == 0)

        if is_left and idx == len(themes_tuples) - 1:
            cell = tbl.cell(row_idx, 0)
            cell_r = tbl.cell(row_idx, 1)
            cell.merge(cell_r)
            set_explicit_cell_width(cell, 10200)
        else:
            cell = tbl.cell(row_idx, col_idx)
            set_explicit_cell_width(cell, 5100)

        set_cell_background(cell, "FFFFFF")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        left_b = {'sz': 24, 'val': 'single', 'color': '1B4D3E'} if is_left else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        right_b = {'sz': 4, 'val': 'single', 'color': '4A6C5C'} if not is_left else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        bottom_b = {'sz': 4, 'val': 'single', 'color': '4A6C5C'} if is_last_row else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        set_cell_border(cell, left=left_b, right=right_b, top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}, bottom=bottom_b)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        if t_head:
            r_th = p.add_run(t_head.replace("PME", "entreprise") + "\n")
            r_th.font.name = FONT_SERIF
            r_th.font.bold = True
            r_th.font.size = Pt(8)
            r_th.font.color.rgb = BRAND_GREEN
            r_tb = p.add_run(t_body.replace("PME", "entreprise"))
            r_tb.font.name = FONT_SANS
            r_tb.font.size = Pt(7.5)
            r_tb.font.color.rgb = BRAND_GREEN

# 3AB. FINANCIAL ROI BOX
def add_financial_roi_box(doc, gain_text, cout_text):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("ANCRAGE FINANCIER & IMPACT ROI SUR LA MARGE DE L'ENTREPRISE")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, CREAM_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("GAIN STRATÉGIQUE OU PROTECTION DE MARGE BRUTE (€) :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = GREEN_TEXT
    r1_b = p1.add_run("- " + gain_text.replace("PME", "entreprise"))
    r1_b.font.name = FONT_SANS
    r1_b.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, CREAM_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("COÛT DE NON-QUALITÉ OU PERTE EN CAS D'INACTION (€) :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RED_TEXT
    r2_b = p2.add_run("- " + cout_text.replace("PME", "entreprise"))
    r2_b.font.name = FONT_SANS
    r2_b.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

# 3B. INTERCONNECTIONS BOX
def add_interconnections_box(doc, inputs_list, outputs_list):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("MATRICE D'INTERCONNEXIONS & FLUX DE DONNÉES ENTRANTES / SORTANTES")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, CREAM_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("DONNÉES ENTRANTES (D'où viennent les informations ?) :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = BRAND_GREEN
    for inp in inputs_list:
        p = c1.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + inp.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, CREAM_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("DONNÉES SORTANTES (Vers quoi cet outil alimente-t-il ?) :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = BRAND_GREEN
    for out in outputs_list:
        p = c2.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + out.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_why_box(doc, gains_list, risques_list):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, WHY_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("POURQUOI CET OUTIL EST-IL ABSOLUMENT CAPITAL POUR VOTRE ENTREPRISE ?")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, WHY_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("Ce que vous gagnez à l'utiliser :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = GREEN_TEXT
    for g in gains_list:
        p = c1.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + g.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, WHY_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Ce que vous risquez sans cet outil :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RED_TEXT
    for r_item in risques_list:
        p = c2.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + r_item.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_guide_box(doc, steps_list):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_explicit_cell_width(cell, 10200)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_border(cell, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    r0 = p0.add_run(f"GUIDE DE REMPLISSAGE PRATIQUE ({len(steps_list)} ÉTAPES)\n")
    r0.font.name = FONT_SERIF
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = BRAND_GREEN
    
    for idx, step in enumerate(steps_list, 1):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r_num = p.add_run(f"Étape {idx} : ")
        r_num.font.name = FONT_SANS
        r_num.font.bold = True
        r_num.font.size = Pt(8)
        r_num.font.color.rgb = BRAND_GREEN
        r_txt = p.add_run(step.replace("PME", "entreprise"))
        r_txt.font.name = FONT_SANS
        r_txt.font.size = Pt(8)
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_next_steps_box(doc, next_list):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_explicit_cell_width(cell, 10200)
    set_cell_background(cell, CREAM_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_border(cell, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                          top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                          bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                          right={'sz': 6, 'val': 'single', 'color': 'C5A065'})
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    r0 = p0.add_run("QUE FAIRE UNE FOIS CET OUTIL COMPLÉTÉ ?\n")
    r0.font.name = FONT_SERIF
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = BRAND_GREEN
    
    for n in next_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run("- " + n.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)
        r.font.color.rgb = BRAND_GREEN
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def format_word_table(doc, headers, rows_data, col_widths_dxa=None, subheaders=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        clean_h = h_text.replace("PME", "entreprise")
        hdr_cells[i].text = clean_h
        if col_widths_dxa and i < len(col_widths_dxa):
            set_explicit_cell_width(hdr_cells[i], col_widths_dxa[i])
        set_cell_background(hdr_cells[i], "1B4D3E")
        set_cell_margins(hdr_cells[i], top=60, bottom=60, left=60, right=60)
        set_cell_border(hdr_cells[i], top={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      bottom={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      left={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      right={'sz': 4, 'val': 'single', 'color': '1B4D3E'})
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE

    if subheaders:
        sub_cells = tbl.add_row().cells
        for i, sub_text in enumerate(subheaders):
            clean_sub = sub_text.replace("PME", "entreprise")
            sub_cells[i].text = clean_sub
            if col_widths_dxa and i < len(col_widths_dxa):
                set_explicit_cell_width(sub_cells[i], col_widths_dxa[i])
            set_cell_background(sub_cells[i], "E8F0EC")
            set_cell_margins(sub_cells[i], top=50, bottom=50, left=60, right=60)
            set_cell_border(sub_cells[i], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                          bottom={'sz': 6, 'val': 'single', 'color': '1B4D3E'},
                                          left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in sub_cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)
                    r.font.bold = True
                    r.font.color.rgb = BRAND_GREEN
                
    for r_idx, row_data in enumerate(rows_data):
        row_cells = tbl.add_row().cells
        bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            clean_val = val.replace("🟢", "").replace("🔴", "").replace("🟠", "").replace("🟡", "").replace("PME", "entreprise").strip()
            row_cells[c_idx].text = clean_val
            if col_widths_dxa and c_idx < len(col_widths_dxa):
                set_explicit_cell_width(row_cells[c_idx], col_widths_dxa[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=60, right=60)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)
                    if "[Votre" in val or "____" in val or "[Saisissez" in val:
                        r.font.italic = True
                        r.font.bold = True
                        r.font.color.rgb = GOLD
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    set_explicit_cell_width(c, 10200)
    set_cell_background(c, THEMES_BG)
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    set_cell_border(c, left={'sz': 20, 'val': 'single', 'color': '1B4D3E'},
                       top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                       bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                       right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = FONT_SERIF
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = BRAND_GREEN
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def save_doc(doc, filename):
    doc.save(os.path.join(docx_dir, filename))

# ----------------------------------------------------
# 1. OUTIL A1: METHODE SMART PRACTICE
# ----------------------------------------------------
def make_outil_a1_template():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A1 (TEMPLATE)", "Grille de Formulation et de Validation des Objectifs SMART", "Trame Opérationnelle - Saisie Directe")
    
    add_section_header(doc, "GRILLE DE FORMULATION ET DE VALIDATION DE VOS OBJECTIFS (ORDRE STRICT S-M-A-R-T)")
    headers1 = ["PÉRIMÈTRE", "AMBITION", "S", "M", "M", "A", "R", "T", "SYNTHÈSE"]
    subheaders1 = ["Service / Équipe", "Idée de départ", "Action concrète à faire", "Situation actuelle (Départ)", "Cible visée (Résultat)", "Responsable & Moyens", "Niveau d'urgence (1-5)", "Fréquence & Échéance", "Phrase d'objectif rédigée"]
    rows1 = [
        ["Ventes & Commercial", "Réduire le délai des devis", "Créer un modèle de chiffrage rapide", "72 heures en moyenne", "Moins de 24h", "Marc (Com) / 2 000 € logiciel", "5 (Très urgent)", "Point hebdo (Cible : 31/03/2026)", "Transmettre 95% des devis en moins de 24h d'ici mars 2026."],
        ["Atelier & Fab", "Diminuer les pièces ratées", "Former les régleurs sur machine CNC", "4,2 % de pièces jetées", "Moins de 1,5 %", "Chef d'atelier / 15h formation", "4 (Urgent)", "Point 1x/mois (Cible : 15/05/2026)", "Réduire les pièces jetées à moins de 1,5% avant le 15 mai 2026."],
        ["Qualité & Client", "Traiter les réclamations", "Mettre en place la fiche de suivi des litiges", "Réponse en 5 jours", "Réponse sous 48h", "Responsable Qualité / Temps interne", "5 (Très urgent)", "Point hebdo (Cible : 31/12/2026)", "Répondre à 100% des réclamations clients sous 48h cette année."],
        ["Ressources Humaines", "Former l'équipe aux postes clés", "Créer le tableau des compétences de l'équipe", "30 % de personnes formées", "75 % de personnes formées", "Responsable RH / Budget 5 000 €", "3 (Moyen)", "Point trimestriel (Cible : 30/09/2026)", "Former 75% de l'équipe aux postes importants avant fin septembre 2026."],
        ["[Service]", "[Idée de départ]", "[Action S]", "[D'où on part]", "[Chiffre visé]", "[Qui & Avec quoi]", "[Urgence 1-5]", "[Point d'étape + Date limite]", "[Phrase finale d'objectif]"],
        ["[Service]", "[Idée de départ]", "[Action S]", "[D'où on part]", "[Chiffre visé]", "[Qui & Avec quoi]", "[Urgence 1-5]", "[Point d'étape + Date limite]", "[Phrase finale d'objectif]"]
    ]
    format_word_table(doc, headers1, rows1, col_widths_dxa=[1100, 1200, 1300, 1100, 1100, 1200, 900, 1000, 1300], subheaders=subheaders1)

    add_next_steps_box(doc, [
        "Inscrire les objectifs validés dans la Feuille de Route de l'entreprise (Outil 06).",
        "Ajouter les chiffres cibles dans le Tableau de Bord de la Direction (Outil 03).",
        "Partager les objectifs avec les responsables d'équipe lors des réunions mensuelles."
    ])
    save_template(doc, "OUTIL_A1_FORMULATION_OBJECTIFS_SMART_TEMPLATE.docx")

def make_outil_a1_guide():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A1 (GUIDE)", "Guide Méthodologique de Formulation des Objectifs SMART", "Notice Méthodologique - Management & ISO 9001:2015")
    
    add_consultant_box(doc, "Ne validez plus aucun objectif flou ! Une ambition vague comme 'Améliorer la qualité' n'aide personne sur le terrain. Ce guide vous accompagne pour décliner chaque ambition en indicateur chiffré, mesurable, doté de moyens et conforme à la clause 6.2 de l'ISO 9001.", [
        ("S = Spécifique (Action claire)", "Définir précisément ce qu'on va faire, sans phrase floue ni ambiguïté."),
        ("M = Mesurable (D'où on part -> Où on va)", "Chiffrer le point de départ (Situation actuelle) et le résultat voulu (Cible)."),
        ("A = Atteignable (Responsable & Moyens)", "Désigner une seule personne responsable et lui donner le temps/budget nécessaire."),
        ("R = Réaliste (Niveau d'urgence)", "S'assurer que l'objectif est utile pour l'entreprise et évaluer son urgence (1 à 5)."),
        ("T = Temporel (Date & Suivi)", "Fixer une date limite ferme et décider quand on fait le point (ex: toutes les semaines).")
    ])
    
    add_interconnections_box(doc, [
        "Politique Qualité & Axes Stratégiques (Outil 04)",
        "Analyse SWOT & PESTEL (Outils 01A/01B)",
        "Synthèse des réclamations clients & non-conformités (Outils 10 & 19)"
    ], [
        "Feuille de Route & Plans d'Actions (Outil 06)",
        "Tableau de Bord KPI Direction (Outil 03)",
        "Revue de Direction & Procès-Verbal (Outil 07)"
    ])
    
    add_why_box(doc, [
        "Des objectifs clairs que tout le monde comprend de la même manière.",
        "Des réunions d'évaluation simples et sans disputes sur les résultats.",
        "La garantie que chaque responsable a les moyens et le temps d'agir.",
        "Conformité garantie lors des audits de certification ISO 9001."
    ], [
        "Des objectifs flous perçus comme irréalistes ou injustes par l'équipe.",
        "Impossibilité de savoir si une action a vraiment fonctionné.",
        "Remarques des auditeurs ISO 9001 sur le manque de suivi des objectifs."
    ])
    
    add_guide_box(doc, [
        "Ouvrir le Template Vierge (OUTIL_A1_FORMULATION_OBJECTIFS_SMART_TEMPLATE.docx) pour saisir vos objectifs.",
        "Compléter le tableau de gauche à droite selon l'ordre strict S -> M -> A -> R -> T.",
        "Transférer les phrases d'objectifs finales validées dans la Feuille de Route de l'entreprise (Outil 06)."
    ])
    save_guide(doc, "OUTIL_A1_FORMULATION_OBJECTIFS_SMART_GUIDE.docx")

# ----------------------------------------------------
# 2. OUTIL A2: COMPTE RENDU REUNION (CLIENT & INTERNE)
# ----------------------------------------------------
def make_outil_a2_template():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A2 (TEMPLATE)", "Fiche & Compte-Rendu de Réunion (Trame Opérationnelle)", "Trame Opérationnelle - Saisie Directe")
    
    add_section_header(doc, "1. CADRAGE DE LA RÉUNION (À REMPLIR EN DÉBUT DE SÉANCE)")
    headers1 = ["Sujet / Intitulé de la Réunion", "Date & Heure", "Lieu / Visioconférence", "Type de Réunion", "Présents (Nom & Service)", "But Principal Attendu"]
    rows1 = [
        ["Revue de Lancement Affaire ALFA", "16/02/2026 à 10h", "Siège Client / Teams", "☑ Client  ☐ Interne  ☐ Chantier", "M. Dupont (Acheteur), M. Delacour (Chef projet)", "Valider les plans 3D et le planning de livraison"],
        ["[Nom de la réunion]", "[Date & Heure]", "[Lieu / Teams]", "☐ Client  ☐ Interne  ☐ Chantier", "[Liste des présents]", "[Objectif principal attendu]"]
    ]
    format_word_table(doc, headers1, rows1, col_widths_dxa=[2000, 1400, 1500, 1500, 2000, 1800])

    add_section_header(doc, "2. RELEVÉ DES DÉCISIONS ET PLAN D'ACTION (CONNECTÉ À L'OUTIL A1)")
    headers2 = ["N°", "Sujet / Décision Retenue", "Action Concrète à Réaliser", "Qui s'en occupe ?", "Date Limite", "Alignée SMART (Outil A1)", "Statut"]
    rows2 = [
        ["1", "Modification nuance inox validée en Inox 316L", "Mettre à jour le plan 3D Indice B et la nomenclature", "Jean (Bureau d'études)", "18/02/2026", "🟢 Validé avec Outil A1", "🟡 En cours"],
        ["2", "Livraison partielle du lot 1 acceptée au 15/03", "Confirmer le planning avec le transporteur express", "Sophie (Logistique)", "20/02/2026", "🟢 Validé avec Outil A1", "🟢 Fait"],
        ["3", "Acompte de 30% débloqué par le client", "Transmettre le RIB et vérifier le virement", "Comptabilité", "17/02/2026", "☐ À formaliser en A1", "🟢 Fait"],
        ["[N°]", "[Décision / Arbitrage]", "[Action concrète de suivi]", "[Responsable]", "[Échéance]", "☐ Utiliser l'Outil A1 si complexe", "[À faire/En cours/Fait]"],
        ["[N°]", "[Décision / Arbitrage]", "[Action concrète de suivi]", "[Responsable]", "[Échéance]", "☐ Utiliser l'Outil A1 si complexe", "[À faire/En cours/Fait]"]
    ]
    format_word_table(doc, headers2, rows2, col_widths_dxa=[500, 2300, 2500, 1400, 1100, 1200, 1200])

    doc.add_page_break()

    add_section_header(doc, "3. MODÈLE D'EMAIL DE DIFFUSION SOUS 24H (À COPIER-COLLER DANS VOTRE MESSAGERIE)")
    
    tbl_email = doc.add_table(rows=1, cols=1)
    tbl_email.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_email.autofit = False
    c_e = tbl_email.cell(0, 0)
    set_explicit_cell_width(c_e, 10200)
    set_cell_background(c_e, CREAM_BG)
    set_cell_margins(c_e, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_e, left={'sz': 24, 'val': 'single', 'color': 'C5A065'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p_e = c_e.paragraphs[0]
    p_e.paragraph_format.space_after = Pt(2)
    r_obj = p_e.add_run("Objet du message : ")
    r_obj.font.name = FONT_SANS
    r_obj.font.bold = True
    r_obj.font.size = Pt(8.5)
    r_obj.font.color.rgb = BRAND_GREEN
    
    r_obj_t = p_e.add_run("[Compte-Rendu] Réunion du [Date] — [Sujet de la Réunion]\n\n")
    r_obj_t.font.name = FONT_SANS
    r_obj_t.font.size = Pt(8.5)
    
    r_body = p_e.add_run(
        "Bonjour à tous,\n\n"
        "Veuillez trouver ci-joint le compte-rendu synthétique de notre réunion du [Date] concernant [Sujet].\n\n"
        "Pour rappel, les principales actions engagées sont :\n"
        "- [Action 1] — Responsable : [Nom] (Échéance : [Date])\n"
        "- [Action 2] — Responsable : [Nom] (Échéance : [Date])\n\n"
        "Merci de nous signaler sous 48h toute remarque ou correction éventuelle.\n\n"
        "Bien cordialement,\n"
        "[Votre Nom / Cabinet Delacour]"
    )
    r_body.font.name = FONT_SANS
    r_body.font.size = Pt(8)
    r_body.font.italic = True
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

    add_next_steps_box(doc, [
        "Pour toute action complexe ou stratégique issue de la réunion, déclinez-la avec l'Outil A1 (Objectifs SMART).",
        "Copier-coller le modèle d'email ci-dessus et joindre ce compte-rendu complété.",
        "Envoyer sous 24h par email aux participants et archiver dans le dossier GED de l'entreprise (02_VENTES_MARGES_ACHATS/COMPTES_RENDUS/)."
    ])
    save_template(doc, "OUTIL_A2_COMPTE_RENDU_DE_REUNION_TEMPLATE.docx")

def make_outil_a2_guide():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A2 (GUIDE)", "Guide d'Animation & de Rédaction des Comptes-Rendus", "Notice Méthodologique - Communication & Suivi d'Affaires")
    
    add_consultant_box(doc, "Ne quittez plus aucune réunion sans consigner les décisions et attribuer les actions ! Un compte-rendu clair et synthétique rédigé séance tenante sécurise vos affaires, évite les malentendus et garantit l'exécution des engagements.", [
        ("1. Cadrage & Ordre du jour", "Fixer la date, les présents et le but clair avant de commencer la réunion."),
        ("2. Décisions Validées", "Consigner factuellement les arbitrages pris pendant la séance (sans bla-bla inutile)."),
        ("3. Actions & Engagements", "Attribuer chaque action à une seule personne responsable avec une date limite ferme."),
        ("4. Diffusion sous 24h", "Envoyer par email le compte-rendu sous 24h aux participants pour validation.")
    ])
    
    add_why_box(doc, [
        "Rassurance et crédibilité immédiates auprès de vos clients et partenaires.",
        "Suivi rigoureux de toutes les décisions et actions engagées sans rien oublier.",
        "Protection contractuelle et preuve écrite en cas de désaccord ou de changement de demande."
    ], [
        "Oubli des décisions prises oralement entraînant des retards ou des erreurs.",
        "Contestation ultérieure des consignes, des tarifs ou des délais par le client.",
        "Discussions répétitives et perte de temps lors des réunions suivantes."
    ])
    
    add_guide_box(doc, [
        "Ouvrir le Template Vierge (OUTIL_A2_COMPTE_RENDU_DE_REUNION_TEMPLATE.docx) en début de réunion.",
        "Saisir directement les décisions et attribuer les responsables séance tenante.",
        "Copier-coller le modèle d'email de la page 2 et diffuser sous 24h."
    ])
    save_guide(doc, "OUTIL_A2_COMPTE_RENDU_DE_REUNION_GUIDE.docx")
# ----------------------------------------------------
# 3. OUTIL A3: GRILLE PARETO 80 20
# ----------------------------------------------------
def make_outil_a3_template():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A3 (TEMPLATE)", "Grille d'Analyse Pareto (80/20) & Priorisation des Actions", "Trame Opérationnelle • Saisie Directe")
    
    add_section_header(doc, "GRILLE D'ANALYSE PARETO ET CUMUL DES CAUSES DE DYSFONCTIONNEMENT")
    headers = ["Rang", "Catégorie de Défaut / Problème", "Fréquence (Nb)", "Coût Financier (€)", "% Du Total", "% Cumulé (Zone 80%)", "Priorité Action"]
    rows = [
        ["1", "Défaut d'ébavurage / Bavures usinage", "42", "8 400 €", "46.7 %", "46.7 % (Zone Vitale)", "PRIORITÉ 1 (Action Immédiate)"],
        ["2", "Erreur d'étiquetage / Colisage", "18", "5 200 €", "28.9 %", "75.6 % (Zone Vitale)", "PRIORITÉ 2 (Action Immédiate)"],
        ["3", "Rayures de surface lors du transport", "8", "2 100 €", "11.7 %", "87.3 % (Zone Secondaire)", "Sous surveillance (Prio 3)"],
        ["4", "Retard de livraison fournisseur", "4", "1 100 €", "6.1 %", "93.4 % (Zone Secondaire)", "Action différée"],
        ["5", "Erreur d'adresse d'expédition", "2", "600 €", "3.3 %", "96.7 % (Zone Secondaire)", "Action différée"],
        ["6", "Autres défauts mineurs", "2", "600 €", "3.3 %", "100.0 %", "Action différée"],
        ["[Rang]", "[Catégorie de problème]", "[Nb]", "[Coût €]", "[% Total]", "[% Cumulé]", "[Niveau de priorité]"],
        ["[Rang]", "[Catégorie de problème]", "[Nb]", "[Coût €]", "[% Total]", "[% Cumulé]", "[Niveau de priorité]"],
        ["[Rang]", "[Catégorie de problème]", "[Nb]", "[Coût €]", "[% Total]", "[% Cumulé]", "[Niveau de priorité]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[600, 2400, 1200, 1400, 1100, 1700, 1800])
    
    add_next_steps_box(doc, [
        "Présenter la courbe de Pareto lors du Comité Qualité mensuel.",
        "Déclencher une analyse 5 Pourquoi (Outil 19) sur la cause N°1 identifiée.",
        "Archiver la grille sous 03_OPERATIONS_PRODUCTION/PARETO_ANALYSIS/."
    ])
    save_template(doc, "OUTIL_A3_GRILLE_PARETO_80_20_TEMPLATE.docx")

def make_outil_a3_guide():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A3 (GUIDE)", "Guide Méthodologique de l'Analyse Pareto (80/20)", "Notice Méthodologique • Analyse & Priorisation des Causes")
    
    add_consultant_box(doc, "Ne vous éparpillez plus ! Appliquez le principe de Pareto (80/20) : 20% des causes produisent 80% des conséquences (coûts, retards, non-conformités). Concentrez vos ressources sur les 2 ou 3 facteurs vitaux.", [
        ("1. Collecte des Données Factuelles", "Lister et comptabiliser la fréquence ou le coût des défauts/problèmes sur une période donnée."),
        ("2. Classement Décroissant", "Trier les types de dysfonctionnements du plus coûteux/fréquent au moins impactant."),
        ("3. Calcul des Pourcentages Cumulés", "Déterminer la part cumulée de chaque cause dans le total global."),
        ("4. Plan d'Action Focalisé (Zone 80%)", "Traiter en priorité absolue les catégories constituant les premiers 80% du total.")
    ])
    
    add_why_box(doc, [
        "Rendement maximal des actions d'amélioration engagées.",
        "Réduction rapide de 80% des coûts de non-qualité avec un effort ciblé.",
        "Utilisation rationnelle et factuelle du budget d'investissement."
    ], [
        "Dispersion de l'énergie de l'équipe sur des détails insignifiants (les 80% de causes secondaires).",
        "Impression d'épuisement sans baisse visible des réclamations clients.",
        "Manque de crédibilité des analyses présentées à la direction."
    ])
    
    add_financial_roi_box(doc,
        "Réduction rapide des coûts de non-qualité par le ciblage prioritaire des causes majeures.",
        "Gaspillage d'énergie et de budget sur des causes secondaires sans impact sur la marge."
    )
    
    add_interconnections_box(doc, [
        "OUTIL 19 (FNC & 5 Pourquoi) : Historique des défauts et coûts de retouches.",
        "OUTIL 27 (Audit Interne) : Registre des écarts décelés."
    ], [
        "OUTIL A5 (Cycle PDCA) : Lancement prioritaire du chantier d'amélioration.",
        "OUTIL 05 (AMDEC Risques) : Recalcul du score de criticité post-action."
    ])
    
    add_guide_box(doc, [
        "Ouvrir le Template Vierge (OUTIL_A3_GRILLE_PARETO_80_20_TEMPLATE.docx).",
        "Renseigner les catégories de défauts observés sur les 3 derniers mois.",
        "Saisir la fréquence ou le coût financier (€) associé à chaque catégorie et trier par ordre décroissant.",
        "Engager le plan d'action immédiat sur les 2 premières lignes du tableau (Zone 80%)."
    ])
    save_guide(doc, "OUTIL_A3_GRILLE_PARETO_80_20_GUIDE.docx")

# ----------------------------------------------------
# 4. OUTIL A4: METHODE 5S ATELIER BUREAU
# ----------------------------------------------------
def make_outil_a4_template():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A4 (TEMPLATE)", "Grille d'Audit & Évaluation 5S (Atelier & Bureau/GED)", "Trame Opérationnelle • Saisie Directe")
    
    add_section_header(doc, "GRILLE D'AUDIT ET D'ÉVALUATION 5S DE ZONE (PHYSIQUE OU NUMÉRIQUE)")
    headers = ["Pilier 5S", "Critère d'Évaluation Factuel", "Conforme (Oui/Non)", "Note /5", "Constat & Écart Décelé", "Action Corrective Immédiate", "Responsable"]
    rows = [
        ["1. Seiri (Débarrasser)", "Absence de pièces rebutées, d'outils HS ou de fichiers périmés au poste", "Oui", "4.5 / 5", "Quelques vieux flacons de produit à jeter", "Évacuer le bac déchets chimie", "Chef d'équipe"],
        ["2. Seiton (Ranger)", "Emplacement identifié par marquage (outils, bacs, sous-dossiers GED)", "Oui", "4.0 / 5", "Marquage sol de la palette d'expédition effacé", "Refaire le traçage peinture jaune", "Maintenance"],
        ["3. Seiso (Nettoyer)", "Poste de travail propre, machines dépoussiérées, pas de fuite", "Oui", "5.0 / 5", "Zone très propre, bac arrosage vidangé", "Rien à signaler", "Opérateur"],
        ["4. Seiketsu (Standard)", "Consignes d'utilisation visuelles affichées et lisibles", "Oui", "4.0 / 5", "Fiche réglage manquante sur tour CNC-03", "Imprimer et afficher la fiche visuelle", "Resp. QSE"],
        ["5. Shitsuke (Pérenniser)", "Audit 5S réalisé chaque mois et actions du mois précédent closes", "Oui", "4.5 / 5", "Actions du mois passé closes à 90%", "Point lors de la Réunion Flash", "Pilote 5S"],
        ["[Pilote 5S]", "[Votre critère]", "[Oui/Non]", "[/5]", "[Constat]", "[Action corrective]", "[Nom]"],
        ["[Pilote 5S]", "[Votre critère]", "[Oui/Non]", "[/5]", "[Constat]", "[Action corrective]", "[Nom]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1500, 2400, 900, 800, 1800, 1800, 1000])
    
    add_next_steps_box(doc, [
        "Afficher la note d'audit 5S au tableau de la zone.",
        "Archiver le compte-rendu sous 03_OPERATIONS_PRODUCTION/AUDITS_5S/."
    ])
    save_template(doc, "OUTIL_A4_METHODE_5S_ATELIER_BUREAU_TEMPLATE.docx")

def make_outil_a4_guide():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A4 (GUIDE)", "Guide de Déploiement & Méthode 5S (Atelier & Bureau/GED)", "Notice Méthodologique • Organisation & Propreté")
    
    add_consultant_box(doc, "Transformez vos espaces de travail physiques et informatiques en modèles d'efficacité ! La méthode 5S élimine le désordre, réduit les temps de recherche et garantit la sécurité au quotidien.", [
        ("1. Seiri (Débarrasser / Éliminer)", "Jeter l'inutile, jeter les fichiers obsolètes, garder uniquement ce qui sert au poste."),
        ("2. Seiton (Ranger / Structurer)", "Une place pour chaque chose, chaque chose à sa place (marquage au sol, arborescence GED)."),
        ("3. Seiso (Nettoyer / Inspecter)", "Nettoyer le poste et les outils, détecter les fuites d'huile ou les bugs système."),
        ("4. Seiketsu (Standardiser) & Shitsuke (Pérenniser)", "Règles visuelles de rangement et audit mensuel pour maintenir le niveau.")
    ])
    
    add_why_box(doc, [
        "Environnement de travail motivant, propre et sécurisé.",
        "Gain de temps massif (tout outil ou fichier disponible en < 30 sec).",
        "Image d'excellence irréprochable transmise aux clients en visite."
    ], [
        "Risque d'accident de travail par encombrement des allées ou sols glissants.",
        "Utilisation par erreur d'outils défectueux ou de pièces obsolètes.",
        "Perte de crédibilité et impression de désordre lors des audits."
    ])
    
    add_guide_box(doc, [
        "Ouvrir le Template Vierge (OUTIL_A4_METHODE_5S_ATELIER_BUREAU_TEMPLATE.docx).",
        "Planifier la journée 'Chantier 5S' dans la zone cible (Atelier ou Bureau/GED).",
        "Éliminer tout l'inutile et définir les emplacements standards.",
        "Réaliser l'audit initial puis reconduire mensuellement (Note cible >= 85/100)."
    ])
    save_guide(doc, "OUTIL_A4_METHODE_5S_ATELIER_BUREAU_GUIDE.docx")

# ----------------------------------------------------
# 5. OUTIL A5: TRAME RESOLUTION PROBLEMES PDCA
# ----------------------------------------------------
def make_outil_a5_template():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A5 (TEMPLATE)", "Trame de Résolution de Problèmes & Cycle PDCA (Deming)", "Trame Opérationnelle • Saisie Directe")
    
    add_section_header(doc, "TRAME DE SUIVI DU CHANTIER D'AMÉLIORATION PDCA")
    headers = ["Phase PDCA", "Étape Methodologique", "Contenu & Actions Déployées", "Indicateur de Succès / Mesure", "Responsable", "Statut Phase"]
    rows = [
        ["1. PLAN", "Analyse Cause Racine (5 Pourquoi)", "Problème : Retards d'usinage. Cause : Mauvais réglage outil CNC", "Délais de réglage < 15 min", "Resp. BE / Prod", "Validé"],
        ["2. DO", "Mise en Œuvre du Plan d'Action", "Achat gabarit de réglage hors machine + formation opérateur", "Gabarit livré et installé", "Chef Prod", "Validé"],
        ["3. CHECK", "Vérification Efficacité (H+30j)", "Mesure du temps réel de réglage sur 20 séries lancées", "Moyenne : 12 min (Cible atteinte)", "Resp. QSE", "Validé"],
        ["4. ACT", "Standardisation & Ancrage", "Mise à jour de l'Instruction Visuelle INS-01 (Outil 14)", "INS-01 Indice B plastifiée", "Resp. QSE", "Clôturé"],
        ["[Phase]", "[Étape]", "[Actions déployées]", "[Indicateur succès]", "[Nom]", "[Statut]"],
        ["[Phase]", "[Étape]", "[Actions déployées]", "[Indicateur succès]", "[Nom]", "[Statut]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1200, 2200, 3200, 1800, 1000, 800])
    
    add_next_steps_box(doc, [
        "Présenter le bilan PDCA en Comité de Direction.",
        "Archiver sous 01_DIRECTION_STRATEGIE/CHANTIERS_PDCA/."
    ])
    save_template(doc, "OUTIL_A5_TRAME_RESOLUTIONS_PROBLEMES_PDCA_TEMPLATE.docx")

def make_outil_a5_guide():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A5 (GUIDE)", "Guide Méthodologique de Résolution de Problèmes (Cycle PDCA)", "Notice Méthodologique • Amélioration Continue")
    
    add_consultant_box(doc, "Ne laissez aucun problème récurrent gâcher le quotidien de vos équipes ! Utilisez la boucle PDCA (Plan, Do, Check, Act) pour traiter méthodiquement tout problème, de sa détection jusqu'à la preuve d'efficacité.", [
        ("1. PLAN (Planifier)", "Identifier le problème factuel, analyser les causes racines (5 Pourquoi) et bâtir le plan d'action."),
        ("2. DO (Déployer)", "Mettre en œuvre les actions correctives sur le terrain avec des responsables et des échéances."),
        ("3. CHECK (Vérifier)", "Mesurer l'efficacité réelle 30 jours plus tard au moyen d'indicateurs chiffrés."),
        ("4. ACT (Standardiser)", "Si l'action est efficace, modifier la procédure officielle pour ancrer le gain définitivement.")
    ])
    
    add_why_box(doc, [
        "Résolution définitive et pérenne des dysfonctionnements.",
        "Garantie de non-récidive des erreurs.",
        "Preuve ultime de démarche d'amélioration continue (ISO 9001 Clause 10.2)."
    ], [
        "Actions correctives abandonnées en cours de route sans vérification.",
        "Retour systématique des mêmes pannes et litiges.",
        "Non-conformité majeure d'audit ISO 9001 (Clause 10.2)."
    ])
    
    add_interconnections_box(doc, [
        "OUTIL 19 (FNC & 5 Pourquoi) : Cause racine identifiée.",
        "OUTIL A3 (Pareto 80/20) : Priorisation du chantier d'amélioration."
    ], [
        "OUTIL 07 (Revue Direction) : Bilan du plan d'action d'amélioration.",
        "OUTIL 27 (Audit Interne) : Preuve d'éradication du dysfonctionnement."
    ])
    
    add_guide_box(doc, [
        "Ouvrir le Template Vierge (OUTIL_A5_TRAME_RESOLUTIONS_PROBLEMES_PDCA_TEMPLATE.docx).",
        "Renseigner les 4 phases du cycle PDCA pour chaque chantier.",
        "Définir le critère de réussite lors de la phase PLAN.",
        "Vérifier formellement le résultat lors de la phase CHECK et mettre à jour la procédure lors de la phase ACT."
    ])
    save_guide(doc, "OUTIL_A5_TRAME_RESOLUTIONS_PROBLEMES_PDCA_GUIDE.docx")

# RUN ALL BUILDERS FOR ANNEX TOOLS
make_outil_a1_template()
make_outil_a1_guide()
make_outil_a2_template()
make_outil_a2_guide()
make_outil_a3_template()
make_outil_a3_guide()
make_outil_a4_template()
make_outil_a4_guide()
make_outil_a5_template()
make_outil_a5_guide()

# GENERATE HTML AND MD FOR ANNEX TOOLS
annex_tools_data = [
    ("OUTIL_A1_METHODE_SMART_PRACTICE", "OUTIL A1 - Méthode SMART", "OUTIL A1 — Méthode Transversale - Management", "Guide d'Application & Grille de Formulation des Objectifs SMART", "Ne validez plus aucun objectif vague. Utilisez la grille SMART pour chiffrer chaque ambition."),
    ("OUTIL_A2_COMPTE_RENDU_REUNION_CLIENT", "OUTIL A2 - Compte-Rendu Réunion Client", "OUTIL A2 — Méthode Transversale - Communication", "Trame Universelle de Réunion & Compte-Rendu RDV Client", "Fixez les engagements mutuels et verrouillez le périmètre pour éliminer les litiges."),
    ("OUTIL_A3_GRILLE_PARETO_80_20", "OUTIL A3 - Grille Pareto 80/20", "OUTIL A3 — Méthode Transversale - Analyse", "Grille d'Analyse Pareto (80/20) & Priorisation des Actions", "Focalisez l'énergie de vos équipes sur les 20% de causes générant 80% des conséquences."),
    ("OUTIL_A4_METHODE_5S_ATELIER_BUREAU", "OUTIL A4 - Méthode 5S Atelier & Bureau", "OUTIL A4 — Méthode Transversale - Organisation", "Guide de Déploiement & Grille d'Audit 5S (Atelier & Bureau/GED)", "Éliminez le désordre, réduisez les temps de recherche et garantissez la propreté."),
    ("OUTIL_A5_TRAME_RESOLUTIONS_PROBLEMES_PDCA", "OUTIL A5 - Résolution Problèmes PDCA", "OUTIL A5 — Méthode Transversale - Amélioration", "Fiche Générique de Résolution de Problèmes (Cycle PDCA)", "Traitez méthodiquement tout dysfonctionnement, de la détection jusqu'à l'ancrage.")
]

for filename_base, title, meta, tool_title, advice in annex_tools_data:
    # HTML File
    html_content = f"""<!DOCTYPE html>
<html lang="fr">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{title} | Cabinet Delacour</title>
  <style>
    body {{ font-family: Arial, sans-serif; background-color: #F9F9F7; color: #152822; padding: 24px; }}
    .container {{ max-width: 1000px; margin: 0 auto; background: #FFF; padding: 32px; border-radius: 4px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); }}
    .header-banner {{ background: #152822; border-bottom: 4px solid #C5A065; color: #F0EAD6; padding: 16px 24px; font-weight: bold; }}
    .tool-title-box {{ background: #F9F9F7; border-left: 5px solid #C5A065; padding: 16px; margin: 20px 0; }}
    .tool-meta {{ color: #C5A065; font-size: 12px; font-weight: bold; text-transform: uppercase; }}
    .tool-title {{ font-family: Georgia, serif; font-size: 22px; color: #152822; margin-top: 4px; }}
    .consultant-box {{ background: #F7F9F8; border-left: 5px solid #4A6C5C; padding: 16px; margin-bottom: 20px; }}
  </style>
</head>
<body>
<div class="container">
  <div class="header-banner">CABINET DELACOUR - OUTILS ANNEXES TRANSVERSAUX</div>
  <div class="tool-title-box">
    <div class="tool-meta">{meta}</div>
    <div class="tool-title">{tool_title}</div>
  </div>
  <div class="consultant-box">
    <strong>Conseil Consultant :</strong> {advice}
  </div>
  <p>Fichier modèle complet disponible en format Word (.docx) dans votre dossier client.</p>
</div>
</body>
</html>"""
    with open(os.path.join(html_dir, f"{filename_base}.html"), "w", encoding="utf-8") as f:
        f.write(html_content)
        
    # MD File
    md_content = f"# 🧰 {tool_title}\n\n> **{meta}**\n\n## 💡 Conseil Consultant\n{advice}\n\n*Consulter la grille complète dans le fichier Word `.docx` rattaché.*"
    with open(os.path.join(md_dir, f"{filename_base}.md"), "w", encoding="utf-8") as f:
        f.write(md_content)

print("Successfully generated all 5 Annex Tools (DOCX, HTML, MD) under OUTILS_ANNEXES_TRANSVERSAUX!")
