import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

base_annex_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/OUTILS_ANNEXES_TRANSVERSAUX"
docx_dir = os.path.join(base_annex_dir, "DOCX")
html_dir = os.path.join(base_annex_dir, "HTML")
md_dir = os.path.join(base_annex_dir, "MD")

os.makedirs(docx_dir, exist_ok=True)
os.makedirs(html_dir, exist_ok=True)
os.makedirs(md_dir, exist_ok=True)

# BRAND COLORS
BRAND_GREEN = RGBColor(27, 77, 62)     # #1B4D3E
DARK_GREEN = RGBColor(21, 40, 34)      # #152822
LIGHT_GREEN = RGBColor(74, 108, 92)    # #4A6C5C
GOLD = RGBColor(197, 160, 101)         # #C5A065
WHITE = RGBColor(255, 255, 255)
GREEN_TEXT = RGBColor(27, 94, 32)      # #1B5E20
RED_TEXT = RGBColor(183, 28, 28)       # #B71C1C
MUTED_TEXT = RGBColor(120, 120, 120)

CREAM_BG = "F9F9F7"
WHY_BG = "FDFBF7"
THEMES_BG = "F7F9F8"
BORDER_ACCENT = "E2DCCC"

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

HEADER_IMG_PATH = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
        else:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)

def create_base_document():
    doc = Document()
    section = doc.sections[0]
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.top_margin = Inches(0.95)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_head.paragraph_format.space_after = Pt(0)
    
    if os.path.exists(HEADER_IMG_PATH):
        r_img = p_head.add_run()
        r_img.add_picture(HEADER_IMG_PATH, width=Inches(7.3))
    
    footer = section.footer
    p_ft = footer.paragraphs[0]
    p_ft.text = ""
    p_ft.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ft.paragraph_format.space_after = Pt(0)
    
    pPr = p_ft._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), '10512')
    tabs.append(tab)
    pPr.append(tabs)
    
    r_l = p_ft.add_run("Cabinet Delacour - delacour.melchior@cabinetdelacour.com")
    r_l.font.name = FONT_SANS
    r_l.font.size = Pt(7.5)
    r_l.font.italic = True
    r_l.font.color.rgb = MUTED_TEXT
    
    r_tab = p_ft.add_run("\t\t")
    
    r_pg = p_ft.add_run("Page ")
    r_pg.font.name = FONT_SANS
    r_pg.font.size = Pt(7.5)
    r_pg.font.italic = True
    r_pg.font.color.rgb = MUTED_TEXT
    
    fld1 = OxmlElement('w:fldSimple')
    fld1.set(qn('w:instr'), 'PAGE')
    p_ft._p.append(fld1)
    
    r_sl = p_ft.add_run(" / ")
    r_sl.font.name = FONT_SANS
    r_sl.font.size = Pt(7.5)
    r_sl.font.italic = True
    r_sl.font.color.rgb = MUTED_TEXT
    
    fld2 = OxmlElement('w:fldSimple')
    fld2.set(qn('w:instr'), 'NUMPAGES')
    p_ft._p.append(fld2)
    
    return doc

def add_header_brand(doc, tool_code, tool_title, sub_meta):
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_title.autofit = False
    c_t = tbl_title.cell(0, 0)
    set_explicit_cell_width(c_t, 10200)
    set_cell_background(c_t, CREAM_BG)
    set_cell_margins(c_t, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_t, left={'sz': 24, 'val': 'single', 'color': 'C5A065'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p_title = c_t.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(2)
    r_cl = p_title.add_run(f"OUTILS ANNEXES TRANSVERSAUX - {tool_code} — {sub_meta}\n")
    r_cl.font.name = FONT_SANS
    r_cl.font.size = Pt(8.5)
    r_cl.font.bold = True
    r_cl.font.color.rgb = GOLD
    
    r_main_t = p_title.add_run(tool_title)
    r_main_t.font.name = FONT_SERIF
    r_main_t.font.size = Pt(14.5)
    r_main_t.font.bold = True
    r_main_t.font.color.rgb = BRAND_GREEN
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_consultant_box(doc, advice_text, themes_tuples):
    while len(themes_tuples) < 4:
        themes_tuples.append(("", ""))

    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, THEMES_BG)
    set_cell_margins(c_title, top=80, bottom=60, left=120, right=120)
    set_cell_border(c_title, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                             top={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                             right={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                             bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("CONSEILS CONSULTANT & PRINCIPES DE FORMULATION MÉTHODOLOGIQUE")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c_posture = tbl.cell(1, 0)
    c_posture_r = tbl.cell(1, 1)
    c_posture.merge(c_posture_r)
    set_explicit_cell_width(c_posture, 10200)
    set_cell_background(c_posture, THEMES_BG)
    set_cell_margins(c_posture, top=40, bottom=60, left=120, right=120)
    set_cell_border(c_posture, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                               right={'sz': 4, 'val': 'single', 'color': '4A6C5C'},
                               bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p_p = c_posture.paragraphs[0]
    p_p.paragraph_format.space_after = Pt(0)
    r_lbl = p_p.add_run("Posture & Animation : ")
    r_lbl.font.name = FONT_SANS
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(8)
    r_lbl.font.color.rgb = BRAND_GREEN
    r_txt = p_p.add_run(advice_text.replace("PME", "entreprise"))
    r_txt.font.name = FONT_SANS
    r_txt.font.italic = True
    r_txt.font.size = Pt(8)

    cards_config = [
        (tbl.cell(2, 0), themes_tuples[0][0], themes_tuples[0][1], True),
        (tbl.cell(2, 1), themes_tuples[1][0], themes_tuples[1][1], False),
        (tbl.cell(3, 0), themes_tuples[2][0], themes_tuples[2][1], True),
        (tbl.cell(3, 1), themes_tuples[3][0], themes_tuples[3][1], False),
    ]

    for idx, (cell, t_head, t_body, is_left) in enumerate(cards_config):
        set_explicit_cell_width(cell, 5100)
        set_cell_background(cell, "FFFFFF")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        left_b = {'sz': 24, 'val': 'single', 'color': '1B4D3E'} if is_left else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        right_b = {'sz': 4, 'val': 'single', 'color': '4A6C5C'} if not is_left else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        bottom_b = {'sz': 4, 'val': 'single', 'color': '4A6C5C'} if idx >= 2 else {'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}
        set_cell_border(cell, left=left_b, right=right_b, top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT}, bottom=bottom_b)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        if t_head:
            r_th = p.add_run(t_head.replace("PME", "entreprise") + "\n")
            r_th.font.name = FONT_SERIF
            r_th.font.bold = True
            r_th.font.size = Pt(8)
            r_th.font.color.rgb = BRAND_GREEN
            r_tb = p.add_run(t_body.replace("PME", "entreprise"))
            r_tb.font.name = FONT_SANS
            r_tb.font.size = Pt(7.5)
            r_tb.font.color.rgb = BRAND_GREEN

# 3AB. FINANCIAL ROI BOX
def add_financial_roi_box(doc, gain_text, cout_text):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("ANCRAGE FINANCIER & IMPACT ROI SUR LA MARGE DE L'ENTREPRISE")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, CREAM_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("GAIN STRATÉGIQUE OU PROTECTION DE MARGE BRUTE (€) :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = GREEN_TEXT
    r1_b = p1.add_run("- " + gain_text.replace("PME", "entreprise"))
    r1_b.font.name = FONT_SANS
    r1_b.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, CREAM_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("COÛT DE NON-QUALITÉ OU PERTE EN CAS D'INACTION (€) :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RED_TEXT
    r2_b = p2.add_run("- " + cout_text.replace("PME", "entreprise"))
    r2_b.font.name = FONT_SANS
    r2_b.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

# 3B. INTERCONNECTIONS BOX
def add_interconnections_box(doc, inputs_list, outputs_list):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, CREAM_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("MATRICE D'INTERCONNEXIONS & FLUX DE DONNÉES ENTRANTES / SORTANTES")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, CREAM_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("DONNÉES ENTRANTES (D'où viennent les informations ?) :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = BRAND_GREEN
    for inp in inputs_list:
        p = c1.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + inp.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, CREAM_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("DONNÉES SORTANTES (Vers quoi cet outil alimente-t-il ?) :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = BRAND_GREEN
    for out in outputs_list:
        p = c2.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + out.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_why_box(doc, gains_list, risques_list):
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_title = tbl.cell(0, 0)
    c_title_r = tbl.cell(0, 1)
    c_title.merge(c_title_r)
    set_explicit_cell_width(c_title, 10200)
    set_cell_background(c_title, WHY_BG)
    set_cell_margins(c_title, top=80, bottom=50, left=120, right=120)
    set_cell_border(c_title, top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                             bottom={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p_t = c_title.paragraphs[0]
    p_t.paragraph_format.space_after = Pt(0)
    r_t = p_t.add_run("POURQUOI CET OUTIL EST-IL ABSOLUMENT CAPITAL POUR VOTRE ENTREPRISE ?")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    c1 = tbl.cell(1, 0)
    set_explicit_cell_width(c1, 5100)
    set_cell_background(c1, WHY_BG)
    set_cell_margins(c1, top=60, bottom=80, left=120, right=120)
    set_cell_border(c1, left={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("Ce que vous gagnez à l'utiliser :\n")
    r1.font.name = FONT_SERIF
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = GREEN_TEXT
    for g in gains_list:
        p = c1.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + g.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    c2 = tbl.cell(1, 1)
    set_explicit_cell_width(c2, 5100)
    set_cell_background(c2, WHY_BG)
    set_cell_margins(c2, top=60, bottom=80, left=120, right=120)
    set_cell_border(c2, right={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                         left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                         top={'sz': 4, 'val': 'single', 'color': 'C5A065'})
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Ce que vous risquez sans cet outil :\n")
    r2.font.name = FONT_SERIF
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RED_TEXT
    for r_item in risques_list:
        p = c2.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("- " + r_item.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_guide_box(doc, steps_list):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_explicit_cell_width(cell, 10200)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_border(cell, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    r0 = p0.add_run("GUIDE DE REMPLISSAGE PAS À PAS (4 ÉTAPES)\n")
    r0.font.name = FONT_SERIF
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = BRAND_GREEN
    
    for idx, step in enumerate(steps_list, 1):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r_num = p.add_run(f"Étape {idx} : ")
        r_num.font.name = FONT_SANS
        r_num.font.bold = True
        r_num.font.size = Pt(8)
        r_num.font.color.rgb = BRAND_GREEN
        r_txt = p.add_run(step.replace("PME", "entreprise"))
        r_txt.font.name = FONT_SANS
        r_txt.font.size = Pt(8)
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_next_steps_box(doc, next_list):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_explicit_cell_width(cell, 10200)
    set_cell_background(cell, CREAM_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_border(cell, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                          top={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                          bottom={'sz': 6, 'val': 'single', 'color': 'C5A065'},
                          right={'sz': 6, 'val': 'single', 'color': 'C5A065'})
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    r0 = p0.add_run("QUE FAIRE UNE FOIS CET OUTIL COMPLÉTÉ ?\n")
    r0.font.name = FONT_SERIF
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = BRAND_GREEN
    
    for n in next_list:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run("- " + n.replace("PME", "entreprise"))
        r.font.name = FONT_SANS
        r.font.size = Pt(8)
        r.font.color.rgb = BRAND_GREEN
        
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def format_word_table(doc, headers, rows_data, col_widths_dxa=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        clean_h = h_text.replace("PME", "entreprise")
        hdr_cells[i].text = clean_h
        if col_widths_dxa and i < len(col_widths_dxa):
            set_explicit_cell_width(hdr_cells[i], col_widths_dxa[i])
        set_cell_background(hdr_cells[i], "1B4D3E")
        set_cell_margins(hdr_cells[i], top=60, bottom=60, left=60, right=60)
        set_cell_border(hdr_cells[i], top={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      bottom={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      left={'sz': 4, 'val': 'single', 'color': '1B4D3E'},
                                      right={'sz': 4, 'val': 'single', 'color': '1B4D3E'})
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE
                
    for r_idx, row_data in enumerate(rows_data):
        row_cells = tbl.add_row().cells
        bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            clean_val = val.replace("🟢", "").replace("🔴", "").replace("🟠", "").replace("🟡", "").replace("PME", "entreprise").strip()
            row_cells[c_idx].text = clean_val
            if col_widths_dxa and c_idx < len(col_widths_dxa):
                set_explicit_cell_width(row_cells[c_idx], col_widths_dxa[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=60, right=60)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)
                    if "[Votre" in val or "____" in val or "[Saisissez" in val:
                        r.font.italic = True
                        r.font.bold = True
                        r.font.color.rgb = GOLD
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def add_section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    set_explicit_cell_width(c, 10200)
    set_cell_background(c, THEMES_BG)
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    set_cell_border(c, left={'sz': 20, 'val': 'single', 'color': '1B4D3E'},
                       top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                       bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                       right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = FONT_SERIF
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = BRAND_GREEN
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(4)

def save_doc(doc, filename):
    doc.save(os.path.join(docx_dir, filename))

# ----------------------------------------------------
# 6. OUTIL A6: ISHIKAWA 5M
# ----------------------------------------------------
def make_outil_a6():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A6", "Diagramme d'Ishikawa (Cause à Effet - 5M)", "Méthode Transversale - Analyse de Causes Racine")
    add_consultant_box(doc, "Analysez les causes d'un problème complexe en explorant les 5 M : Matière, Matériel, Méthode, Main-d'œuvre et Milieu. Ne traitez pas le symptôme, éliminez la cause racine !", [
        ("1. Matière", "Qualité des composants, conformité des matières premières, lots sous-traités."),
        ("2. Matériel", "État des machines, usure des outils, réglages, métrologie et maintenance."),
        ("3. Méthode", "Procédures, gammes d'usinage, instructions visuelles et clarté des consignes."),
        ("4. Main-d'œuvre & Milieu", "Formation, compétences, fatigue, bruit, éclairage et température atelier.")
    ])
    add_why_box(doc, [
        "Vision globale à 360° sans oublier d'axe explicatif.",
        "Implication collective de l'équipe pluridisciplinaire.",
        "Détection efficace des causes cachées."
    ], [
        "Recherche de boucs émissaires au lieu de chercher les failles système.",
        "Réaction impulsive sans traiter l'origine réelle du problème."
    ])
    add_financial_roi_box(doc,
        "Élimination des dysfonctionnements complexes et des pannes inter-services.",
        "Accumulation de surcoûts d'exploitation par non-traitement de la cause racine."
    )
    add_interconnections_box(doc, [
        "OUTIL 19 (FNC & 5 Pourquoi) : Dysfonctionnement complexe inter-services.",
        "OUTIL 10 (Réclamations Client) : Problème majeur remonté par un client."
    ], [
        "OUTIL A5 (Cycle PDCA) : Lancement des chantiers d'actions par axe 5M.",
        "OUTIL 05 (AMDEC Risques) : Mise à jour de la cartographie des risques."
    ])
    add_guide_box(doc, [
        "Poser le problème central (Effet) clairement à droite.",
        "Renseigner les causes potentielles dans les 5 branches M.",
        "Sélectionner les 3 causes majeures par vote ou mesures factuelles.",
        "Lancer le plan d'action préventif associé."
    ])
    doc.add_page_break()
    add_section_header(doc, "MATRICE D'ANALYSE ISHIKAWA (LES 5M)")
    headers = ["Axe 5M", "Causes Potentielles Identifiées", "Facteur Critique (Oui/Non)", "Cause Racine Retenue", "Action Corrective Associée", "Responsable"]
    rows = [
        ["1. Matière", "Variabilité de la nuance inox reçue du fournisseur", "Oui", "Absence de certificat 3.1 contrôlé", "Contrôle 100% certificats à la réception", "Resp. Achats"],
        ["2. Matériel", "Jeu sur le roulement de la broche CNC-02", "Oui", "Dépassement du délai de maintenance", "Remplacement roulement broche", "Resp. Maintenance"],
        ["3. Méthode", "Instructions d'ébavurage floues sur la gamme", "Non", "Gamme non mise à jour Indice B", "Rédaction fiche visuelle Outil 14", "Resp. BE"],
        ["4. Main-d'œuvre", "Intérimaire non formé au contrôle 3D", "Oui", "Absence d'émargement du livret d'accueil", "Formation 2h au poste par tuteur", "Chef Prod"],
        ["5. Milieu", "Éclairage insuffisant sur le poste métrologie", "Non", "Ampoule défectueuse au poste", "Remplacement néon LED haute précision", "Services Généraux"],
        ["[Axe 5M]", "[Causes identifiées]", "[Oui/Non]", "[Cause racine]", "[Action corrective]", "[Nom]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1500, 2600, 1100, 2000, 2000, 1000])
    add_next_steps_box(doc, [
        "Valider le diagramme avec l'équipe projet.",
        "Archiver sous 03_OPERATIONS_PRODUCTION/ISHIKAWA_ANALYSIS/."
    ])
    save_doc(doc, "OUTIL_A6_ISHIKAWA_5M.docx")

# ----------------------------------------------------
# 7. OUTIL A7: MATRICE EISENHOWER
# ----------------------------------------------------
def make_outil_a7():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A7", "Matrice d'Eisenhower (Gestion du Temps & Priorités)", "Méthode Transversale - Management & Efficacité Personnel")
    add_consultant_box(doc, "Ne laissez plus l'urgence quotidienne étouffer les projets stratégiques ! Triez vos tâches selon leur urgence et leur importance pour vous concentrer sur l'essentiel.", [
        ("1. Urgent & Important (À FAIRE)", "Gestion de crise, réclamation client bloquante, arrêt machine."),
        ("2. Non Urgent mais Important (À PLANIFIER)", "Stratégie, formation, prévention, rédaction de procédures (Zone de valeur)."),
        ("3. Urgent mais Non Important (À DÉLÉGUER)", "Interruptions, appels opportuns, saisies secondaires."),
        ("4. Non Urgent & Non Important (À ÉLIMINER)", "Tâches chronophages inutiles, bruits de couloir, emails secondaires.")
    ])
    add_why_box(doc, [
        "Sérénité et réduction du stress des dirigeants et managers.",
        "Avancement réel des projets structurants à forte valeur ajoutée.",
        "Gestion optimale du temps et respect des échéances."
    ], [
        "Mode 'pompier' permanent où l'on subit les urgences des autres.",
        "Abandon des projets de fond indispensables au développement."
    ])
    add_guide_box(doc, [
        "Lister l'ensemble des tâches en attente sur votre bureau.",
        "Positionner chaque tâche dans l'un des 4 cadrans d'Eisenhower.",
        "Planifier immédiatement les tâches du Cadran 2 dans votre agenda.",
        "Déléguer le Cadran 3 et supprimer le Cadran 4."
    ])
    doc.add_page_break()
    add_section_header(doc, "MATRICE DE TRI ET PLANIFICATION DES PRIORITÉS (EISENHOWER)")
    headers = ["Cadran", "Intitulé de la Tâche / Projet", "Niveau Urgence", "Niveau Importance", "Décision (Faire/Planifier/Déléguer/Jeter)", "Responsable / Date"]
    rows = [
        ["Cadran 1", "Traiter la FNC client ALFA sur lot vannes", "Urgent", "Important", "À FAIRE Immédiatement (Aujourd'hui)", "Resp. QSE / 16 Fév"],
        ["Cadran 2", "Rédiger le plan de polyvalence binômes", "Non Urgent", "Important", "À PLANIFIER dans l'agenda (Mardi 14h)", "Resp. RH / 20 Fév"],
        ["Cadran 3", "Saisie manuelle des bons de livraison ERP", "Urgent", "Non Important", "À DÉLÉGUER à l'assistant logistique", "Assistant / 17 Fév"],
        ["Cadran 4", "Rangement vieux prospectus commerciaux", "Non Urgent", "Non Important", "À ÉLIMINER / Reporter", "Supprimé"],
        ["[Cadran]", "[Description tâche]", "[Urgence]", "[Importance]", "[Décision]", "[Nom & Date]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1200, 2600, 1100, 1100, 2700, 1500])
    add_next_steps_box(doc, [
        "Mettre à jour la matrice chaque lundi matin.",
        "Archiver sous 01_DIRECTION_STRATEGIE/GESTION_TEMPS/."
    ])
    save_doc(doc, "OUTIL_A7_MATRICE_EISENHOWER_PRIORITES.docx")

# ----------------------------------------------------
# 8. OUTIL A8: GRILLE QQOQCCP
# ----------------------------------------------------
def make_outil_a8():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A8", "Fiche d'Analyse QQOQCCP (Cadrage Complet de Sujet)", "Méthode Transversale - Cadrage & Description de Projet")
    add_consultant_box(doc, "Cadrez n'importe quelle situation, cahier des charges ou problème en vous posant les 7 questions fondamentales : Qui, Quoi, Où, Quand, Comment, Combien et Pourquoi.", [
        ("Qui ?", "Qui est concerné ? Qui a détecté ? Qui doit agir ?"),
        ("Quoi ?", "De quoi s'agit-il exactement ? Quel est le produit ou le problème ?"),
        ("Où & Quand ?", "Où le problème est-il apparu ? À quel moment précis (date, heure) ?"),
        ("Comment, Combien & Pourquoi ?", "Comment cela s'est-il produit ? Combien de pièces/€ impactés ? Pourquoi agir ?")
    ])
    add_why_box(doc, [
        "Description factuelle et complète sans oubli d'information.",
        "Cadrage rapide des besoins lors du lancement d'un projet.",
        "Gain de temps en évitant les aller-retours d'explications."
    ], [
        "Informations incomplètes nécessitant 5 emails de relance.",
        "Prise de décision basée sur des impressions au lieu de faits."
    ])
    add_interconnections_box(doc, [
        "Cadrage initial d'un projet, réclamation ou cahier des charges client."
    ], [
        "OUTIL 08 (Revue Devis) : Spécifications nettoyées et factuelles.",
        "OUTIL 15 (Jalons APQP BE) : Exigences de conception validées."
    ])
    add_guide_box(doc, [
        "Ouvrir la fiche dès la survenue d'un sujet à cadrer.",
        "Répondre avec précision aux 7 questions QQOQCCP.",
        "Faire valider la synthèse par les parties prenantes.",
        "Diffuser la fiche comme document de référence du projet."
    ])
    doc.add_page_break()
    add_section_header(doc, "GRILLE FACTUELLE DE CADRAGE QQOQCCP")
    headers = ["Question", "Sous-Question de Précision", "Information Factuelle Renseignée", "Source / Preuve Factuelle"]
    rows = [
        ["1. QUI ?", "Qui est impacté ? Qui a découvert ?", "Client Grand Compte ALFA / Contrôleur Métrologie atelier", "Bon de livraison N° BL-884"],
        ["2. QUOI ?", "Quel est le problème / sujet précis ?", "Cote de diamètre hors tolérance (+0.08mm sur axe inox)", "Rapport de contrôle 3D"],
        ["3. OÙ ?", "Dans quel atelier / machine / dossier ?", "Poste de tournage CNC-05 (Atelier Usinage)", "Fiche suiveuse OF-2026-44"],
        ["4. QUAND ?", "À quelle date / heure / fréquence ?", "Le 15/02/2026 lors de la première pièce de réglage du matin", "Horodatage machine 08h15"],
        ["5. COMMENT ?", "Dans quelles conditions cela s'est produit ?", "Pression de serrage du mandrin trop élevée déformant la pièce", "Constat régleur au poste"],
        ["6. COMBIEN ?", "Combien de pièces / Coût financier ?", "12 pièces impactées / Coût de reprise estimé à 450 € HT", "Comptabilité analytique"],
        ["7. POURQUOI ?", "Pourquoi est-il impératif d'agir ?", "Éviter l'arrêt de ligne d'assemblage du client prévu vendredi", "Planning client engagé"],
        ["[Question]", "[Précision]", "[Information factuelle]", "[Preuve]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1500, 2200, 4300, 2200])
    add_next_steps_box(doc, [
        "Diffuser la fiche de cadrage à l'équipe projet.",
        "Archiver sous 03_OPERATIONS_PRODUCTION/CADRAGE_PROJECTS/."
    ])
    save_doc(doc, "OUTIL_A8_GRILLE_QQOQCCP.docx")

# ----------------------------------------------------
# 9. OUTIL A9: CHECKLIST ONBOARDING OFFBOARDING
# ----------------------------------------------------
def make_outil_a9():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A9", "Checklist Onboarding & Offboarding Salarié", "Méthode Transversale - Gestion RH & Sécurité SI")
    add_consultant_box(doc, "Sécurisez l'arrivée et le départ de chaque salarié ! Évitez l'oubli de remise des EPI ou des accès informatiques, et garantissez la restitution immédiate des badges et fichiers informatiques au départ.", [
        ("1. Arrivée (Onboarding J-7 à J+30)", "Préparation du poste, création compte email/ERP, remise EPI, Livret d'Accueil, Parrainage."),
        ("2. Départ (Offboarding J-14 à J+1)", "Désactivation des accès serveurs/VPN, restitution PC/Téléphone, badge, entretien de départ."),
        ("3. Transmission des Fichiers & GED", "Transfert de propriété des dossiers et fichiers du serveur au responsable de service."),
        ("4. Traçabilité & Visas RH", "Émargement de la checklist par le salarié, le Responsable IT et le Responsable RH.")
    ])
    add_why_box(doc, [
        "Intégration réussie et valorisante pour le nouveau collaborateur.",
        "Sécurité totale des données de l'entreprise (zéro compte actif après départ).",
        "Traçabilité rigoureuse exigée par l'ISO 9001 et le RGPD."
    ], [
        "Compte informatique resté actif 6 mois après le départ d'un salarié.",
        "Matériel de l'entreprise non restitué (PC portable, téléphone, outils).",
        "Perte de fichiers importants restés sur le bureau local d'un ancien salarié."
    ])
    add_interconnections_box(doc, [
        "OUTIL 20 (Organigramme RACI) : Rôles et tuteurs désignés.",
        "OUTIL 22 (Polyvalence) : Profil de compétences requis."
    ], [
        "OUTIL 23 (Livret d'accueil) : Causerie d'intégration émargée.",
        "Sécurité des accès SI & restitution du matériel."
    ])
    add_guide_box(doc, [
        "Déclencher la checklist dès la signature du contrat ou la notification de départ.",
        "Cocher chaque point de contrôle au fur et a mesure de l'avancement.",
        "Faire signer la fiche par le salarié et le Responsable IT/RH.",
        "Classer la checklist au dossier RH du salarié."
    ])
    doc.add_page_break()
    add_section_header(doc, "CHECKLIST DE TRAÇABILITÉ DES ARRIVÉES ET DÉPARTS DE PERSONNEL")
    headers = ["Phase", "Étape / Point de Contrôle", "Action Réalisée", "Responsable", "Échéance", "Visa / Validation"]
    rows = [
        ["Onboarding (Arrivée)", "1. Matériel informatique & Téléphone préparé", "PC portable configuré avec accès ERP", "Resp. IT", "J-2 avant arrivée", "Validé IT"],
        ["Onboarding (Arrivée)", "2. Remise des EPI & Badge d'accès", "Chaussures de sécurité, lunettes, badge remis", "Resp. QSE", "Jour J à 08h30", "Émargé Salarié"],
        ["Onboarding (Arrivée)", "3. Causerie Sécurité & Livret d'Accueil (Outil 23)", "Entretien 30 min réalisé par tuteur", "Tuteur / RH", "Jour J à 10h00", "Émargé Salarié"],
        ["Offboarding (Départ)", "4. Désactivation des comptes SI & VPN", "Compte email clôturé, accès ERP coupés", "Resp. IT", "Jour D à 17h00", "Validé IT"],
        ["Offboarding (Départ)", "5. Restitution matériel & Badges", "PC, Téléphone, Badge, Clés restitués", "Resp. RH", "Jour D à 16h30", "Émargé Salarié"],
        ["Offboarding (Départ)", "6. Transfert de propriété dossiers GED", "Fichiers locaux transférés sur serveur", "Chef de Service", "Jour D-1", "Validé Chef"],
        ["[Phase]", "[Point de contrôle]", "[Action réalisée]", "[Nom]", "[Date]", "[Visa]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1800, 2800, 2400, 1200, 1000, 1000])
    add_next_steps_box(doc, [
        "Intégrer la checklist au dossier administratif du personnel.",
        "Archiver sous 04_RH_SAVOIR_GED_IT/DOSSIERS_RH/ONBOARDING/."
    ])
    save_doc(doc, "OUTIL_A9_CHECKLIST_ONBOARDING_OFFBOARDING.docx")

# ----------------------------------------------------
# 10. OUTIL A10: RAPPORT INCIDENT SIGNALEMENT RISQUE
# ----------------------------------------------------
def make_outil_a10():
    doc = create_base_document()
    add_header_brand(doc, "OUTIL A10", "Fiche de Signalement 'Presque-Accident' & Risque Opérationnel", "Méthode Transversale - Sécurité HSE & Prévention Qualité")
    add_consultant_box(doc, "Détectez les situations dangereuses avant qu'elles n'entraînent un accident de travail ou un sinistre qualité ! Favorisez la remontée d'information terrain bienveillante sans punition.", [
        ("1. Notion de 'Presque-Accident'", "Événement imprévu qui aurait pu entraîner une blessure ou un dégât matériel si les circonstances avaient été légèrement différentes."),
        ("2. Déclaration Simple en 2 Minutes", "Formulaire court accessible à tout salarié à l'atelier ou dans les bureaux."),
        ("3. Analyse Immédiate par le Référent HSE", "Évaluation de la gravité potentielle et mise en place d'une protection temporaire sous 24h."),
        ("4. Action Corrective & Remerciement", "Correction du problème et valorisation du salarié déclarant lors de la Réunion Flash.")
    ])
    add_why_box(doc, [
        "Réduction drastique des accidents de travail et des arrêts maladie.",
        "Culture de prévention proactive partagée par l'ensemble des salariés.",
        "Conformité avec les exigences de sécurité au travail et ISO 45001 / 9001."
    ], [
        "Survenue d'un accident grave évitable qui avait été pressenti par les équipes.",
        "Climat d'insécurité et rétention d'information par peur de sanctions."
    ])
    add_interconnections_box(doc, [
        "Observations directes des salariés sur le terrain et à l'atelier."
    ], [
        "OUTIL 24 (Réunion Flash) : Communication sécurité hebdomadaire.",
        "OUTIL 05 (AMDEC) : Mise à jour de la cartographie des risques HSE."
    ])
    add_guide_box(doc, [
        "Remplir la fiche dès qu'une situation dangereuse ou un presque-accident est constaté.",
        "Déposer la fiche dans la boîte à idées / signalements ou transmettre au QSE.",
        "Mettre en place la mesure de protection sous 24h.",
        "Présenter l'action corrective lors de la Réunion Flash du lundi."
    ])
    doc.add_page_break()
    add_section_header(doc, "REGISTRE DE SIGNALEMENT DES PRESQUE-ACCIDENTS ET SITUATIONS DANGEREUSES")
    headers = ["N° Signalement", "Date", "Lieu / Zone", "Description de la Situation Dangereuse", "Gravité Potentielle", "Mesure Immédiate Prise", "Action Corrective", "Responsable & Échéance"]
    rows = [
        ["SIG-2026-01", "10/02", "Zone Stockage Matière", "Huile au sol sous le chariot élévateur (risque de glissade)", "Moyenne (Chute)", "Épandage d'absorbant immédiatement", "Réparation fuite hydraulique chariot", "Resp. Maintenance / 11 Fév"],
        ["SIG-2026-02", "14/02", "Atelier Meulage", "Protecteur plexiglas de la meuleuse desserré", "Élevée (Projectiles)", "Mise hors service temporaire machine", "Fixation renforcée du protecteur", "Régleur / 14 Fév"],
        ["SIG-2026-03", "16/02", "Bureau d'Études", "Câble réseau traversant l'allée de circulation", "Faible (Trébuchement)", "Mise en place d'un passe-câble au sol", "Câblage définitif sous plinthe", "Services Généraux / 18 Fév"],
        ["[N°]", "[Date]", "[Zone]", "[Description du risque]", "[Gravité]", "[Mesure immédiate]", "[Action corrective]", "[Nom & Date]"]
    ]
    format_word_table(doc, headers, rows, col_widths_dxa=[1200, 900, 1400, 2400, 1100, 1600, 1600, 1000])
    add_next_steps_box(doc, [
        "Remercier le déclarant lors de la Réunion Flash 5 min.",
        "Archiver sous 04_RH_SAVOIR_GED_IT/SIGNALEMENTS_HSE/."
    ])
    save_doc(doc, "OUTIL_A10_RAPPORT_INCIDENT_SIGNALEMENT_RISQUE.docx")

# RUN BUILDERS A6-A10
make_outil_a6()
make_outil_a7()
make_outil_a8()
make_outil_a9()
make_outil_a10()

# HTML & MD FOR A6-A10
annex_tools_data_2 = [
    ("OUTIL_A6_ISHIKAWA_5M", "OUTIL A6 - Diagramme Ishikawa 5M", "OUTIL A6 — Méthode Transversale - Analyse", "Diagramme d'Ishikawa (Cause à Effet - 5M)", "Analysez les causes d'un problème en explorant les 5 M : Matière, Matériel, Méthode, Main-d'œuvre et Milieu."),
    ("OUTIL_A7_MATRICE_EISENHOWER_PRIORITES", "OUTIL A7 - Matrice Eisenhower", "OUTIL A7 — Méthode Transversale - Management", "Matrice d'Eisenhower (Gestion du Temps & Priorités)", "Ne laissez plus l'urgence quotidienne étouffer les projets stratégiques."),
    ("OUTIL_A8_GRILLE_QQOQCCP", "OUTIL A8 - Fiche QQOQCCP", "OUTIL A8 — Méthode Transversale - Cadrage", "Fiche d'Analyse QQOQCCP (Cadrage Complet de Sujet)", "Cadrez n'importe quelle situation ou projet en posant les 7 questions fondamentales."),
    ("OUTIL_A9_CHECKLIST_ONBOARDING_OFFBOARDING", "OUTIL A9 - Onboarding Offboarding", "OUTIL A9 — Méthode Transversale - Gestion RH", "Checklist Onboarding & Offboarding Salarié", "Sécurisez l'arrivée et le départ de chaque salarié et verrouillez les accès informatiques."),
    ("OUTIL_A10_RAPPORT_INCIDENT_SIGNALEMENT_RISQUE", "OUTIL A10 - Presque-Accident & Risques", "OUTIL A10 — Méthode Transversale - Sécurité HSE", "Fiche de Signalement 'Presque-Accident' & Risque Opérationnel", "Détectez les situations dangereuses avant qu'elles n'entraînent un accident.")
]

for filename_base, title, meta, tool_title, advice in annex_tools_data_2:
    html_content = f"""<!DOCTYPE html>
<html lang="fr">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{title} | Cabinet Delacour</title>
  <style>
    body {{ font-family: Arial, sans-serif; background-color: #F9F9F7; color: #152822; padding: 24px; }}
    .container {{ max-width: 1000px; margin: 0 auto; background: #FFF; padding: 32px; border-radius: 4px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); }}
    .header-banner {{ background: #152822; border-bottom: 4px solid #C5A065; color: #F0EAD6; padding: 16px 24px; font-weight: bold; }}
    .tool-title-box {{ background: #F9F9F7; border-left: 5px solid #C5A065; padding: 16px; margin: 20px 0; }}
    .tool-meta {{ color: #C5A065; font-size: 12px; font-weight: bold; text-transform: uppercase; }}
    .tool-title {{ font-family: Georgia, serif; font-size: 22px; color: #152822; margin-top: 4px; }}
    .consultant-box {{ background: #F7F9F8; border-left: 5px solid #4A6C5C; padding: 16px; margin-bottom: 20px; }}
  </style>
</head>
<body>
<div class="container">
  <div class="header-banner">CABINET DELACOUR - OUTILS ANNEXES TRANSVERSAUX</div>
  <div class="tool-title-box">
    <div class="tool-meta">{meta}</div>
    <div class="tool-title">{tool_title}</div>
  </div>
  <div class="consultant-box">
    <strong>Conseil Consultant :</strong> {advice}
  </div>
  <p>Fichier modèle complet disponible en format Word (.docx) dans votre dossier client.</p>
</div>
</body>
</html>"""
    with open(os.path.join(html_dir, f"{filename_base}.html"), "w", encoding="utf-8") as f:
        f.write(html_content)
        
    md_content = f"# 🧰 {tool_title}\n\n> **{meta}**\n\n## 💡 Conseil Consultant\n{advice}\n\n*Consulter la grille complète dans le fichier Word `.docx` rattaché.*"
    with open(os.path.join(md_dir, f"{filename_base}.md"), "w", encoding="utf-8") as f:
        f.write(md_content)

print("Successfully generated Annex Tools A6 to A10 (DOCX, HTML, MD)!")
