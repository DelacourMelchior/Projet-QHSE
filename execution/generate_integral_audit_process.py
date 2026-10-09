#!/usr/bin/env python3
"""
Générateur du Processus Intégral d'Intervention Audit SMQ en 4 Phases Strictes (Cabinet Delacour)
Aligné sur les 5 Sessions Métiers et sans parenthèses ni émojis dans les titres.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
THEMES_BG    = "F4F6F4"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def generate_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(7.1))
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - Processus d'Intervention Audit SMQ en 4 Phases\t\tPage ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        frun2 = fp.add_run(" / ")
        frun2.font.name = FONT_SANS
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED_GRAY

        f_xml5 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml5)
        f_xml6 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml6)
        f_xml7 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml7)
        f_xml8 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml8)
        
        tab_xml = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:pos="10512"/></w:tabs>' % nsdecls('w'))
        fp._p.get_or_add_pPr().append(tab_xml)

    # Title Banner
    tbl_b = doc.add_table(rows=1, cols=1)
    tbl_b.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b.autofit = False
    cb = tbl_b.cell(0, 0)
    set_explicit_cell_width(cb, 10200)
    set_cell_background(cb, HEADER_BG)
    set_cell_margins(cb, top=140, bottom=140, left=160, right=160)
    set_cell_border(cb, top={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        bottom={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        left={'sz': 4, 'val': 'single', 'color': BORDER_GOLD},
                        right={'sz': 4, 'val': 'single', 'color': BORDER_GOLD})
    
    pb = cb.paragraphs[0]
    pb.paragraph_format.space_after = Pt(2)
    rb1 = pb.add_run("PROCESSUS D'INTERVENTION AUDIT ET LIVRAISON DU KIT QUALITÉ\n")
    rb1.font.name = FONT_SERIF
    rb1.font.size = Pt(14)
    rb1.font.bold = True
    rb1.font.color.rgb = ACCENT_GOLD
    
    rb2 = pb.add_run("Manuel Opérationnel d'Intervention : Du Cadrage Pré-Audit à la Remise du Rapport & Kit d'Outils")
    rb2.font.name = FONT_SANS
    rb2.font.size = Pt(9.5)
    rb2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

    # Note Périmètre
    b_note = doc.add_table(rows=1, cols=1)
    b_note.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_n = b_note.cell(0, 0)
    set_explicit_cell_width(c_n, 10200)
    set_cell_background(c_n, CREAM_BG)
    set_cell_margins(c_n, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_n, left={'sz': 20, 'val': 'single', 'color': '1B4D3E'},
                          top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                          right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    p_n = c_n.paragraphs[0]
    p_n.paragraph_format.space_after = Pt(2)
    r_nh = p_n.add_run("CLARIFICATION DU PÉRIMÈTRE DE LA PRESTATION AUDIT :\n")
    r_nh.font.name = FONT_SERIF
    r_nh.font.bold = True
    r_nh.font.size = Pt(9.5)
    r_nh.font.color.rgb = BRAND_GREEN
    
    r_nb = p_n.add_run(
        "La prestation d'audit s'exécute strictement en 4 Phases chronologiques et s'achève avec la livraison du Rapport Complet d'Audit + Plan d'Action Sur-Mesure et la transmission du Pack des Outils Qualité Delacour."
    )
    r_nb.font.name = FONT_SANS
    r_nb.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Process Table (4 Phases & 5 Sessions Sur-Site)
    headers_p = ["Phase", "Timing / Horizon", "Actions Opérationnelles Consultant", "Livrable & Support Transmis au Client"]
    tbl_p = doc.add_table(rows=1, cols=len(headers_p))
    tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p.autofit = False
    widths_p = [2000, 1800, 3800, 2600]
    for i, h in enumerate(headers_p):
        tbl_p.rows[0].cells[i].text = h
        set_explicit_cell_width(tbl_p.rows[0].cells[i], widths_p[i])
        set_cell_background(tbl_p.rows[0].cells[i], HEADER_BG)
        set_cell_margins(tbl_p.rows[0].cells[i], top=60, bottom=60, left=50, right=50)
        p = tbl_p.rows[0].cells[i].paragraphs[0]
        for r in p.runs:
            r.font.name = FONT_SANS
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_rows = [
        ["PHASE 1\nPré-Audit Off-Site", "J-15 à J-2", "- Envoi du Mail Universel de Cadrage Pré-Audit.\n- Réception des 8 documents clés du client.\n- Analyse Off-Site en 30-45 min selon la grille de lecture.\n- Pré-remplissage de 30-40% de la Grille d'Audit SMQ.", "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR\n&\nCHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR"],
        ["PHASE 2\nPréparation J-1", "J-1", "- Personnalisation des Slides de Réunion d'Ouverture.\n- Pré-identification des 3 Points Rouges.\n- Pré-sélection des outils recommandés du Kit Delacour.", "SLIDES_REUNION_OUVERTURE_DELACOUR\n&\nGUIDE_METHODOLOGIQUE_EXPLOITATION_PRE_AUDIT_DELACOUR"],
        ["PHASE 3\nJournée sur Site", "Jour J\n(08h30 - 17h45)", "- 08h30-09h00 : Réunion d'ouverture.\n- 09h00-16h45 : Audit des 5 Sessions Métiers (Direction, Commercial/Achats, Prod, RH, Data/IT).\n- 16h45-17h15 : Consolidation du score & synthèse.\n- 17h15-17h45 : Restitution Flash devant le Dirigeant.", "SLIDES_REUNION_OUVERTURE_DELACOUR\nSUPPORT_RESTITUTION_FLASH_DELACOUR\nGRILLE_D_AUDIT_SMQ_DELACOUR"],
        ["PHASE 4\nClôture & Remise Livrables", "J+2 à J+5", "- Finalisation du Rapport d'Audit Complet et Plan d'Action Sur-Mesure.\n- Transmission officielle au client du Rapport d'Audit et du Pack des Outils Qualité Delacour.\n- Clôture officielle de la prestation d'audit.", "RAPPORT_AUDIT_SMQ_ET_PLAN_D_ACTION_DELACOUR.docx\n&\nPACK COMPLET DES OUTILS QUALITÉ DELACOUR"]
    ]
    for r_idx, row in enumerate(p_rows):
        row_cells = tbl_p.add_row().cells
        bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            set_explicit_cell_width(row_cells[c_idx], widths_p[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=50, right=50)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(8)

    doc.save(output_path)

def generate_html(output_path):
    html = """<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Processus d'Intervention Audit SMQ en 4 Phases — Cabinet Delacour</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 30px; background-color: #fcfcfc; color: #222; line-height: 1.6; }
        h1 { font-family: Georgia, serif; color: #1B4D3E; margin-bottom: 5px; }
        .subtitle { font-style: italic; color: #C5A065; margin-bottom: 20px; }
        .box { background-color: #F9F9F7; border-left: 4px solid #1B4D3E; padding: 15px; margin-bottom: 20px; border-radius: 4px; }
        table { width: 100%; border-collapse: collapse; margin-bottom: 20px; font-size: 13px; }
        th { background-color: #1B4D3E; color: white; padding: 10px; text-align: left; }
        td { padding: 8px; border: 1px solid #D0D7D1; }
        tr:nth-child(even) { background-color: #F9F9F7; }
    </style>
</head>
<body>
    <h1>PROCESSUS D'INTERVENTION AUDIT ET LIVRAISON DU KIT QUALITÉ</h1>
    <div class="subtitle">Manuel Opérationnel d'Intervention Audit SMQ — Cabinet Delacour</div>
    
    <div class="box">
        <strong>CLARIFICATION DU PÉRIMÈTRE DE LA PRESTATION AUDIT :</strong><br><br>
        La prestation d'audit s'exécute strictement en 4 Phases chronologiques et s'achève avec la livraison du Rapport Complet d'Audit + Plan d'Action Sur-Mesure et la transmission du Pack des Outils Qualité Delacour.
    </div>

    <table>
        <thead>
            <tr><th>Phase</th><th>Timing</th><th>Actions Opérationnelles Consultant</th><th>Livrables & Supports Transmis</th></tr>
        </thead>
        <tbody>
            <tr>
                <td><strong>PHASE 1</strong><br>Pré-Audit Off-Site</td>
                <td>J-15 à J-2</td>
                <td>- Envoi du Mail Universel de Cadrage.<br>- Réception des 8 documents clés.<br>- Analyse Off-Site en 30-45 min.<br>- Pré-remplissage de 30-40% de la Grille.</td>
                <td>TEMPLATE_EMAIL_PREPARATOIRE<br>CHECKLIST_PREPARATOIRE</td>
            </tr>
            <tr>
                <td><strong>PHASE 2</strong><br>Préparation J-1</td>
                <td>J-1</td>
                <td>- Personnalisation des Slides de Réunion d'Ouverture.<br>- Pré-identification des 3 Points Rouges.<br>- Pré-sélection des outils du Kit Delacour.</td>
                <td>SLIDES_REUNION_OUVERTURE<br>GUIDE_METHODOLOGIQUE_EXPLOITATION</td>
            </tr>
            <tr>
                <td><strong>PHASE 3</strong><br>Journée sur Site</td>
                <td>Jour J<br>(08h30 - 17h45)</td>
                <td>- 08h30-09h00 : Réunion d'ouverture.<br>- 09h00-16h45 : Audit des 5 Sessions Métiers (Direction, Commercial/Achats, Prod, RH, Data/IT).<br>- 16h45-17h15 : Consolidation & Synthèse.<br>- 17h15-17h45 : Restitution Flash.</td>
                <td>SLIDES_REUNION_OUVERTURE<br>SUPPORT_RESTITUTION_FLASH<br>GRILLE_D_AUDIT_SMQ</td>
            </tr>
            <tr>
                <td><strong>PHASE 4</strong><br>Clôture & Livrables</td>
                <td>J+2 à J+5</td>
                <td>- Finalisation du Rapport d'Audit Complet.<br>- Transmission du Rapport et du Kit d'Outils.<br>- Clôture de la prestation.</td>
                <td>RAPPORT_AUDIT_SMQ_ET_PLAN_D_ACTION<br>PACK COMPLET DES OUTILS QUALITÉ</td>
            </tr>
        </tbody>
    </table>
</body>
</html>
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

def generate_md(output_path):
    md = """# PROCESSUS D'INTERVENTION AUDIT ET LIVRAISON DU KIT QUALITÉ
> **Manuel Opérationnel d'Intervention : Du Cadrage Pré-Audit à la Remise du Rapport & Kit d'Outils**

---

### CLARIFICATION DU PÉRIMÈTRE DE LA PRESTATION AUDIT
La prestation d'audit s'exécute strictement en 4 Phases chronologiques et s'achève avec la livraison du Rapport Complet d'Audit + Plan d'Action Sur-Mesure et la transmission du Pack des Outils Qualité Delacour.

---

### TABLEAU SYNTHÉTIQUE DES 4 PHASES D'AUDIT

| Phase | Timing / Horizon | Actions Opérationnelles Consultant | Livrable & Support Transmis au Client |
| :--- | :--- | :--- | :--- |
| **PHASE 1 : Pré-Audit Off-Site** | J-15 à J-2 | - Envoi du Mail de Cadrage Pré-Audit.<br>- Réception des 8 documents clés du client.<br>- Analyse Off-Site en 30-45 min.<br>- Pré-remplissage de 30-40% de la Grille d'Audit SMQ. | `TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR`<br>`CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR` |
| **PHASE 2 : Préparation J-1** | J-1 | - Personnalisation des Slides de Réunion d'Ouverture.<br>- Pré-identification des 3 Points Rouges.<br>- Pré-sélection des outils recommandés. | `SLIDES_REUNION_OUVERTURE_DELACOUR`<br>`GUIDE_METHODOLOGIQUE_EXPLOITATION_PRE_AUDIT_DELACOUR` |
| **PHASE 3 : Journée sur Site** | Jour J (08h30 - 17h45) | - 08h30-09h00 : Réunion d'ouverture.<br>- 09h00-16h45 : Audit des 5 Sessions Métiers (Direction, Commercial/Achats, Prod, RH, Data/IT).<br>- 16h45-17h15 : Consolidation & Synthèse.<br>- 17h15-17h45 : Restitution Flash devant le Dirigeant. | `SLIDES_REUNION_OUVERTURE_DELACOUR`<br>`SUPPORT_RESTITUTION_FLASH_DELACOUR`<br>`GRILLE_D_AUDIT_SMQ_DELACOUR` |
| **PHASE 4 : Clôture & Livrables** | J+2 à J+5 | - Finalisation du Rapport d'Audit Complet.<br>- Transmission officielle du Rapport et du Pack d'Outils.<br>- Clôture officielle de la prestation d'audit. | `RAPPORT_AUDIT_SMQ_ET_PLAN_D_ACTION_DELACOUR.docx`<br>`PACK COMPLET DES OUTILS QUALITÉ DELACOUR` |
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(md)

def main():
    dest_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    os.makedirs(os.path.join(dest_dir, "DOCX"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "HTML"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "MD"), exist_ok=True)

    generate_docx(os.path.join(dest_dir, "DOCX", "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "HTML", "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "MD", "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.md"))

    generate_docx(os.path.join(dest_dir, "02_CADRAGE_ET_PRE_AUDIT", "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "PROCESS_INTEGRAL_AUDIT_SMQ_DELACOUR.md"))

    print("✅ Processus intégral d'audit aligné sur les 5 Sessions Métiers généré avec succès !")

if __name__ == "__main__":
    main()
