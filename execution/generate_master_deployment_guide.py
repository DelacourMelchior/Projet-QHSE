import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
THEMES_BG    = "F4F6F4"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def create_master_guide():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(7.1))
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - delacour.melchior@cabinetdelacour.com\t\tPage ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        frun2 = fp.add_run(" / ")
        frun2.font.name = FONT_SANS
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED_GRAY

        f_xml5 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml5)
        f_xml6 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml6)
        f_xml7 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml7)
        f_xml8 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml8)

    # Title Banner
    tbl_b = doc.add_table(rows=1, cols=1)
    tbl_b.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b.autofit = False
    cb = tbl_b.cell(0, 0)
    set_explicit_cell_width(cb, 10200)
    set_cell_background(cb, HEADER_BG)
    set_cell_margins(cb, top=140, bottom=140, left=160, right=160)
    set_cell_border(cb, top={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        bottom={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        left={'sz': 4, 'val': 'single', 'color': BORDER_GOLD},
                        right={'sz': 4, 'val': 'single', 'color': BORDER_GOLD})
    
    pb = cb.paragraphs[0]
    pb.paragraph_format.space_after = Pt(2)
    rb1 = pb.add_run("MANUEL DE DÉPLOIEMENT & GOUVERNANCE DU SMQ DELACOUR\n")
    rb1.font.name = FONT_SERIF
    rb1.font.size = Pt(14)
    rb1.font.bold = True
    rb1.font.color.rgb = ACCENT_GOLD
    
    rb2 = pb.add_run("Trame Cadre Adaptable sur 3 Phases & Intégration des Dispositifs Client")
    rb2.font.name = FONT_SANS
    rb2.font.size = Pt(9.5)
    rb2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

    # Section 1: Strategic Vision
    p_sec1 = doc.add_paragraph()
    r_sec1 = p_sec1.add_run("1. VISION STRATÉGIQUE, SOUPLESSE & VALORISATION DE L'EXISTANT")
    r_sec1.font.name = FONT_SERIF
    r_sec1.font.size = Pt(12)
    r_sec1.font.bold = True
    r_sec1.font.color.rgb = BRAND_GREEN

    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_after = Pt(4)
    r_t1 = p_t1.add_run(
        "Ce manuel constitue la trame cadre de déploiement d'un Système de Management de la Qualité (SMQ) axé sur la performance. "
        "Principe d'intervention du Cabinet Delacour : chaque entreprise est unique et dispose d'outils, d'historique et de pratiques déjà fonctionnels. "
        "L'objectif n'est jamais de faire table rase, mais d'analyser l'existant, de valoriser ce qui fonctionne et de déployer les outils complémentaires sur-mesure."
    )
    r_t1.font.name = FONT_SANS
    r_t1.font.size = Pt(9)

    # Table 3 Phases Roadmap
    p_sec2 = doc.add_paragraph()
    r_sec2 = p_sec2.add_run("2. FEUILLE DE ROUTE DE DÉPLOIEMENT MODULABLE (3 PHASES)")
    r_sec2.font.name = FONT_SERIF
    r_sec2.font.size = Pt(12)
    r_sec2.font.bold = True
    r_sec2.font.color.rgb = BRAND_GREEN

    headers = ["Phase & Horizon", "Objectif Stratégique Adapté", "Outils (Existant Client ou Kit Delacour)", "Livrables Attendus & Validation"]
    rows = [
        ["PHASE 1 (Mois 1)\nFondations & Governance", "Organiser la gouvernance, clarifier les rôles et intégrer la GED existante.", "[ Conserver les outils de stratégie existants / Déployer Outils 01 à 07 & Outil 20 RACI ]", "Organigramme RACI émargé par la Direction.\nArborescence GED validée.\n1er Tableau de Bord d'activité."],
        ["PHASE 2 (Mois 2-3)\nFlux Ventes & Fab", "Sécuriser les devis, encadrer la sous-traitance et formaliser les pratiques clés.", "[ Intégrer les grilles devis/fournisseurs existantes ou Outils 08, 09, 14, 16 ]", "Procédures Ventes & Achats adaptées au contexte.\nGammes visuelles validées aux postes."],
        ["PHASE 3 (Mois 4+)\nAmélioration & Audit", "Résoudre les défaillances récurrentes, automatiser les KPIs et piloter les audits.", "[ Intégrer l'ERP client ou Déployer Outils 19, 26, 27 & Outils Annexes A1-A10 ]", "Traitement des non-conformités aux causes racines.\nRevues de direction régulières et préparation d'audit."]
    ]

    tbl_r = doc.add_table(rows=len(rows)+1, cols=4)
    tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_r.autofit = False

    widths = [2200, 2600, 2600, 2800]
    hdr_cells = tbl_r.rows[0].cells
    for i, h_text in enumerate(headers):
        set_explicit_cell_width(hdr_cells[i], widths[i])
        set_cell_background(hdr_cells[i], HEADER_BG)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = FONT_SERIF
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl_r.rows[r_idx+1].cells
        bg = CREAM_BG if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            set_explicit_cell_width(row_cells[c_idx], widths[c_idx])
            set_cell_background(row_cells[c_idx], bg)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = FONT_SANS
            r.font.size = Pt(8)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_after = Pt(6)

    # Governance Rules
    p_sec3 = doc.add_paragraph()
    r_sec3 = p_sec3.add_run("3. RÈGLES DE GOUVERNANCE & SOUPLESSE MANAGÉRIALE")
    r_sec3.font.name = FONT_SERIF
    r_sec3.font.size = Pt(12)
    r_sec3.font.bold = True
    r_sec3.font.color.rgb = BRAND_GREEN

    rules = [
        "1. Principe de Sobriété Documentaire : Aucun document n'est créé si un outil existant couvre déjà efficacement l'exigence.",
        "2. Règle du Seuil de Rentabilité : Adaptable selon les arbitrages financiers définis avec la Direction du client.",
        "3. Règle de l'Amélioration Continue : Priorité au traitement des dysfonctionnements réels du terrain.",
        "4. Règle des Rituels d'Animation : Adaptation du rythme des réunions de suivi selon la disponibilité et la culture interne de l'entreprise."
    ]
    for rule in rules:
        pr = doc.add_paragraph()
        pr.paragraph_format.space_after = Pt(2)
        rr = pr.add_run(rule)
        rr.font.name = FONT_SANS
        rr.font.size = Pt(8.5)

    # Section 4: Conduite du Changement
    p_sec4 = doc.add_paragraph()
    r_sec4 = p_sec4.add_run("4. STRATÉGIE D'ALIGNEMENT ORGANISATIONNEL & CONDUITE DU CHANGEMENT")
    r_sec4.font.name = FONT_SERIF
    r_sec4.font.size = Pt(12)
    r_sec4.font.bold = True
    r_sec4.font.color.rgb = BRAND_GREEN

    headers_ch = ["Niveau d'Adhésion", "Cible Collaborateurs", "Modalité d'Accompagnement Adaptée", "Engagement Attendu"]
    tbl_ch = doc.add_table(rows=1, cols=len(headers_ch))
    tbl_ch.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ch.autofit = False
    col_w_ch = [2000, 2400, 3200, 2600]
    for i, h in enumerate(headers_ch):
        tbl_ch.rows[0].cells[i].text = h
        set_explicit_cell_width(tbl_ch.rows[0].cells[i], col_w_ch[i])
        set_cell_background(tbl_ch.rows[0].cells[i], HEADER_BG)
        set_cell_margins(tbl_ch.rows[0].cells[i], top=60, bottom=60, left=50, right=50)
        p = tbl_ch.rows[0].cells[i].paragraphs[0]
        for r in p.runs:
            r.font.name = FONT_SANS
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    ch_rows = [
        ["1. Information & Sensibilisation", "Ensemble du Personnel", "Présentation des objectifs de la démarche et valorisation des acquis.", "Compréhension de la démarche et adhésion des équipes."],
        ["2. Co-construction Terrain", "Opérateurs & Chefs de Poste", "Ateliers participatifs de mise à jour des modes opératoires.", "Consignes de travail simples rédigées et validées avec le terrain."],
        ["3. Accompagnement Manager", "Chefs de Services / Cadres", "Formation au pilotage par la preuve et à l'animation des rituels d'équipe.", "Autonomie managériale et traitement réactif des irritants."],
        ["4. Ancrage & Pérennité", "CODIR & Direction", "Revues de gouvernance adaptées au rythme de l'entreprise.", "Pérennisation du système Qualité et amélioration continue."]
    ]
    for r_idx, row in enumerate(ch_rows):
        row_cells = tbl_ch.add_row().cells
        bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            set_explicit_cell_width(row_cells[c_idx], col_w_ch[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=50, right=50)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(8)

    output_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    os.makedirs(os.path.join(output_dir, "DOCX"), exist_ok=True)
    os.makedirs(os.path.join(output_dir, "HTML"), exist_ok=True)
    os.makedirs(os.path.join(output_dir, "MD"), exist_ok=True)

    docx_path = os.path.join(output_dir, "DOCX", "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.docx")
    doc.save(docx_path)
    print(f"Saved Master Guide DOCX to {docx_path}")

    docx_root_path = os.path.join(output_dir, "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.docx")
    doc.save(docx_root_path)

    # Generate MD version
    md_content = f"""# MANUEL DE DÉPLOIEMENT & GOUVERNANCE DU SMQ DELACOUR
**Trame Cadre Adaptable sur 3 Phases & Intégration des Dispositifs Client**

---

## 1. VISION STRATÉGIQUE, SOUPLESSE & VALORISATION DE L'EXISTANT

Ce manuel constitue la trame cadre de déploiement d'un Système de Management de la Qualité (SMQ) axé sur la performance. Principe d'intervention du Cabinet Delacour : chaque entreprise est unique et dispose d'outils, d'historique et de pratiques déjà fonctionnels. L'objectif n'est jamais de faire table rase, mais d'analyser l'existant, de valoriser ce qui fonctionne et de déployer les outils complémentaires sur-mesure.

---

## 2. FEUILLE DE ROUTE DÉPLOIEMENT EN 3 PHASES MODULABLES

| Phase & Horizon | Objectif Stratégique Adapté | Outils (Existant Client ou Kit Delacour) | Livrables Attendus & Validation |
| :--- | :--- | :--- | :--- |
| **PHASE 1 (Mois 1)**<br>Fondations & Governance | Organiser la gouvernance, clarifier les rôles et intégrer la GED existante. | `[ Conserver les outils de stratégie existants / Déployer Outils 01 à 07 & Outil 20 RACI ]` | Organigramme RACI émargé.<br>Arborescence GED validée.<br>1er Tableau de Bord d'activité. |
| **PHASE 2 (Mois 2-3)**<br>Flux Ventes & Fab | Sécuriser les devis, encadrer la sous-traitance et formaliser les pratiques clés. | `[ Intégrer les grilles devis/fournisseurs existantes ou Outils 08, 09, 14, 16 ]` | Procédures Ventes & Achats adaptées au contexte.<br>Gammes visuelles validées aux postes. |
| **PHASE 3 (Mois 4+)**<br>Amélioration & Audit | Traiter les défaillances récurrentes, automatiser les KPIs et piloter les audits. | `[ Intégrer l'ERP client ou Déployer Outils 19, 26, 27 & Outils Annexes A1-A10 ]` | Traitement des non-conformités aux causes racines.<br>Revues de direction régulières et préparation d'audit. |

---

## 3. RÈGLES DE GOUVERNANCE & SOUPLESSE MANAGÉRIALE
1. **Sobriété Documentaire** : Aucun document n'est créé si un outil existant couvre déjà efficacement l'exigence.
2. **Souplesse Financière** : Ratios et seuils d'arbitrage adaptés aux spécificités du secteur d'activité du client.
3. **Pérennisation** : Priorité absolue au traitement des irritants réels du terrain.

---

## 4. STRATÉGIE D'ALIGNEMENT ORGANISATIONNEL & CONDUITE DU CHANGEMENT

| Niveau d'Adhésion | Cible Collaborateurs | Modalité d'Accompagnement Adaptée | Engagement Attendu |
| :--- | :--- | :--- | :--- |
| **1. Information** | Ensemble du Personnel | Présentation des objectifs & valorisation de l'existant | Adhésion des équipes |
| **2. Co-construction** | Opérateurs & Chefs de Poste | Ateliers participatifs de rédaction des modes opératoires | Consignes validées par le terrain |
| **3. Accompagnement** | Chefs de Services | Formation au pilotage & rituels d'équipe | Autonomie managériale |
| **4. Ancrage** | CODIR & Direction | Revues de gouvernance adaptées au rythme de l'entreprise | Pérennisation du SMQ |
"""
    md_path = os.path.join(output_dir, "MD", "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.md")
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(md_content)
    print(f"Saved Master Guide MD to {md_path}")

    md_root_path = os.path.join(output_dir, "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.md")
    with open(md_root_path, 'w', encoding='utf-8') as f:
        f.write(md_content)

    # Generate HTML version
    html_content = f"""<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Manuel de Déploiement & Gouvernance SMQ Delacour</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 30px; background-color: #fcfcfc; color: #222; }}
        h1 {{ font-family: Georgia, serif; color: #1B4D3E; margin-bottom: 5px; }}
        h2 {{ font-family: Georgia, serif; color: #1B4D3E; border-bottom: 2px solid #C5A065; padding-bottom: 5px; margin-top: 25px; }}
        .subtitle {{ font-style: italic; color: #C5A065; margin-bottom: 20px; }}
        .box {{ background-color: #F9F9F7; border-left: 4px solid #1B4D3E; padding: 15px; margin-bottom: 20px; border-radius: 4px; }}
        table {{ width: 100%; border-collapse: collapse; margin-bottom: 20px; font-size: 13px; }}
        th {{ background-color: #1B4D3E; color: white; padding: 10px; text-align: left; }}
        td {{ padding: 8px; border: 1px solid #D0D7D1; }}
        tr:nth-child(even) {{ background-color: #F9F9F7; }}
    </style>
</head>
<body>
    <h1>MANUEL DE DÉPLOIEMENT & GOUVERNANCE DU SMQ DELACOUR</h1>
    <div class="subtitle">Trame Cadre Adaptable sur 3 Phases & Intégration des Dispositifs Client</div>
    
    <div class="box">
        <strong>1. VISION STRATÉGIQUE, SOUPLESSE & VALORISATION DE L'EXISTANT</strong><br><br>
        Ce manuel constitue la trame cadre de déploiement d'un Système de Management de la Qualité (SMQ) axé sur la performance.<br>
        Principe d'intervention du Cabinet Delacour : chaque entreprise est unique et dispose d'outils et de pratiques déjà fonctionnels. L'objectif n'est jamais de faire table rase, mais d'analyser l'existant, de valoriser ce qui fonctionne et de déployer les outils complémentaires sur-mesure.
    </div>

    <h2>2. FEUILLE DE ROUTE DE DÉPLOIEMENT MODULABLE (3 PHASES)</h2>
    <table>
        <thead>
            <tr><th>Phase & Horizon</th><th>Objectif Stratégique Adapté</th><th>Outils (Existant Client ou Delacour)</th><th>Livrables Attendus</th></tr>
        </thead>
        <tbody>
            <tr><td><strong>PHASE 1 (Mois 1)</strong><br>Fondations & Governance</td><td>Organiser la gouvernance et intégrer la GED existante</td><td>[ Conserver outils existants / Outils 01 à 07 & 20 ]</td><td>Arborescence GED & Organigramme RACI</td></tr>
            <tr><td><strong>PHASE 2 (Mois 2-3)</strong><br>Flux Ventes & Fab</td><td>Sécuriser devis & encadrer la sous-traitance</td><td>[ Intégrer grilles existantes ou Outils 08, 09, 14, 16 ]</td><td>Procédures Ventes/Achats & Gammes Visuelles</td></tr>
            <tr><td><strong>PHASE 3 (Mois 4+)</strong><br>Amélioration & Audit</td><td>Résoudre défaillances & piloter les audits</td><td>[ Intégrer ERP client ou Outils 19, 26, 27 ]</td><td>Revues de direction & Préparation d'Audit</td></tr>
        </tbody>
    </table>

    <h2>3. STRATÉGIE D'ALIGNEMENT ORGANISATIONNEL & CONDUITE DU CHANGEMENT</h2>
    <table>
        <thead>
            <tr><th>Niveau d'Adhésion</th><th>Cible Collaborateurs</th><th>Modalité d'Accompagnement Adaptée</th><th>Engagement Attendu</th></tr>
        </thead>
        <tbody>
            <tr><td>1. Information</td><td>Ensemble du Personnel</td><td>Présentation des objectifs & valorisation de l'existant</td><td>Adhésion des équipes</td></tr>
            <tr><td>2. Co-construction</td><td>Opérateurs & Chefs de Poste</td><td>Ateliers participatifs de mise à jour des modes opératoires</td><td>Consignes validées terrain</td></tr>
            <tr><td>3. Accompagnement</td><td>Chefs de Services</td><td>Formation au pilotage & rituels d'équipe</td><td>Autonomie managériale</td></tr>
            <tr><td>4. Ancrage</td><td>CODIR & Direction</td><td>Revues de gouvernance adaptées au rythme entreprise</td><td>Pérennisation du SMQ</td></tr>
        </tbody>
    </table>
</body>
</html>
"""
    html_path = os.path.join(output_dir, "HTML", "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.html")
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"Saved Master Guide HTML to {html_path}")

    html_root_path = os.path.join(output_dir, "GUIDE_DEPLOIEMENT_ET_GOUVERNANCE_DELACOUR.html")
    with open(html_root_path, 'w', encoding='utf-8') as f:
        f.write(html_content)

if __name__ == "__main__":
    create_master_guide()
