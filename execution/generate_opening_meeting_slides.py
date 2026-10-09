#!/usr/bin/env python3
"""
Générateur de la Présentation / Support de Réunion d'Ouverture d'Audit SMQ (Cabinet Delacour)
Aligné à 100% avec le planning à 5 Sessions Métiers. Style corporate sans émojis et sans parenthèses dans les titres.
Génère le support sous 3 formes : HTML, DOCX (Paysage), MD.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
THEMES_BG    = "F4F6F4"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

SLIDES_DATA = [
    {
        "number": "01",
        "title": "RÉUNION D'OUVERTURE D'AUDIT SMQ & PERFORMANCE",
        "subtitle": "Méthode Cabinet Delacour - Alignement ISO 9001 & Chasse aux Coûts Cachés",
        "type": "cover",
        "content": [
            "Entreprise cliente : [Nom de votre entreprise]",
            "Date de l'audit : [DD/MM/YYYY]",
            "Horaires : 08h30 - 17h45",
            "Dirigeant / Décideur d'entreprise : [Nom, Prénom & Fonction]",
            "Auditeur référent : Cabinet Delacour (delacour.melchior@cabinetdelacour.com)"
        ]
    },
    {
        "number": "02",
        "title": "OBJECTIFS & ESPRIT DE LA MISSION D'AUDIT",
        "subtitle": "Allier la conformité ISO 9001 et la rentabilité concrète du terrain",
        "type": "dual_cards",
        "card1_title": "Objectif 1 : Conformité Normative",
        "card1_body": "- Évaluer la maturité de l'organisation face aux exigences clés d'ISO 9001:2015.\n- Identifier les points forts et sécuriser les processus.\n- Préparer sereinement les étapes de certification.",
        "card2_title": "Objectif 2 : Performance & Valorisation de l'Existant",
        "card2_body": "- S'appuyer sur vos outils et procédures déjà en place sans tout réinventer.\n- Déceler et stopper les pertes d'argent invisibles et fuites de marge.\n- Co-construire une feuille de route 100% sur-mesure pour votre entreprise."
    },
    {
        "number": "03",
        "title": "LES 3 PILIERS DE ROBUSTESSE CABINET DELACOUR",
        "subtitle": "Le socle méthodologique d'évaluation de la performance globale",
        "type": "three_pillars",
        "pillars": [
            ("PILIER 1", "Identification & Maîtrise des Risques", "Détecter et stopper les fuites de marge, les litiges et les coûts cachés de non-qualité."),
            ("PILIER 2", "Structuration & Souveraineté du Savoir", "Sécuriser le capital intellectuel, supprimer la dépendance aux hommes-clés et fluidifier les opérations."),
            ("PILIER 3", "Pilotage Data & Agilité Décisionnelle", "Remplacer l'intuition par la certitude mathématique, neutraliser les contrats toxiques et réagir sous 24h.")
        ]
    },
    {
        "number": "04",
        "title": "DÉROULEMENT & PLANNING DE LA JOURNÉE D'AUDIT",
        "subtitle": "Planning structuré en 5 Sessions Métiers de 08h30 à 17h45",
        "type": "timeline",
        "schedule": [
            ("08h30 - 09h00", "Réunion d'Ouverture & Cadrage", "Direction & Responsable Qualité"),
            ("09h00 - 10h30", "Session 1 : Direction & Pilotage Stratégique", "Direction Générale & COPIL (1h30)"),
            ("10h30 - 11h45", "Session 2 : Commercial, Offre & Approvisionnements", "Direction Commerciale & Achats (1h15)"),
            ("11h45 - 12h45", "Pause Déjeuner & Échanges informels", "Pause repas (1 heure)"),
            ("12h45 - 14h45", "Session 3 : Opérations, Production & Qualité Terrain", "Responsables Production, Chefs d'Atelier (2h00)"),
            ("14h45 - 15h45", "Session 4 : RH, Compétences & Souveraineté du Savoir", "Responsable RH & Managers (1h00)"),
            ("15h45 - 16h45", "Session 5 : Data, IT & Agilité Décisionnelle", "Responsable IT, Data & SI (1h00)"),
            ("16h45 - 17h15", "Consolidation du Score & Synthèse", "Auditeur seul (30 min)"),
            ("17h15 - 17h45", "Restitution Flash & Points Rouges 48h", "Direction Générale & Responsable Qualité")
        ]
    },
    {
        "number": "05",
        "title": "MÉTHODOLOGIE, ALIGNEMENT & RÈGLES DU JEU",
        "subtitle": "Une démarche participative, factuelle et axée sur l'élimination des frictions",
        "type": "list_cards",
        "items": [
            ("Intégration de vos Outils Existants", "Nous analysons et valorisons les dispositifs déjà déployés dans l'entreprise pour construire un SMQ pragmatique et fluide."),
            ("Plan d'Action Sur-Mesure", "Chaque entreprise est unique : la feuille de route d'exécution est adaptée à votre contexte, vos équipes et vos ressources."),
            ("Examen Factuel & Échelle 0 à 3", "Évaluation objective basée sur les preuves réelles (documents, enregistrements) selon le barème : 0 (Écart Majeur) à 3 (Excellence & ROI)."),
            ("Confidentialité Totale", "Le Cabinet Delacour garantit le secret professionnel absolu sur l'ensemble des données, processus et échanges de la mission.")
        ]
    },
    {
        "number": "06",
        "title": "PROCHAINES ÉTAPES & LIVRABLES RESTITUÉS",
        "subtitle": "Des réponses immédiates et un plan de bataille à court terme",
        "type": "dual_cards",
        "card1_title": "Restitution Flash à 17h15 aujourd'hui",
        "card1_body": "- Présentation des 3 à 5 'Points Rouges' (vulnérabilités majeures nécessitant une intervention sous 48h).\n- Restitution orale du niveau global de maturité par session.\n- Échanges à chaud avec la Direction.",
        "card2_title": "Rapport Complet & Plan d'Action sous 5 jours",
        "card2_body": "- Remise de la Grille d'Audit complète et formalisée.\n- Feuille de route d'exécution structurée :\n  - Court terme (15 jours) : Colmatage des fuites de cash.\n  - Moyen terme (3 mois) : Boîte à outils documentaire.\n  - Long terme (6 mois) : Dashboard Data & Audit Blanc."
    },
    {
        "number": "07",
        "title": "QUESTIONS RÉPONSES & LANCEMENT",
        "subtitle": "Avez-vous des questions avant de démarrer la première session d'interview ?",
        "type": "conclusion",
        "text": "Merci pour votre accueil et votre confiance.\n\nNous enchaînons directement avec la SESSION 1 : DIRECTION & PILOTAGE STRATÉGIQUE (Direction Générale & COPIL)."
    }
]

def generate_docx(output_path):
    doc = docx.Document()
    
    # Format Paysage (Landscape)
    for section in doc.sections:
        new_width, new_height = section.page_height, section.page_width
        section.page_width = new_width
        section.page_height = new_height
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(10.0))

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - Support Réunion d'Ouverture Audit SMQ\t\tDiapositive ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        tab_xml = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:pos="15000"/></w:tabs>' % nsdecls('w'))
        fp._p.get_or_add_pPr().append(tab_xml)

    for slide in SLIDES_DATA:
        p_num = doc.add_paragraph()
        p_num.paragraph_format.space_after = Pt(0)
        r_n = p_num.add_run(f"SLIDE {slide['number']}")
        r_n.font.name = FONT_SANS
        r_n.font.size = Pt(9)
        r_n.font.bold = True
        r_n.font.color.rgb = ACCENT_GOLD

        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(slide["title"])
        r_t.font.name = FONT_SERIF
        r_t.font.size = Pt(16)
        r_t.font.bold = True
        r_t.font.color.rgb = BRAND_GREEN

        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(10)
        r_s = p_sub.add_run(slide["subtitle"])
        r_s.font.name = FONT_SANS
        r_s.font.size = Pt(10)
        r_s.font.italic = True
        r_s.font.color.rgb = MUTED_GRAY

        if slide["type"] == "cover":
            box = doc.add_table(rows=1, cols=1)
            box.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = box.cell(0, 0)
            set_explicit_cell_width(c, 14500)
            set_cell_background(c, CREAM_BG)
            set_cell_margins(c, top=140, bottom=140, left=160, right=160)
            set_cell_border(c, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                               top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                               bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                               right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            cp = c.paragraphs[0]
            for line in slide["content"]:
                r = cp.add_run(line + "\n")
                r.font.name = FONT_SANS
                r.font.size = Pt(9.5)

        elif slide["type"] == "dual_cards":
            tbl = doc.add_table(rows=1, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c1, c2 = tbl.cell(0, 0), tbl.cell(0, 1)
            set_explicit_cell_width(c1, 7100)
            set_explicit_cell_width(c2, 7100)
            set_cell_background(c1, CREAM_BG)
            set_cell_background(c2, CREAM_BG)
            set_cell_margins(c1, top=120, bottom=120, left=140, right=140)
            set_cell_margins(c2, top=120, bottom=120, left=140, right=140)
            set_cell_border(c1, top={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                bottom={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                left={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                right={'sz': 6, 'val': 'single', 'color': BORDER_GOLD})
            set_cell_border(c2, top={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                bottom={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                left={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                right={'sz': 6, 'val': 'single', 'color': BORDER_GOLD})
            
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(slide["card1_title"] + "\n\n")
            r1.font.name = FONT_SERIF
            r1.font.bold = True
            r1.font.size = Pt(10.5)
            r1.font.color.rgb = BRAND_GREEN
            r1b = p1.add_run(slide["card1_body"])
            r1b.font.name = FONT_SANS
            r1b.font.size = Pt(9)

            p2 = c2.paragraphs[0]
            r2 = p2.add_run(slide["card2_title"] + "\n\n")
            r2.font.name = FONT_SERIF
            r2.font.bold = True
            r2.font.size = Pt(10.5)
            r2.font.color.rgb = BRAND_GREEN
            r2b = p2.add_run(slide["card2_body"])
            r2b.font.name = FONT_SANS
            r2b.font.size = Pt(9)

        elif slide["type"] == "three_pillars":
            tbl = doc.add_table(rows=1, cols=3)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (tag, p_t, p_d) in enumerate(slide["pillars"]):
                cell = tbl.cell(0, i)
                set_explicit_cell_width(cell, 4700)
                set_cell_background(cell, CREAM_BG)
                set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
                set_cell_border(cell, top={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                      bottom={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                      left={'sz': 6, 'val': 'single', 'color': BORDER_GOLD},
                                      right={'sz': 6, 'val': 'single', 'color': BORDER_GOLD})
                p = cell.paragraphs[0]
                rt = p.add_run(tag + "\n" + p_t + "\n\n")
                rt.font.name = FONT_SERIF
                rt.font.bold = True
                rt.font.size = Pt(10)
                rt.font.color.rgb = BRAND_GREEN
                rd = p.add_run(p_d)
                rd.font.name = FONT_SANS
                rd.font.size = Pt(8.5)

        elif slide["type"] == "timeline":
            tbl = doc.add_table(rows=1, cols=3)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ["Horaires", "Séquence / Étape d'Audit", "Participants Requis"]
            widths = [2500, 7000, 5000]
            for i, h in enumerate(headers):
                tbl.rows[0].cells[i].text = h
                set_explicit_cell_width(tbl.rows[0].cells[i], widths[i])
                set_cell_background(tbl.rows[0].cells[i], HEADER_BG)
                set_cell_margins(tbl.rows[0].cells[i], top=50, bottom=50, left=50, right=50)
                p = tbl.rows[0].cells[i].paragraphs[0]
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(8.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            for r_idx, (t_time, t_step, t_part) in enumerate(slide["schedule"]):
                row_cells = tbl.add_row().cells
                bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
                row_cells[0].text = t_time
                row_cells[1].text = t_step
                row_cells[2].text = t_part
                for i_c in range(3):
                    set_explicit_cell_width(row_cells[i_c], widths[i_c])
                    set_cell_background(row_cells[i_c], bg_hex)
                    set_cell_margins(row_cells[i_c], top=40, bottom=40, left=50, right=50)
                    set_cell_border(row_cells[i_c], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                     bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                     left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                     right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
                    for p in row_cells[i_c].paragraphs:
                        for r in p.runs:
                            r.font.name = FONT_SANS
                            r.font.size = Pt(8)

        elif slide["type"] == "list_cards":
            tbl = doc.add_table(rows=2, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, (item_t, item_b) in enumerate(slide["items"]):
                r_i = idx // 2
                c_i = idx % 2
                cell = tbl.cell(r_i, c_i)
                set_explicit_cell_width(cell, 7100)
                set_cell_background(cell, CREAM_BG)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                set_cell_border(cell, top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                      bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                      left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                      right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
                p = cell.paragraphs[0]
                rt = p.add_run(item_t + "\n")
                rt.font.name = FONT_SERIF
                rt.font.bold = True
                rt.font.size = Pt(9.5)
                rt.font.color.rgb = BRAND_GREEN
                rb = p.add_run(item_b)
                rb.font.name = FONT_SANS
                rb.font.size = Pt(8.5)

        elif slide["type"] == "conclusion":
            box = doc.add_table(rows=1, cols=1)
            box.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = box.cell(0, 0)
            set_explicit_cell_width(c, 14500)
            set_cell_background(c, CREAM_BG)
            set_cell_margins(c, top=140, bottom=140, left=160, right=160)
            set_cell_border(c, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                               top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                               bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                               right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            cp = c.paragraphs[0]
            r = cp.add_run(slide["text"])
            r.font.name = FONT_SANS
            r.font.size = Pt(10)

        # Fin de slide (saut de page sauf la dernière)
        if slide["number"] != "07":
            doc.add_page_break()

    doc.save(output_path)

def generate_html(output_path):
    html = """<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Slides Réunion d'Ouverture d'Audit SMQ — Cabinet Delacour</title>
    <style>
        body { font-family: Arial, sans-serif; background-color: #111; color: #fff; margin: 0; padding: 20px; display: flex; flex-direction: column; align-items: center; }
        .slide { background-color: #ffffff; color: #222; width: 960px; min-height: 540px; padding: 40px; margin-bottom: 30px; border-radius: 8px; box-shadow: 0 10px 25px rgba(0,0,0,0.5); box-sizing: border-box; }
        .slide-num { font-size: 12px; font-weight: bold; color: #C5A065; text-transform: uppercase; margin-bottom: 5px; }
        h1 { font-family: Georgia, serif; color: #1B4D3E; margin-top: 0; margin-bottom: 5px; font-size: 24px; }
        .subtitle { font-style: italic; color: #555; margin-bottom: 25px; font-size: 14px; }
        .box { background-color: #F9F9F7; border-left: 4px solid #1B4D3E; border: 1px solid #D0D7D1; border-left-width: 4px; border-left-color: #1B4D3E; padding: 20px; border-radius: 4px; }
        .grid-2 { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; }
        .grid-3 { display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 15px; }
        .card { background-color: #F9F9F7; border: 1px solid #C5A065; padding: 15px; border-radius: 4px; }
        .card h3 { font-family: Georgia, serif; color: #1B4D3E; margin-top: 0; margin-bottom: 8px; font-size: 16px; }
        table { width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 13px; }
        th { background-color: #1B4D3E; color: white; padding: 8px; text-align: left; }
        td { padding: 8px; border: 1px solid #D0D7D1; }
        tr:nth-child(even) { background-color: #F9F9F7; }
    </style>
</head>
<body>
"""
    for slide in SLIDES_DATA:
        html += f'    <div class="slide">\n'
        html += f'        <div class="slide-num">Slide {slide["number"]}</div>\n'
        html += f'        <h1>{slide["title"]}</h1>\n'
        html += f'        <div class="subtitle">{slide["subtitle"]}</div>\n'

        if slide["type"] == "cover":
            html += '        <div class="box">\n'
            for line in slide["content"]:
                html += f'            <p style="margin:8px 0;">{line}</p>\n'
            html += '        </div>\n'

        elif slide["type"] == "dual_cards":
            html += '        <div class="grid-2">\n'
            html += f'            <div class="card"><h3>{slide["card1_title"]}</h3><p>{slide["card1_body"].replace(chr(10), "<br>")}</p></div>\n'
            html += f'            <div class="card"><h3>{slide["card2_title"]}</h3><p>{slide["card2_body"].replace(chr(10), "<br>")}</p></div>\n'
            html += '        </div>\n'

        elif slide["type"] == "three_pillars":
            html += '        <div class="grid-3">\n'
            for tag, p_t, p_d in slide["pillars"]:
                html += f'            <div class="card"><h3>{tag}<br>{p_t}</h3><p>{p_d}</p></div>\n'
            html += '        </div>\n'

        elif slide["type"] == "timeline":
            html += '        <table><thead><tr><th>Horaires</th><th>Séquence / Étape</th><th>Participants Requis</th></tr></thead><tbody>\n'
            for t_time, t_step, t_part in slide["schedule"]:
                html += f'            <tr><td><strong>{t_time}</strong></td><td>{t_step}</td><td>{t_part}</td></tr>\n'
            html += '        </tbody></table>\n'

        elif slide["type"] == "list_cards":
            html += '        <div class="grid-2">\n'
            for item_t, item_b in slide["items"]:
                html += f'            <div class="card"><h3>{item_t}</h3><p>{item_b}</p></div>\n'
            html += '        </div>\n'

        elif slide["type"] == "conclusion":
            html += f'        <div class="box"><p style="font-size:16px;">{slide["text"].replace(chr(10), "<br>")}</p></div>\n'

        html += '    </div>\n'

    html += "</body>\n</html>\n"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

def generate_md(output_path):
    md = "# SLIDES DE RÉUNION D'OUVERTURE D'AUDIT SMQ & PERFORMANCE\n> **Cabinet Delacour - Support de présentation au format paysage**\n\n---\n\n"
    for slide in SLIDES_DATA:
        md += f"## SLIDE {slide['number']} : {slide['title']}\n"
        md += f"*{slide['subtitle']}*\n\n"
        if slide["type"] == "cover":
            for line in slide["content"]:
                md += f"- {line}\n"
        elif slide["type"] == "dual_cards":
            md += f"### {slide['card1_title']}\n{slide['card1_body']}\n\n### {slide['card2_title']}\n{slide['card2_body']}\n"
        elif slide["type"] == "three_pillars":
            for tag, p_t, p_d in slide["pillars"]:
                md += f"### {tag} : {p_t}\n{p_d}\n\n"
        elif slide["type"] == "timeline":
            md += "| Horaires | Séquence / Étape d'Audit | Participants Requis |\n| :--- | :--- | :--- |\n"
            for t_time, t_step, t_part in slide["schedule"]:
                md += f"| **{t_time}** | {t_step} | {t_part} |\n"
        elif slide["type"] == "list_cards":
            for item_t, item_b in slide["items"]:
                md += f"### {item_t}\n{item_b}\n\n"
        elif slide["type"] == "conclusion":
            md += f"{slide['text']}\n"
        md += "\n---\n\n"
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(md)

def main():
    dest_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    os.makedirs(os.path.join(dest_dir, "DOCX"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "HTML"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "MD"), exist_ok=True)

    generate_docx(os.path.join(dest_dir, "DOCX", "SLIDES_REUNION_OUVERTURE_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "HTML", "SLIDES_REUNION_OUVERTURE_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "MD", "SLIDES_REUNION_OUVERTURE_DELACOUR.md"))

    generate_docx(os.path.join(dest_dir, "03_INTERVENTION_SUR_SITE_JOUR_J", "SLIDES_REUNION_OUVERTURE_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "SLIDES_REUNION_OUVERTURE_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "SLIDES_REUNION_OUVERTURE_DELACOUR.md"))

    print("✅ Slides de réunion d'ouverture 100% alignés avec le planning à 5 Sessions Métiers générés avec succès !")

if __name__ == "__main__":
    main()
