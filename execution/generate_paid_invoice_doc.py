#!/usr/bin/env python3
"""
Générateur de la Facture Définitive Acquittée d'Audit SMQ (Cabinet Delacour)
Format Word (.docx) directement dans 00_AUDIT_ET_GOUVERNANCE/FACTURE_ACQUITTEE_AUDIT_DELACOUR.docx
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"
PAID_GREEN_BG= "E8F5E9"
PAID_GREEN   = RGBColor(0x2E, 0x7D, 0x32)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def generate_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(7.1))
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - Facture Acquittée Audit SMQ\t\tPage ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        frun2 = fp.add_run(" / ")
        frun2.font.name = FONT_SANS
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED_GRAY

        f_xml5 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml5)
        f_xml6 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml6)
        f_xml7 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml7)
        f_xml8 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml8)
        
        tab_xml = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:pos="10512"/></w:tabs>' % nsdecls('w'))
        fp._p.get_or_add_pPr().append(tab_xml)

    # Title Banner
    tbl_b = doc.add_table(rows=1, cols=1)
    tbl_b.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b.autofit = False
    cb = tbl_b.cell(0, 0)
    set_explicit_cell_width(cb, 10200)
    set_cell_background(cb, HEADER_BG)
    set_cell_margins(cb, top=140, bottom=140, left=160, right=160)
    set_cell_border(cb, top={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        bottom={'sz': 12, 'val': 'single', 'color': BORDER_GOLD},
                        left={'sz': 4, 'val': 'single', 'color': BORDER_GOLD},
                        right={'sz': 4, 'val': 'single', 'color': BORDER_GOLD})
    
    pb = cb.paragraphs[0]
    pb.paragraph_format.space_after = Pt(2)
    rb1 = pb.add_run("FACTURE DÉFINITIVE ACQUITTÉE\n")
    rb1.font.name = FONT_SERIF
    rb1.font.size = Pt(14)
    rb1.font.bold = True
    rb1.font.color.rgb = ACCENT_GOLD
    
    rb2 = pb.add_run("Prestation d'Audit SMQ & Performance Opérationnelle - Règlement Intégral Perçu")
    rb2.font.name = FONT_SANS
    rb2.font.size = Pt(9.5)
    rb2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

    # Info Table (Emetteur & Client)
    tbl_inf = doc.add_table(rows=1, cols=2)
    tbl_inf.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_inf.autofit = False
    c_left = tbl_inf.cell(0, 0)
    c_right = tbl_inf.cell(0, 1)
    set_explicit_cell_width(c_left, 5000)
    set_explicit_cell_width(c_right, 5200)
    set_cell_background(c_left, CREAM_BG)
    set_cell_background(c_right, "FFFFFF")
    set_cell_margins(c_left, top=100, bottom=100, left=120, right=120)
    set_cell_margins(c_right, top=100, bottom=100, left=120, right=120)
    set_cell_border(c_left, left={'sz': 16, 'val': 'single', 'color': '1B4D3E'},
                            top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                            bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                            right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    set_cell_border(c_right, left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                             right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})

    p_l = c_left.paragraphs[0]
    p_l.paragraph_format.space_after = Pt(2)
    r_lh = p_l.add_run("ÉMETTEUR (CABINET DELACOUR) :\n")
    r_lh.font.name = FONT_SERIF
    r_lh.font.bold = True
    r_lh.font.size = Pt(9)
    r_lh.font.color.rgb = BRAND_GREEN
    r_lb = p_l.add_run("Cabinet Delacour - Stratégie & Performance\nReprésenté par : Melchior Delacour\nEmail : delacour.melchior@cabinetdelacour.com\nN° Siret : [Saisir N° Siret]\nN° TVA Intracommunautaire : [Saisir TVA]")
    r_lb.font.name = FONT_SANS
    r_lb.font.size = Pt(8)

    p_r = c_right.paragraphs[0]
    p_r.paragraph_format.space_after = Pt(2)
    r_rh = p_r.add_run("CLIENT & FACTURATION :\n")
    r_rh.font.name = FONT_SERIF
    r_rh.font.bold = True
    r_rh.font.size = Pt(9)
    r_rh.font.color.rgb = BRAND_GREEN
    r_rb = p_r.add_run("Raison Sociale : [Nom du Client / Entreprise]\nÀ l'attention de : [Nom du Dirigeant / Décideur]\nAdresse : [Adresse du siège client]\nN° Facture : FAC-2026-[001]\nRef. Devis : DEV-2026-[001]\nDate de Facturation : [Date du jour]")
    r_rb.font.name = FONT_SANS
    r_rb.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Details Table
    headers_d = ["Désignation des Prestations Effectuées", "Phases d'Intervention Incluses", "Prix HT (€)"]
    tbl_d = doc.add_table(rows=1, cols=len(headers_d))
    tbl_d.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_d.autofit = False
    col_w_d = [4500, 4200, 1500]
    for i, h in enumerate(headers_d):
        tbl_d.rows[0].cells[i].text = h
        set_explicit_cell_width(tbl_d.rows[0].cells[i], col_w_d[i])
        set_cell_background(tbl_d.rows[0].cells[i], HEADER_BG)
        set_cell_margins(tbl_d.rows[0].cells[i], top=60, bottom=60, left=50, right=50)
        p = tbl_d.rows[0].cells[i].paragraphs[0]
        for r in p.runs:
            r.font.name = FONT_SANS
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    d_row = tbl_d.add_row().cells
    d_row[0].text = "PRESTATION D'AUDIT SMQ & PERFORMANCE OPÉRATIONNELLE (Forfait Intégral 4 Phases)"
    d_row[1].text = (
        "- Phase 1 : Cadrage pré-audit off-site & revue des 5 documents clés.\n"
        "- Phase 2 : Pré-remplissage de la grille & préparation réunion d'ouverture.\n"
        "- Phase 3 : Intervention d'audit sur-site (08h30-15h00) & Restitution Flash à chaud.\n"
        "- Phase 4 : Rapport d'Audit Complet + Plan d'Action Sur-Mesure + Pack des 27 Outils Qualité Delacour."
    )
    d_row[2].text = "[ Saisir Prix HT ] €"

    for c_idx in range(3):
        set_explicit_cell_width(d_row[c_idx], col_w_d[c_idx])
        set_cell_background(d_row[c_idx], "FFFFFF")
        set_cell_margins(d_row[c_idx], top=60, bottom=60, left=50, right=50)
        set_cell_border(d_row[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                       bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                       left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                       right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
        p = d_row[c_idx].paragraphs[0]
        for r in p.runs:
            r.font.name = FONT_SANS
            r.font.size = Pt(8)

    # Totals Table
    tbl_tot = doc.add_table(rows=4, cols=2)
    tbl_tot.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tot.autofit = False
    
    t_labels = ["Total Hors Taxes (HT) :", "TVA (20,0 %) :", "Montant Total TTC :", "Montant Perçu / Réglé :"]
    t_vals   = ["[ Montant HT ] €", "[ Montant TVA ] €", "[ Montant TTC ] €", "[ Montant TTC Perçu ] €"]
    
    for r_idx in range(4):
        c_l = tbl_tot.rows[r_idx].cells[0]
        c_v = tbl_tot.rows[r_idx].cells[1]
        set_explicit_cell_width(c_l, 7200)
        set_explicit_cell_width(c_v, 3000)
        bg = PAID_GREEN_BG if r_idx == 3 else (HEADER_BG if r_idx == 2 else CREAM_BG)
        set_cell_background(c_l, bg)
        set_cell_background(c_v, bg)
        set_cell_margins(c_l, top=50, bottom=50, left=100, right=100)
        set_cell_margins(c_v, top=50, bottom=50, left=100, right=100)
        
        pl = c_l.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rl = pl.add_run(t_labels[r_idx])
        rl.font.name = FONT_SANS
        rl.font.size = Pt(8.5)
        rl.font.bold = True
        rl.font.color.rgb = PAID_GREEN if r_idx == 3 else (RGBColor(0xFF, 0xFF, 0xFF) if r_idx == 2 else BRAND_GREEN)
        
        pv = c_v.paragraphs[0]
        pv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rv = pv.add_run(t_vals[r_idx])
        rv.font.name = FONT_SANS
        rv.font.size = Pt(8.5)
        rv.font.bold = True
        rv.font.color.rgb = PAID_GREEN if r_idx == 3 else (RGBColor(0xFF, 0xFF, 0xFF) if r_idx == 2 else BRAND_GREEN)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Paid Box Status
    b_paid = doc.add_table(rows=1, cols=1)
    b_paid.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_pd = b_paid.cell(0, 0)
    set_explicit_cell_width(c_pd, 10200)
    set_cell_background(c_pd, PAID_GREEN_BG)
    set_cell_margins(c_pd, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_pd, left={'sz': 24, 'val': 'single', 'color': '2E7D32'},
                           top={'sz': 4, 'val': 'single', 'color': '2E7D32'},
                           bottom={'sz': 4, 'val': 'single', 'color': '2E7D32'},
                           right={'sz': 4, 'val': 'single', 'color': '2E7D32'})
    ppd = c_pd.paragraphs[0]
    ppd.paragraph_format.space_after = Pt(2)
    r_pdh = ppd.add_run("🟢 ATTESTATION DE RÈGLEMENT — FACTURE ACQUITTÉE :\n")
    r_pdh.font.name = FONT_SERIF
    r_pdh.font.bold = True
    r_pdh.font.size = Pt(9.5)
    r_pdh.font.color.rgb = PAID_GREEN
    
    txt_pd = (
        "- Statut du Paiement : FACTURE INTÉGRALEMENT ACQUITTÉE (PAYÉE)\n"
        "- Mode de Règlement : Virement Bancaire sur le compte du Cabinet Delacour\n"
        "- Date de Réception des Fonds : [ Saisir la date du virement client ]\n"
        "- Référence Transaction / Virement : [ Saisir référence virement bancaire ]\n"
        "- Solde Restant Dû : 0,00 € (Aucun reliquat restant à payer)"
    )
    r_pdb = ppd.add_run(txt_pd)
    r_pdb.font.name = FONT_SANS
    r_pdb.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Signatures
    tbl_s = doc.add_table(rows=1, cols=1)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s.autofit = False
    cs = tbl_s.cell(0, 0)
    set_explicit_cell_width(cs, 10200)
    set_cell_background(cs, CREAM_BG)
    set_cell_margins(cs, top=80, bottom=80, left=100, right=100)
    set_cell_border(cs, top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                        bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                        left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                        right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    ps = cs.paragraphs[0]
    ps.add_run(
        "Pour valoir ce que de droit — Facture acquittée valant quittance officielle de paiement.\n"
        "Cabinet Delacour - Melchior Delacour (Cachet & Signature)"
    ).font.size = Pt(8)

    doc.save(output_path)
    print(f"✅ Facture Acquittée Word générée dans : {output_path}")

def main():
    dest_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    output_path = os.path.join(dest_dir, "FACTURE_ACQUITTEE_AUDIT_DELACOUR.docx")
    generate_docx(output_path)

if __name__ == "__main__":
    main()
