#!/usr/bin/env python3
"""
Générateur du Document de Préparation Pré-Audit Client (Cabinet Delacour)
Version corporate sans émojis. Formulaire granulaire item par item sans horaires indicatifs.
Génère CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR en 3 formats (DOCX, HTML, MD).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
THEMES_BG    = "F4F6F4"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

# Structure sans émojis (8 items exacts)
CHECKLIST_ITEMS_BY_SESSION = [
    {
        "session_title": "SÉQUENCE 1 : GOUVERNANCE & PILOTAGE STRATÉGIQUE",
        "items": [
            ["1.1", "Organigramme à jour (ou liste des collaborateurs par service).", "PDF / Word / Org", "[  ] Prêt\n[  ] Non formalisé"],
            ["1.2", "Déclaration de politique / vision ou objectifs de l'année (si formalisés).", "PDF / Word", "[  ] Prêt\n[  ] Non formalisé"]
        ]
    },
    {
        "session_title": "SÉQUENCE 2 : VENTES, MARGES & ACHATS",
        "items": [
            ["2.1", "2 à 3 devis récents représentatifs avec les commandes clients associées.", "Devis PDF / ERP", "[  ] Prêt\n[  ] Non formalisé"],
            ["2.2", "Liste des principaux fournisseurs et sous-traitants clés.", "Excel / ERP / PDF", "[  ] Prêt\n[  ] Non formalisé"]
        ]
    },
    {
        "session_title": "SÉQUENCE 3 : OPÉRATIONS, PRODUCTION & QUALITÉ",
        "items": [
            ["3.1", "Un exemple de gamme de fabrication, consigne de travail ou mode opératoire.", "Fiches papier / ERP", "[  ] Prêt\n[  ] Non formalisé"],
            ["3.2", "Suivi des réclamations clients, non-conformités ou litiges récents (si existant).", "Excel / ERP / Word", "[  ] Prêt\n[  ] Non formalisé"]
        ]
    },
    {
        "session_title": "SÉQUENCE 4 : RH, COMPÉTENCES & GED",
        "items": [
            ["4.1", "Fiches de poste existantes, grille RACI ou matrice des compétences.", "Word / Excel / PDF", "[  ] Prêt\n[  ] Non formalisé"],
            ["4.2", "Vue / capture de l'arborescence de classement de vos fichiers sur votre serveur ou Cloud.", "Capture d'écran / GED", "[  ] Prêt\n[  ] Non formalisé"]
        ]
    }
]

def generate_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(7.1))
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - Checklist Préparatoire Client\t\tPage ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        frun2 = fp.add_run(" / ")
        frun2.font.name = FONT_SANS
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED_GRAY

        f_xml5 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml5)
        f_xml6 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml6)
        f_xml7 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml7)
        f_xml8 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml8)
        
        tab_xml = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:pos="10512"/></w:tabs>' % nsdecls('w'))
        fp._p.get_or_add_pPr().append(tab_xml)

    # Title
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run("CHECKLIST PRÉPARATOIRE & DOSSIER PRÉ-AUDIT CLIENT")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    r_s = p_sub.add_run("Guide à l'usage de l'entreprise - Éléments documentaires à rassembler pour l'audit")
    r_s.font.name = FONT_SANS
    r_s.font.size = Pt(10.5)
    r_s.font.italic = True
    r_s.font.color.rgb = ACCENT_GOLD

    # Instructions Box
    b_inst = doc.add_table(rows=1, cols=1)
    b_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_in = b_inst.cell(0, 0)
    set_explicit_cell_width(c_in, 10200)
    set_cell_background(c_in, CREAM_BG)
    set_cell_margins(c_in, top=100, bottom=100, left=140, right=140)
    set_cell_border(c_in, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                           top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                           bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                           right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    cp = c_in.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    r_head = cp.add_run("CONSIGNES PRATIQUES POUR L'ENTREPRISE AUDITÉE :\n")
    r_head.font.name = FONT_SERIF
    r_head.font.bold = True
    r_head.font.size = Pt(9.5)
    r_head.font.color.rgb = BRAND_GREEN

    txt_in = (
        "- Ne rédigez aucun document spécifique pour l'audit : rassemblez uniquement vos fichiers de travail habituels (papier ou numérique).\n"
        "- Mettez à disposition ces éléments dans un dossier partagé (ou sur clé USB / réseau) ou préparez-les sur écran pour le jour J.\n"
        "- Pour chaque document ci-dessous, cochez la case 'Prêt' ou 'Non formalisé'."
    )
    r_body = cp.add_run(txt_in)
    r_body.font.name = FONT_SANS
    r_body.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 1: Documents par Séquence d'Audit
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(4)
    r_s1 = p_s1.add_run("1. LISTE DES ÉLÉMENTS DOCUMENTAIRES ITEM PAR ITEM")
    r_s1.font.name = FONT_SERIF
    r_s1.font.size = Pt(12)
    r_s1.font.bold = True
    r_s1.font.color.rgb = BRAND_GREEN

    headers_doc = ["N°", "Document / Élément à Rassembler", "Format Habituel", "Statut (Case à cocher)"]
    col_w_doc = [600, 5400, 2200, 2000]

    for session in CHECKLIST_ITEMS_BY_SESSION:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_after = Pt(2)
        r_sec = p_sec.add_run(session["session_title"])
        r_sec.font.name = FONT_SERIF
        r_sec.font.size = Pt(10)
        r_sec.font.bold = True
        r_sec.font.color.rgb = BRAND_GREEN

        tbl_doc = doc.add_table(rows=1, cols=len(headers_doc))
        tbl_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_doc.autofit = False

        for i, h in enumerate(headers_doc):
            tbl_doc.rows[0].cells[i].text = h
            set_explicit_cell_width(tbl_doc.rows[0].cells[i], col_w_doc[i])
            set_cell_background(tbl_doc.rows[0].cells[i], HEADER_BG)
            set_cell_margins(tbl_doc.rows[0].cells[i], top=50, bottom=50, left=50, right=50)
            p = tbl_doc.rows[0].cells[i].paragraphs[0]
            for r in p.runs:
                r.font.name = FONT_SANS
                r.font.size = Pt(7.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for r_idx, item in enumerate(session["items"]):
            row_cells = tbl_doc.add_row().cells
            bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(item):
                row_cells[c_idx].text = val
                set_explicit_cell_width(row_cells[c_idx], col_w_doc[c_idx])
                set_cell_background(row_cells[c_idx], bg_hex)
                set_cell_margins(row_cells[c_idx], top=40, bottom=40, left=50, right=50)
                set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                  bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                  left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                                  right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
                for p in row_cells[c_idx].paragraphs:
                    for r in p.runs:
                        r.font.name = FONT_SANS
                        r.font.size = Pt(7.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 2: Checklist Logistique
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_after = Pt(4)
    r_s2 = p_s2.add_run("2. CHECKLIST LOGISTIQUE & ACCÈS POUR L'AUDIT")
    r_s2.font.name = FONT_SERIF
    r_s2.font.size = Pt(12)
    r_s2.font.bold = True
    r_s2.font.color.rgb = BRAND_GREEN

    headers_log = ["Élément Logistique", "Exigence d'Organisation", "Responsable Client", "Validation"]
    tbl_log = doc.add_table(rows=1, cols=len(headers_log))
    tbl_log.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_log.autofit = False
    col_w_log = [2500, 4300, 1900, 1500]
    for i, h in enumerate(headers_log):
        tbl_log.rows[0].cells[i].text = h
        set_explicit_cell_width(tbl_log.rows[0].cells[i], col_w_log[i])
        set_cell_background(tbl_log.rows[0].cells[i], HEADER_BG)
        set_cell_margins(tbl_log.rows[0].cells[i], top=50, bottom=50, left=50, right=50)
        p = tbl_log.rows[0].cells[i].paragraphs[0]
        for r in p.runs:
            r.font.name = FONT_SANS
            r.font.size = Pt(7.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    log_rows = [
        ["Salle de Réunion", "Une salle calme disponible avec écran/projecteur.", "[ Nom du responsable ]", "[  ] OK"],
        ["Accès Réseau / Wi-Fi", "Accès Wi-Fi invité sécurisé pour l'auditeur.", "[ Service IT / Admin ]", "[  ] OK"],
        ["Disponibilité Intervenants", "Confirmation de la présence des responsables selon le planning.", "[ Responsable Qualité ]", "[  ] OK"],
        ["Visite d'Atelier / Bureaux", "Consignes de sécurité et Équipements de Protection (EPI) si requis.", "[ Responsable Sécurité ]", "[  ] OK"]
    ]
    for r_idx, row in enumerate(log_rows):
        row_cells = tbl_log.add_row().cells
        bg_hex = "F9F9F7" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            set_explicit_cell_width(row_cells[c_idx], col_w_log[c_idx])
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=40, bottom=40, left=50, right=50)
            set_cell_border(row_cells[c_idx], top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              left={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                                              right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
            for p in row_cells[c_idx].paragraphs:
                for r in p.runs:
                    r.font.name = FONT_SANS
                    r.font.size = Pt(7.5)

    doc.save(output_path)

def generate_html(output_path):
    html = """<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Checklist Préparatoire Pré-Audit Client — Cabinet Delacour</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 30px; background-color: #fcfcfc; color: #222; }
        h1 { font-family: Georgia, serif; color: #1B4D3E; margin-bottom: 5px; }
        h2 { font-family: Georgia, serif; color: #1B4D3E; border-bottom: 2px solid #C5A065; padding-bottom: 5px; margin-top: 25px; }
        h3 { font-family: Georgia, serif; color: #1B4D3E; margin-top: 15px; margin-bottom: 5px; }
        .subtitle { font-style: italic; color: #C5A065; margin-bottom: 20px; }
        .box { background-color: #F9F9F7; border-left: 4px solid #1B4D3E; padding: 15px; margin-bottom: 20px; border-radius: 4px; }
        table { width: 100%; border-collapse: collapse; margin-bottom: 15px; font-size: 12px; }
        th { background-color: #1B4D3E; color: white; padding: 8px; text-align: left; }
        td { padding: 6px; border: 1px solid #D0D7D1; }
        tr:nth-child(even) { background-color: #F9F9F7; }
    </style>
</head>
<body>
    <h1>CHECKLIST PRÉPARATOIRE & DOSSIER PRÉ-AUDIT CLIENT</h1>
    <div class="subtitle">Guide à l'usage de l'entreprise - Éléments documentaires à rassembler pour l'audit</div>
    
    <div class="box">
        <strong>CONSIGNES PRATIQUES POUR L'ENTREPRISE AUDITÉE :</strong><br><br>
        - Ne rédigez aucun document spécifique : rassemblez uniquement vos documents et fichiers de travail habituels.<br>
        - Regroupez les éléments dans un dossier partagé (ou sur clé USB / réseau) ou préparez-les sur écran.<br>
        - Pour chaque document ci-dessous, cochez la case 'Prêt' ou 'Non formalisé'.
    </div>

    <h2>1. LISTE DES ÉLÉMENTS DOCUMENTAIRES ITEM PAR ITEM</h2>
"""
    for session in CHECKLIST_ITEMS_BY_SESSION:
        html += f"    <h3>{session['session_title']}</h3>\n"
        html += """    <table>
        <thead>
            <tr><th>N°</th><th>Document / Élément à Rassembler</th><th>Format Habituel</th><th>Statut (Prêt / Non formalisé)</th></tr>
        </thead>
        <tbody>\n"""
        for item in session["items"]:
            html += f"            <tr><td><strong>{item[0]}</strong></td><td>{item[1]}</td><td>{item[2]}</td><td>{item[3].replace(chr(10), '<br>')}</td></tr>\n"
        html += "        </tbody>\n    </table>\n"

    html += """
    <h2>2. CHECKLIST LOGISTIQUE & ACCÈS POUR L'AUDIT</h2>
    <table>
        <thead>
            <tr><th>Élément Logistique</th><th>Exigence d'Organisation</th><th>Responsable Client</th><th>Validation</th></tr>
        </thead>
        <tbody>
            <tr><td>Salle de Réunion</td><td>Une salle calme disponible avec écran/projecteur.</td><td>[ Nom ]</td><td>[ ] OK</td></tr>
            <tr><td>Accès Réseau / Wi-Fi</td><td>Accès Wi-Fi invité sécurisé pour l'auditeur.</td><td>[ IT / Admin ]</td><td>[ ] OK</td></tr>
            <tr><td>Disponibilité Intervenants</td><td>Confirmation de la présence des responsables selon le planning.</td><td>[ Qualité ]</td><td>[ ] OK</td></tr>
        </tbody>
    </table>
</body>
</html>
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

def generate_md(output_path):
    md = """# CHECKLIST PRÉPARATOIRE & DOSSIER PRÉ-AUDIT CLIENT
> **Cabinet Delacour - Formulaire granulaire 100% aligné avec l'Email Préparatoire**

---

## CONSIGNES PRATIQUES POUR L'ENTREPRISE AUDITÉE
- **Ne rédigez aucun document spécifique pour l'audit** : rassemblez uniquement vos documents et fichiers de travail habituels (papier ou numérique).
- **Mettez à disposition ces éléments dans un dossier partagé** (ou sur clé USB / réseau) ou préparez-les sur écran pour le jour J.
- **Pour chaque point ci-dessous**, cochez la case `Prêt` ou `Non formalisé`.

---

## 1. LISTE DES ÉLÉMENTS DOCUMENTAIRES ITEM PAR ITEM

"""
    for session in CHECKLIST_ITEMS_BY_SESSION:
        md += f"### {session['session_title']}\n\n"
        md += "| N° | Document / Élément à Rassembler | Format Habituel | Statut (Case à cocher) |\n"
        md += "| :-: | :--- | :--- | :---: |\n"
        for item in session["items"]:
            statut_md = item[3].replace("\n", "<br>")
            md += f"| **{item[0]}** | {item[1]} | {item[2]} | `{statut_md}` |\n"
        md += "\n"

    md += """---

## 2. CHECKLIST LOGISTIQUE & ACCÈS POUR L'AUDIT

| Élément Logistique | Exigence d'Organisation | Responsable Interne Client | Validation |
| :--- | :--- | :--- | :---: |
| **Salle de Réunion** | Une salle calme disponible avec écran/projecteur | `[ Nom du responsable ]` | `[  ] OK` |
| **Accès Réseau / Wi-Fi** | Accès Wi-Fi invité sécurisé pour l'auditeur | `[ Service IT / Admin ]` | `[  ] OK` |
| **Disponibilité Intervenants** | Confirmation de la présence des responsables selon le planning | `[ Responsable Qualité ]` | `[  ] OK` |
| **Visite d'Atelier / Bureaux** | Consignes de sécurité et Équipements de Protection (EPI) si requis | `[ Responsable Sécurité ]` | `[  ] OK` |
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(md)

def main():
    dest_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    os.makedirs(os.path.join(dest_dir, "DOCX"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "HTML"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "MD"), exist_ok=True)

    generate_docx(os.path.join(dest_dir, "DOCX", "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "HTML", "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "MD", "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.md"))

    generate_docx(os.path.join(dest_dir, "02_CADRAGE_ET_PRE_AUDIT", "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.md"))

    print("✅ Checklist Préparatoire sans émojis générée avec succès !")

if __name__ == "__main__":
    main()
