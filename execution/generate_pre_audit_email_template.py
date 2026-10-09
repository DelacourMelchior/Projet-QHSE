#!/usr/bin/env python3
"""
Générateur du Template d'Email Préparatoire (Cabinet Delacour)
Version corporate sans émojis. Aligné à 100% (miroir parfait 8 items) avec la CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.
Génère TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR dans les 3 formats (DOCX, HTML, MD).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_SERIF = "Georgia"
FONT_SANS = "Arial"

BRAND_GREEN = RGBColor(0x1B, 0x4D, 0x3E)
ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x65)
TEXT_DARK   = RGBColor(0x22, 0x22, 0x22)
MUTED_GRAY  = RGBColor(0x55, 0x55, 0x55)

CREAM_BG     = "F9F9F7"
HEADER_BG    = "1B4D3E"
THEMES_BG    = "F4F6F4"
BORDER_GOLD  = "C5A065"
BORDER_ACCENT= "D0D7D1"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_explicit_cell_width(cell, width_dxa):
    cell.width = width_dxa
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            color = edge_data.get('color', 'auto')
            sz = edge_data.get('sz', '4')
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def generate_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        header_img_path = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/.tmp/header_brand_unified_hd.png"
        if os.path.exists(header_img_path):
            hrun = hp.add_run()
            hrun.add_picture(header_img_path, width=Inches(7.1))
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Cabinet Delacour - Email de Demande Documentaire Pré-Audit\t\tPage ")
        frun.font.name = FONT_SANS
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = MUTED_GRAY

        f_xml = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml)
        f_xml2 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml2)
        f_xml3 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml3)
        f_xml4 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml4)

        frun2 = fp.add_run(" / ")
        frun2.font.name = FONT_SANS
        frun2.font.size = Pt(8.5)
        frun2.font.color.rgb = MUTED_GRAY

        f_xml5 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml5)
        f_xml6 = parse_xml(r'<w:r %s><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>' % nsdecls('w'))
        fp._p.append(f_xml6)
        f_xml7 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml7)
        f_xml8 = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        fp._p.append(f_xml8)
        
        tab_xml = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:pos="10512"/></w:tabs>' % nsdecls('w'))
        fp._p.get_or_add_pPr().append(tab_xml)

    # Title
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run("TEMPLATE D'EMAIL DE DEMANDE DOCUMENTAIRE PRE-AUDIT")
    r_t.font.name = FONT_SERIF
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = BRAND_GREEN

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_s = p_sub.add_run("Modèle de courriel aligné à 100% (miroir exact) avec la Checklist Préparatoire Client")
    r_s.font.name = FONT_SANS
    r_s.font.size = Pt(10.5)
    r_s.font.italic = True
    r_s.font.color.rgb = ACCENT_GOLD

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_cell = box.cell(0, 0)
    set_explicit_cell_width(b_cell, 10200)
    set_cell_background(b_cell, CREAM_BG)
    set_cell_margins(b_cell, top=140, bottom=140, left=160, right=160)
    set_cell_border(b_cell, left={'sz': 24, 'val': 'single', 'color': '1B4D3E'},
                            top={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                            bottom={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT},
                            right={'sz': 4, 'val': 'single', 'color': BORDER_ACCENT})
    cp = b_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    
    txt_email = (
        "OBJET : Préparation documentaire pré-audit — [Nom de votre entreprise]\n\n"
        "Bonjour [Prénom],\n\n"
        "Afin de préparer au mieux notre journée d'audit du [Date de l'audit] et de fluidifier notre intervention sur site sans faire perdre de temps à vos équipes, "
        "je réalise une étude documentaire préalable à distance.\n\n"
        "Vous trouverez ci-joint notre Checklist Préparatoire indicative (CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx).\n\n"
        "Pourriez-vous nous partager (par retour de mail ou via un lien Cloud/serveur partagé) les éléments habituels disponibles suivants d'ici le [Date limite] ?\n\n"
        "1. Séquence Gouvernance & Pilotage Stratégique :\n"
        "- 1.1 Organigramme à jour (ou liste des collaborateurs par service).\n"
        "- 1.2 Déclaration de politique / vision ou objectifs de l'année (si formalisés).\n\n"
        "2. Séquence Ventes, Marges & Approvisionnements :\n"
        "- 2.1 2 à 3 devis récents représentatifs avec les commandes clients associées.\n"
        "- 2.2 Liste des principaux fournisseurs et sous-traitants clés.\n\n"
        "3. Séquence Opérations, Production & Qualité :\n"
        "- 3.1 Un exemple de gamme de fabrication, consigne de travail ou mode opératoire.\n"
        "- 3.2 Suivi des réclamations clients, non-conformités ou litiges récents (si existant).\n\n"
        "4. Séquence RH, Compétences & GED :\n"
        "- 4.1 Fiches de poste existantes, grille RACI ou matrice des compétences.\n"
        "- 4.2 Vue / capture de l'arborescence de classement de vos fichiers sur votre serveur ou Cloud.\n\n"
        "Rappel important : Aucun document n'est à rédiger spécifiquement pour cet audit. Transmettez-nous uniquement vos fichiers et outils de travail quotidiens, "
        "tels qu'ils existent. Si un élément reste oral ou non formalisé, cochez simplement la case correspondante sur la checklist ci-jointe ou nous le consulterons ensemble sur écran le jour J.\n\n"
        "Merci d'avance pour votre collaboration et à votre disposition pour tout échange d'ici notre rencontre.\n\n"
        "Bien cordialement,\n\n"
        "Melchior Delacour\n"
        "Cabinet Delacour - Stratégie & Performance\n"
        "delacour.melchior@cabinetdelacour.com"
    )
    r_em = cp.add_run(txt_email)
    r_em.font.name = FONT_SANS
    r_em.font.size = Pt(8.5)

    doc.save(output_path)

def generate_html(output_path):
    html = """<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <title>Template Email Demande Documentaire Pré-Audit — Cabinet Delacour</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 30px; background-color: #fcfcfc; color: #222; line-height: 1.6; }
        h1 { font-family: Georgia, serif; color: #1B4D3E; margin-bottom: 5px; }
        .subtitle { font-style: italic; color: #C5A065; margin-bottom: 20px; }
        .box { background-color: #F9F9F7; border-left: 4px solid #1B4D3E; border: 1px solid #D0D7D1; border-left-width: 4px; border-left-color: #1B4D3E; padding: 20px; margin-bottom: 20px; border-radius: 4px; }
        ul { margin-top: 5px; margin-bottom: 15px; }
    </style>
</head>
<body>
    <h1>TEMPLATE D'EMAIL DE DEMANDE DOCUMENTAIRE PRE-AUDIT</h1>
    <div class="subtitle">Modèle aligné à 100% avec la Checklist Préparatoire Client (8 éléments en miroir exact)</div>
    
    <div class="box">
        <strong>OBJET :</strong> Préparation documentaire pré-audit — [Nom de votre entreprise]<br><br>
        Bonjour <strong>[Prénom]</strong>,<br><br>
        Afin de préparer au mieux notre journée d'audit du <strong>[Date de l'audit]</strong> et de fluidifier notre intervention sur site sans faire perdre de temps à vos équipes, je réalise une étude documentaire préalable à distance.<br><br>
        Vous trouverez ci-joint notre Checklist Préparatoire indicative (CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx).<br><br>
        Pourriez-vous nous partager (par retour de mail ou via un lien Cloud/serveur partagé) les éléments habituels disponibles suivants d'ici le <strong>[Date limite]</strong> ?<br><br>
        
        <strong>1. Séquence Gouvernance & Pilotage Stratégique :</strong>
        <ul>
            <li><strong>1.1</strong> Organigramme à jour (ou liste des collaborateurs par service).</li>
            <li><strong>1.2</strong> Déclaration de politique / vision ou objectifs de l'année (si formalisés).</li>
        </ul>

        <strong>2. Séquence Ventes, Marges & Approvisionnements :</strong>
        <ul>
            <li><strong>2.1</strong> 2 à 3 devis récents représentatifs avec les commandes clients associées.</li>
            <li><strong>2.2</strong> Liste des principaux fournisseurs et sous-traitants clés.</li>
        </ul>

        <strong>3. Séquence Opérations, Production & Qualité :</strong>
        <ul>
            <li><strong>3.1</strong> Un exemple de gamme de fabrication, consigne de travail ou mode opératoire.</li>
            <li><strong>3.2</strong> Suivi des réclamations clients, non-conformités ou litiges récents (si existant).</li>
        </ul>

        <strong>4. Séquence RH, Compétences & GED :</strong>
        <ul>
            <li><strong>4.1</strong> Fiches de poste existantes, grille RACI ou matrice des compétences.</li>
            <li><strong>4.2</strong> Vue / capture de l'arborescence de classement de vos fichiers sur votre serveur ou Cloud.</li>
        </ul>

        <em>Rappel important : Aucun document n'est à rédiger spécifiquement pour cet audit. Transmettez-nous uniquement vos fichiers et outils de travail quotidiens, tels qu'ils existent. Si un élément reste oral ou non formalisé, cochez simplement la case correspondante sur la checklist ci-jointe ou nous le consulterons ensemble sur écran le jour J.</em><br><br>

        Merci d'avance pour votre collaboration et à votre disposition pour tout échange d'ici notre rencontre.<br><br>
        Bien cordialement,<br><br>
        <strong>Melchior Delacour</strong><br>
        Cabinet Delacour - Stratégie & Performance<br>
        delacour.melchior@cabinetdelacour.com
    </div>
</body>
</html>
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)

def generate_md(output_path):
    md = """# TEMPLATE D'EMAIL DE DEMANDE DOCUMENTAIRE PRE-AUDIT
> **Cabinet Delacour - Modèle aligné à 100% avec la Checklist Préparatoire Client (8 éléments en miroir exact)**

---

### INFORMATIONS DU MESSAGE
- **OBJET** : `Préparation documentaire pré-audit — [Nom de votre entreprise]`
- **DESTINATAIRE** : `[Dirigeant / Décideur d'entreprise ou Responsable Qualité]`
- **PIÈCE JOINTE** : `CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx`

---

### CORPS DU MESSAGE (À COPIER-COLLER DIRECTEMENT) :

Bonjour **[Prénom]**,

Afin de préparer au mieux notre journée d'audit du **[Date de l'audit]** et de fluidifier notre intervention sur site sans faire perdre de temps à vos équipes, je réalise une étude documentaire préalable à distance.

Vous trouverez ci-joint notre Checklist Préparatoire indicative (`CHECKLIST_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx`).

Pourriez-vous nous partager (par retour de mail ou via un lien Cloud/serveur partagé) les éléments habituels disponibles suivants d'ici le **[Date limite]** ?

**1. Séquence Gouvernance & Pilotage Stratégique :**
- **1.1** Organigramme à jour (ou liste des collaborateurs par service).
- **1.2** Déclaration de politique / vision ou objectifs de l'année (si formalisés).

**2. Séquence Ventes, Marges & Approvisionnements :**
- **2.1** 2 à 3 devis récents représentatifs avec les commandes clients associées.
- **2.2** Liste des principaux fournisseurs et sous-traitants clés.

**3. Séquence Opérations, Production & Qualité :**
- **3.1** Un exemple de gamme de fabrication, consigne de travail ou mode opératoire.
- **3.2** Suivi des réclamations clients, non-conformités ou litiges récents (si existant).

**4. Séquence RH, Compétences & GED :**
- **4.1** Fiches de poste existantes, grille RACI ou matrice des compétences.
- **4.2** Vue / capture de l'arborescence de classement de vos fichiers sur votre serveur ou Cloud.

*Rappel important : Aucun document n'est à rédiger spécifiquement pour cet audit. Transmettez-nous uniquement vos fichiers et outils de travail quotidiens, tels qu'ils existent. Si un élément reste oral ou non formalisé, cochez simplement la case correspondante sur la checklist ci-jointe ou nous le consulterons ensemble sur écran le jour J.*

Merci d'avance pour votre collaboration et à votre disposition pour tout échange d'ici notre rencontre.

Bien cordialement,

**Melchior Delacour**  
Cabinet Delacour - Stratégie & Performance  
`delacour.melchior@cabinetdelacour.com`
"""
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(md)

def main():
    dest_dir = "/Users/delacourmelchior/Site Internet Cabinet Delacour/Projet-QHSE/OUTILS_QUALITE_DELACOUR/00_AUDIT_ET_GOUVERNANCE"
    os.makedirs(os.path.join(dest_dir, "DOCX"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "HTML"), exist_ok=True)
    os.makedirs(os.path.join(dest_dir, "MD"), exist_ok=True)

    generate_docx(os.path.join(dest_dir, "DOCX", "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "HTML", "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "MD", "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.md"))

    generate_docx(os.path.join(dest_dir, "02_CADRAGE_ET_PRE_AUDIT", "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.docx"))
    generate_html(os.path.join(dest_dir, "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.html"))
    generate_md(os.path.join(dest_dir, "TEMPLATE_EMAIL_PREPARATOIRE_PRE_AUDIT_DELACOUR.md"))

    print("✅ Email préparatoire sans émojis généré avec succès !")

if __name__ == "__main__":
    main()
